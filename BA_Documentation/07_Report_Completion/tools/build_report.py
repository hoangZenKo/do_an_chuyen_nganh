from pathlib import Path
from collections import Counter
from copy import deepcopy
import hashlib
import json
import re

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.table import Table
from PIL import Image

ROOT = Path('/Users/hoang/Rental_room_management')
WORK = ROOT / 'BA_Documentation/07_Report_Completion'
SOURCE = Path('/Users/hoang/Desktop/ĐỒ ÁN/Bao_cao_Quan_ly_Nha_tro_V3_2026-10-28.docx')
DESTINATION = SOURCE.with_name('Bao_cao_Quan_ly_Nha_tro_V4_2026-10-29_BA.docx')
DATA = json.loads((WORK / 'data/business-report-content.json').read_text())
CATALOG = json.loads((ROOT / 'BA_Documentation/02_Functional_Decomposition_Use_Case_Catalog/data/catalog.json').read_text())
DIAGRAMS = json.loads((WORK / 'data/report-diagrams.json').read_text())
CASES = {item['id']: item for item in DATA['useCases']}
GROUPS = {item['id']: item for item in DATA['groups']}
ACTORS = DATA['actors']
SHORT_STATUS = {'Xác minh tĩnh': 'Đã đối chiếu', 'Một phần': 'Một phần', 'Đã triển khai — DEMO': 'Thử nghiệm', 'Đã triển khai — MÔ PHỎNG': 'Mô phỏng'}
REDIRECTS = {
    'UC-SYS-04': ['UC-SYS-03'], 'UC-DIS-06': ['UC-DIS-05', 'UC-DIS-06.2'],
    'UC-DIS-11': ['UC-TNT-05', 'UC-TNT-12.2', 'UC-TNT-12.3'],
    'UC-DIS-13': ['UC-DIS-13.1', 'UC-DIS-13.2'],
    'UC-OPT-04.2': ['UC-OPT-04.1'], 'UC-TNT-13': ['UC-TNT-04'],
    'UC-BIL-08.2': ['UC-BIL-04', 'UC-TEN-03'], 'UC-BIL-12': ['UC-BIL-12.1', 'UC-BIL-12.2'],
    'UC-BIL-18': ['UC-BIL-18.1', 'UC-BIL-18.2', 'UC-BIL-18.3'],
    'UC-BIL-21': ['UC-BIL-20'], 'UC-OPS-08': ['UC-BIL-18.3', 'UC-SUP-05.2'],
    'UC-TEN-07': ['UC-TEN-07.1', 'UC-TEN-07.2', 'UC-TEN-07.3'],
    'UC-TEN-08': ['UC-TEN-08.1', 'UC-TEN-08.2'],
    'UC-TEN-09': ['UC-TEN-09.2', 'UC-TEN-09.3', 'UC-SUP-04'],
    'UC-ADM-04.2': ['UC-ADM-04.1'], 'UC-ADM-07': ['UC-ADM-07.1', 'UC-SYS-03'],
    'UC-ADM-13': ['UC-ADM-13.1', 'UC-ADM-13.2', 'UC-ADM-13.3'],
    'UC-ADM-17': ['UC-ADM-17.1', 'UC-ADM-17.2'],
    'UC-SUP-02': ['UC-SUP-02.1', 'UC-SUP-02.2'],
    'UC-SUP-03': ['UC-SUP-01', 'UC-SUP-01.1', 'UC-SUP-03.2', 'UC-SUP-03.3'],
    'UC-SUP-05': ['UC-SUP-05.1', 'UC-SUP-05.2'],
}
SCOPE_CHANGES = {
    'UC-ADM-11': ['UC-ADM-11', 'UC-ADM-11.1'],
    'UC-ADM-12': ['UC-ADM-12', 'UC-ADM-12.1'],
    'UC-ADM-15': ['UC-ADM-15', 'UC-ADM-15.1'],
    'UC-ADM-18': ['UC-ADM-18', 'UC-ADM-18.1'],
    'UC-SUP-01': ['UC-SUP-01', 'UC-SUP-01.1'],
}

class Report:
    def __init__(self):
        self.doc = Document(SOURCE)
        self.original = list(self.doc.element.body)
        self.anchor = None
        self.counter = 10000
        if 'Heading 5' not in self.doc.styles:
            self.doc.styles.add_style('Heading 5', WD_STYLE_TYPE.PARAGRAPH)
        self.doc.styles['Heading 5'].base_style = self.doc.styles['Heading 4']
        self.doc.styles['Heading 5'].font.size = Pt(12)
        self.doc.styles['Heading 5'].font.bold = True
        self.doc.styles['Heading 5'].font.italic = False
        if 'BA Heading 4' not in self.doc.styles:
            self.doc.styles.add_style('BA Heading 4', WD_STYLE_TYPE.PARAGRAPH)
        self.doc.styles['BA Heading 4'].base_style = self.doc.styles['Heading 4']
        self.doc.styles['BA Heading 4'].font.bold = True
        self.doc.styles['BA Heading 4'].font.italic = False
        if 'Caption' not in self.doc.styles:
            self.doc.styles.add_style('Caption', WD_STYLE_TYPE.PARAGRAPH)
        self.doc.styles['Caption'].base_style = self.doc.styles['Normal']
        self.original_hash = hashlib.sha256(SOURCE.read_bytes()).hexdigest()

    def move(self, element):
        self.anchor.addprevious(element)
        return element

    def paragraph(self, text='', style=None, size=None, bold=None, before=False, center=False):
        para = self.doc.add_paragraph(style=style)
        if text:
            run = para.add_run(text)
            if bold is not None:
                run.bold = bold
            if size:
                run.font.size = Pt(size)
        para.paragraph_format.space_after = Pt(6)
        if before:
            para.paragraph_format.page_break_before = True
        if center:
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if style and style.startswith('Heading'):
            para.paragraph_format.keep_with_next = True
        self.move(para._p)
        return para

    def heading(self, text, level, before=False):
        para = self.paragraph(text, 'BA Heading 4' if level == 4 else 'Heading ' + str(level), before=before)
        para.paragraph_format.keep_with_next = True
        return para

    def table(self, headings, rows, widths, size=11.5):
        table = self.doc.add_table(rows=1, cols=len(headings))
        table.style = 'Table Grid'
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        for column, width in zip(table.columns, widths):
            column.width = Inches(width)
        header = table.rows[0]
        repeat = OxmlElement('w:tblHeader')
        header._tr.get_or_add_trPr().append(repeat)
        for index, title in enumerate(headings):
            header.cells[index].text = title
        for row in rows:
            cells = table.add_row().cells
            for index, value in enumerate(row):
                cells[index].text = str(value)
        for row_index, row in enumerate(table.rows):
            for cell_index, cell in enumerate(row.cells):
                cell.width = Inches(widths[cell_index])
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
                margins = OxmlElement('w:tcMar')
                for side, amount in [('top',70),('bottom',70),('left',80),('right',80)]:
                    edge = OxmlElement('w:' + side)
                    edge.set(qn('w:w'), str(amount))
                    edge.set(qn('w:type'), 'dxa')
                    margins.append(edge)
                cell._tc.get_or_add_tcPr().append(margins)
                for para in cell.paragraphs:
                    para.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    para.paragraph_format.first_line_indent = Pt(0)
                    para.paragraph_format.keep_with_next = False
                    para.paragraph_format.keep_together = False
                    para.paragraph_format.widow_control = True
                    para.paragraph_format.line_spacing = 1.08
                    para.paragraph_format.space_after = Pt(3)
                    para.paragraph_format.space_before = Pt(0)
                    for run in para.runs:
                        run.font.name = 'Times New Roman'
                        run.font.size = Pt(size)
                        run.bold = row_index == 0 or (len(headings) == 2 and cell_index == 0)
                        run.font.color.rgb = RGBColor(0,0,0)
        self.move(table._tbl)
        self.paragraph('', size=3).paragraph_format.space_after = Pt(1)
        return table

    def replace(self, start, stop, builder):
        self.anchor = self.original[start]
        builder()
        for element in self.original[start:stop + 1]:
            if element.getparent() is not None:
                element.getparent().remove(element)

    def set_original_text(self, index, text):
        Paragraph(self.original[index], self.doc).text = text

    def requirements(self):
        self.paragraph('Phạm vi hiện trạng gồm 162 Use Case thuộc 16 nhóm nghiệp vụ và 5 tác nhân con người. Mỗi Use Case thể hiện một mục tiêu của người sử dụng; các thao tác nhập, xem trước, xác nhận và ghi nhận được tổ chức trong cùng mục tiêu khi phù hợp. Danh mục, đặc tả và sơ đồ được thống nhất theo cùng mã Use Case.')
        self.table(['Nhóm', 'Nội dung nghiệp vụ', 'Số UC'], [(g['id'],g['name'],g['ucCount']) for g in DATA['groups']], [.65,4.85,.75])
        self.heading('2.1.1 Tác nhân và phạm vi truy cập',3)
        descriptions = {
            'A01': 'Chưa đăng nhập; tìm phòng công khai, đăng ký, đăng nhập, nhận lời mời hoặc khôi phục tài khoản theo điều kiện tương ứng.',
            'A02': 'Quản lý hoạt động cho thuê trong phạm vi của mình; giao việc và cấp quyền cho nhân viên. Không tự có quyền quản trị toàn nền tảng.',
            'A03': 'Thực hiện công việc được chủ trọ giao, theo quyền nghiệp vụ và khu được phân công. Không mặc nhiên có toàn bộ quyền của chủ trọ.',
            'A04': 'Tìm phòng sau khi đăng ký; chỉ sử dụng dữ liệu thuê của mình khi có hồ sơ hoặc hợp đồng được liên kết. Đây là hai bối cảnh của cùng vai trò, không phải hai tài khoản quyền tách biệt.',
            'A05': 'Quản trị tài khoản, nội dung và hoạt động nền tảng theo bộ quyền được cấp. Không kế thừa quyền chủ trọ; các nghiệp vụ nhạy cảm có điều kiện xác minh hoặc phân tách người thực hiện.',
        }
        self.table(['Mã', 'Tác nhân', 'Phạm vi tham gia'], [(key,ACTORS[key],descriptions[key]) for key in ACTORS], [.55,1.5,4.2])
        self.paragraph('Một người có thể tham gia các bối cảnh nghiệp vụ khác nhau; báo cáo không suy ra quyền mới chỉ từ việc người đó đồng thời là người thuê hoặc chủ trọ ngoài thực tế. Quyền sử dụng được xác định theo tài khoản, phân công, quan hệ với hồ sơ và điều kiện của từng Use Case.')
        self.paragraph('Các tác nhân phối hợp chỉ tham gia khi chức năng có sử dụng dịch vụ tương ứng. Không xem việc có kết nối hoặc thông tin thiết lập là bằng chứng đã giao nhận thành công ngoài thực tế.')
        self.table(['Mã', 'Tác nhân phối hợp', 'Vai trò'], [
            ('E01','Đơn vị gửi thư hoặc thông báo','Chuyển lời mời, liên kết khôi phục hoặc thông báo theo kênh đã thiết lập.'),
            ('E02','Đơn vị cung cấp Zalo ZNS','Giao thông báo qua Zalo ZNS khi sử dụng kênh này.'),
            ('E03','Đơn vị kiểm tra an toàn tệp','Kiểm tra tệp hỗ trợ trước khi được tiếp nhận.'),
            ('E04','Dịch vụ trợ lý bằng trí tuệ nhân tạo','Hỗ trợ trả lời khi được bật; có phương án trả lời thay thế.'),
            ('E05','Đơn vị lưu giữ kết quả báo cáo','Lưu kết quả khi sử dụng phương thức lưu giữ ngoài hệ thống.'),
        ], [.55,2.1,3.6])
        self.heading('2.1.2 Thuật ngữ và quy tắc nghiệp vụ',3)
        self.table(['Thuật ngữ', 'Ý nghĩa sử dụng trong báo cáo'], [
            ('Hồ sơ người thuê','Thông tin quản lý khách thuê; không đồng nghĩa đã có tài khoản truy cập.'),
            ('Tài khoản người thuê','Thông tin đăng nhập; chỉ truy cập dữ liệu thuê khi được liên kết đúng hồ sơ.'),
            ('Khoản thu và phân bổ','Khoản tiền được ghi nhận cho người thuê và phần tiền được sử dụng để thanh toán hóa đơn; không cộng hai lần cùng một khoản.'),
            ('Quyết toán trả phòng','Đối chiếu nghĩa vụ cuối kỳ, tiền cọc, số tiền phải thu hoặc hoàn; xác nhận quyết toán khác ghi nhận đã hoàn tiền.'),
            ('Thông tin lần cập nhật','Căn cứ kiểm tra người dùng đang thao tác trên thông tin mới nhất; khi hồ sơ đã đổi, phải xem lại trước khi xác nhận.'),
            ('Đã đối chiếu','Hành vi đã được đối chiếu với hiện trạng triển khai; không đồng nghĩa đã nghiệm thu đầy đủ trong sử dụng thực tế.'),
            ('Một phần','Mục tiêu đã được đáp ứng một phần hoặc còn phụ thuộc giao nhận, dịch vụ chưa được xác nhận thực tế.'),
            ('Thử nghiệm và mô phỏng','Thanh toán thẻ chỉ thử nghiệm; trao đổi Zalo chỉ mô phỏng. Không diễn giải thành thu tiền hoặc gửi tin thực.'),
        ],[1.65,4.6])
        self.table(['Quy tắc chung', 'Nội dung'], [
            ('Phạm vi truy cập','Chủ trọ và nhân viên chỉ sử dụng dữ liệu thuộc phạm vi được giao; người thuê chỉ sử dụng dữ liệu gắn với mình.'),
            ('Không xóa lịch sử tùy ý','Lưu trữ, vô hiệu hóa và xóa là các hành vi khác nhau; chỉ xóa khi đáp ứng các điều kiện của từng loại hồ sơ.'),
            ('Phân tách quyết định','Xét duyệt và thực thi yêu cầu vòng đời thuê là hai mục tiêu; người lập và người duyệt đối soát phải độc lập; xét kháng nghị không do người quyết định ban đầu tự thực hiện.'),
            ('Ghi nhận tài chính','Không coi chứng từ gửi lên là đã thu tiền; không coi lệnh hoàn tiền hoặc kết quả đối chiếu là bằng chứng chuyển tiền thực tế.'),
            ('Ngừng nhận tệp chưa an toàn','Tệp hỗ trợ không được nhận khi việc kiểm tra an toàn chưa thực hiện được hoặc không đạt.'),
        ], [1.65,4.6])

    def scenarios(self):
        self.heading('2.4.1 Danh mục và đặc tả Use Case',3)
        self.paragraph('Phần này mô tả 162 mục tiêu nghiệp vụ trong 16 nhóm. Mã Use Case được giữ để liên kết danh mục, đặc tả, sơ đồ và bảng phân quyền; tiền tố mã kế thừa lịch sử đánh số, không quyết định nhóm nghiệp vụ hiện tại.')
        self.paragraph('Mỗi Use Case được trình bày bằng một bảng liên tục gồm 20 thuộc tính. B1, B2… là các bước của luồng chính; A1, A2… là luồng thay thế hợp lệ; E1, E2… là ngoại lệ. Mỗi nhánh nêu bước phát sinh và kết quả xử lý. Các điều kiện đã đăng nhập hoặc đã được cấp quyền không được hiểu là các bước nghiệp vụ bắt buộc phải thực hiện lại trong mọi lần sử dụng.')
        self.paragraph('Mức hoàn thiện gồm 151 Use Case đã đối chiếu hiện trạng, 9 Use Case đáp ứng một phần, 1 Use Case thanh toán thẻ thử nghiệm và 1 Use Case trao đổi Zalo mô phỏng. Các giới hạn được giữ ngay trong đặc tả để không trình bày thử nghiệm thành chức năng đã sẵn sàng vận hành đầy đủ.')
        self.heading('Danh mục Use Case theo nhóm nghiệp vụ',4)
        for group in DATA['groups']:
            self.heading(group['id'] + ' ' + group['name'],4)
            records = [u for u in DATA['useCases'] if u['groupId'] == group['id']]
            self.table(['Mã UC','Tên Use Case','Tác nhân chính','Mức hoàn thiện'], [(u['id'],u['title'],'; '.join(ACTORS[a] for a in u['primaryActorIds']),SHORT_STATUS[u['status']]) for u in records], [1.25,2.45,1.65,.9],size=11)
        for group in DATA['groups']:
            self.heading('Đặc tả ' + group['id'] + ' ' + group['name'],4,before=True)
            for usecase in [u for u in DATA['useCases'] if u['groupId'] == group['id']]:
                self.heading(usecase['id'] + ' ' + usecase['title'],5)
                self.table(['Thuộc tính', 'Đặc tả nghiệp vụ ' + usecase['id']], list(usecase['fields'].items()), [1.55,4.7],size=12)

    def diagrams(self):
        self.heading('2.4.2 Mô hình Use Case',3,before=True)
        self.paragraph('Sơ đồ tổng quan trình bày 12 mục tiêu đại diện, không thay thế toàn bộ danh mục. Các sơ đồ theo nhóm bao phủ đủ 162 Use Case. Một số nhóm được chia thành nhiều hình để dễ đọc; việc chia hình không tạo thêm Use Case hoặc thay đổi ranh giới nghiệp vụ.')
        self.paragraph('Đường nối biểu thị tác nhân tham gia mục tiêu, không thể hiện trình tự xử lý. Đường có nhãn giới hạn chỉ cho phép phần việc được ghi rõ, không phải toàn bộ Use Case. Điều kiện quyền, khu được giao và hồ sơ của chính mình phải đọc cùng đặc tả và ma trận ở Phụ lục B. Không giả định quản trị viên kế thừa quyền chủ trọ.')
        for entry in DIAGRAMS:
            title = entry['title'].replace(' · ', ' ').replace('MFA','xác thực nhiều yếu tố').replace('SLA','thời hạn hỗ trợ').replace('Admin','quản trị')
            self.heading(title,4,before=True)
            file = WORK / 'diagrams' / entry['file'].replace('.puml','.png')
            width,height = Image.open(file).size
            maximum_height = 8.0 if entry['kind'] == 'overall' else 7.8
            factor = min(6.25 / width, maximum_height / height)
            para = self.paragraph(center=True)
            para.paragraph_format.keep_with_next = True
            para.paragraph_format.space_after = Pt(2)
            para.add_run().add_picture(str(file),width=Inches(width*factor),height=Inches(height*factor))
            for shape in para._p.xpath('.//wp:docPr'):
                shape.set('descr',title + '; ' + ', '.join(entry['ucIds']))
            self.paragraph('Hình 2.0 ' + title,style='Caption',size=11,center=True)
            self.paragraph('Use Case trong hình: ' + ', '.join(entry['ucIds']) + '.',size=10.5)

    def appendix_a(self):
        self.heading('Phụ lục A Đối chiếu danh mục Use Case',1,before=True)
        self.paragraph('Lần hoàn thiện này thay phần Use Case của báo cáo ngày 28/10/2026 gồm 131 mục thuộc 9 nhóm bằng danh mục chuẩn hóa 162 mục thuộc 16 nhóm. Đây là thay đổi cách xác định mục tiêu và phạm vi dựa trên bộ phân tích hiện trạng; chênh lệch số lượng không được hiểu là bổ sung 31 chức năng mới cho ứng dụng.')
        self.table(['Nội dung','Trước hoàn thiện','Sau hoàn thiện'],[
            ('Tác nhân con người','4 tác nhân; chưa đưa nhân viên vào phạm vi.','5 tác nhân; bổ sung nhân viên theo quyền và khu được giao.'),
            ('Danh mục','131 mục thuộc 9 nhóm.','162 Use Case thuộc 16 nhóm nghiệp vụ.'),
            ('Đặc tả','Có mô tả theo cách triển khai, một số mục tiêu gộp hoặc tách chưa phù hợp.','Một bảng 20 thuộc tính mỗi Use Case; mô tả mục tiêu, điều kiện, luồng và kết quả bằng ngôn ngữ nghiệp vụ.'),
            ('Sơ đồ và phân quyền','Theo danh mục cũ.','Đồng bộ với mã mới; có ma trận tác nhân, không gán quyền toàn phần cho phạm vi tham gia giới hạn.'),
            ('Giới hạn sử dụng','Dễ nhầm thanh toán thử và trao đổi mô phỏng với hoạt động thật.','Giữ riêng 151 đã đối chiếu, 9 một phần, 1 thanh toán thử và 1 trao đổi mô phỏng.'),
        ],[1.2,2.3,2.75])
        self.heading('A.1 Đối chiếu 131 mục của bản báo cáo trước',2)
        self.paragraph('Bảng này hướng dẫn tìm nội dung tương ứng, không bảo lưu nguyên đặc tả cũ. Mục có nhiều mã đích đã được phân định lại mục tiêu; các bước nhập, xem trước, xác nhận, xuất kết quả hoặc duy trì đăng nhập không mặc nhiên là Use Case độc lập.')
        old = []
        for element in self.original[247:640]:
            if element.tag == qn('w:p'):
                para = Paragraph(element,self.doc)
                if re.match(r'^UC-',para.text):
                    old.append((para.text.split()[0],para.text.partition(' ')[2]))
        rows=[]
        for identifier,title in old:
            targets = SCOPE_CHANGES.get(identifier) or REDIRECTS.get(identifier) or [identifier]
            assert all(target in CASES for target in targets),identifier
            action = 'Giữ mã, rà soát lại mục tiêu và quyền.' if targets == [identifier] else ('Phân định các mục tiêu hoặc phạm vi riêng.' if len(targets)>1 else 'Đưa về mục tiêu nghiệp vụ tương ứng; không lập mục độc lập.')
            if identifier == 'UC-SYS-04':
                action = 'Duy trì đăng nhập là điều kiện sử dụng, không phải mục tiêu nghiệp vụ độc lập.'
            rows.append((identifier,', '.join(targets),action))
        self.table(['Mã trước','Mã hiện tại','Hướng dẫn đối chiếu'],rows,[1.1,2.4,2.75],size=11)
        self.heading('A.2 Đối chiếu 67 mục của bản thảo ban đầu',2)
        self.paragraph('Giữ dấu vết đối chiếu bản thảo ban đầu; cột mã hiện tại đã được chuyển về danh mục 162 Use Case. Các mã ở cột đầu là mã lịch sử, không dùng để xác định chức năng còn hiệu lực.')
        old_table=Table(self.original[1162],self.doc)
        rows=[]
        for row in old_table.rows[1:]:
            target_ids=[]
            for identifier in re.findall(r'UC-[A-Z]+-\d+(?:\.\d+)*',row.cells[2].text):
                target_ids.extend(SCOPE_CHANGES.get(identifier) or REDIRECTS.get(identifier) or [identifier])
            target_ids=list(dict.fromkeys(target_ids))
            assert all(target in CASES for target in target_ids),target_ids
            rows.append((row.cells[0].text,row.cells[1].text,', '.join(target_ids)))
        self.table(['Mã bản thảo','Mục tiêu trong bản thảo','Mã hiện tại'],rows,[1.1,2.6,2.55],size=11)

    def appendix_b(self):
        self.heading('Phụ lục B Ma trận tác nhân và Use Case',1,before=True)
        self.paragraph('P: tác nhân chính. P*: tác nhân chính có điều kiện về quyền, sở hữu hoặc trạng thái. V: chỉ tham gia biến thể giới hạn, không thực hiện toàn bộ mục tiêu. B: bị chặn thực hiện. Dấu gạch ngang: không xác định là tác nhân thực hiện mục tiêu này; không dùng để suy ra quyền.')
        self.paragraph('A01: Khách vãng lai; A02: Chủ trọ; A03: Nhân viên vận hành; A04: Người tìm phòng hoặc người thuê; A05: Quản trị viên nền tảng. Điều kiện chi tiết được nêu tại mục Phân quyền và bảo mật trong từng đặc tả. Ma trận không thay thế việc xét điều kiện của từng lần thực hiện.')
        catalog_cases={u['id']:u for u in CATALOG['useCases']}
        for group in DATA['groups']:
            self.heading(group['id'] + ' ' + group['name'],3)
            rows=[]
            for usecase in [u for u in DATA['useCases'] if u['groupId']==group['id']]:
                matrix=catalog_cases[usecase['id']]['matrix']
                symbols=[matrix[a].split(' — ')[0] if matrix[a] != '—' else '—' for a in ACTORS]
                rows.append([usecase['id'],*symbols])
            self.table(['Mã UC',*ACTORS.keys()],rows,[1.5,.95,.95,.95,.95,.95],size=11)
        self.heading('B.1 Tác nhân phối hợp theo từng Use Case',2)
        self.paragraph('Chỉ những Use Case có tác nhân phối hợp trong danh mục chuẩn hóa được liệt kê. Sự tham gia có thể tùy kênh hoặc phương thức được chọn; không tự xác nhận đã giao nhận thành công thực tế.')
        self.table(['Mã UC','Mục tiêu','Tác nhân phối hợp'],[(u['id'],u['title'],', '.join(u['supportingActorIds'])) for u in DATA['useCases'] if u['supportingActorIds']],[1.05,3.6,1.6],size=11)

    def appendix_c(self):
        self.heading('Phụ lục C Phạm vi đối chiếu và tài liệu kèm theo',1,before=True)
        self.paragraph('Phần Use Case được hoàn thiện dựa trên bộ tài liệu phân tích hiện trạng, danh mục đã chuẩn hóa, sơ đồ và đặc tả chi tiết cập nhật ngày 29/10/2026. Các bảng trong báo cáo sử dụng cách diễn đạt nghiệp vụ; bằng chứng triển khai được giữ riêng để đối chiếu khi cần, không đưa vào các bước xử lý của người dùng.')
        self.table(['Tài liệu','Mục đích'],[
            ('Phân tích hiện trạng','Phạm vi chức năng đã xác minh, chức năng một phần và các điểm còn chưa đủ căn cứ.'),
            ('Phân rã chức năng và danh mục Use Case','Nguồn thống nhất mã, nhóm mục tiêu và tác nhân.'),
            ('Bộ sơ đồ Use Case','Đối chiếu tác nhân, phạm vi tham gia và mức hoàn thiện.'),
            ('Bộ đặc tả chi tiết Use Case','Căn cứ soát điều kiện, luồng nghiệp vụ, ngoại lệ và giới hạn.'),
            ('Tài nguyên hoàn thiện báo cáo','Bản nội dung nghiệp vụ, bảng đối chiếu tên và kết quả kiểm tra tính nhất quán.'),
        ],[2.15,4.1])
        self.paragraph('Đợt hoàn thiện này không sửa ứng dụng và không tạo thêm kết quả nghiệm thu. Các kết quả kiểm thử trong mục 2.6 được giữ theo đợt thực hiện đã ghi ở đó; không dùng việc hoàn thiện tài liệu để khẳng định mọi giao diện, giao nhận bên ngoài hoặc thanh toán thực tế đã được kiểm tra thành công.')
        self.paragraph('Các chương khác, mô hình lớp, sơ đồ tương tác và ảnh giao diện được giữ theo báo cáo trước. Việc chuẩn hóa phần Use Case không được hiểu là đã dựng lại hoặc nghiệm thu mới các phần đó. Những chỗ tham chiếu phạm vi tác nhân và số lượng Use Case được cập nhật để tránh mâu thuẫn.')
        self.paragraph('Thông tin cần người thực hiện bổ sung trước khi nộp: họ tên, mã số sinh viên, lớp và giảng viên hướng dẫn. Báo cáo không tự suy đoán các thông tin này.')

    def bookmark(self,para,name):
        self.counter+=1
        start=OxmlElement('w:bookmarkStart'); start.set(qn('w:id'),str(self.counter)); start.set(qn('w:name'),name)
        end=OxmlElement('w:bookmarkEnd'); end.set(qn('w:id'),str(self.counter))
        para._p.insert(0,start); para._p.append(end)

    def navigation(self):
        for element in self.original[15:56] + self.original[58:153]:
            if element.getparent() is not None:
                element.getparent().remove(element)
        pagefile=WORK / 'data/page-map.json'
        pages=json.loads(pagefile.read_text()) if pagefile.exists() else {}
        headings=[]; captions=[]; figure=0
        for para in list(self.doc.paragraphs):
            if re.match(r'^Hình 2\.\d+ ',para.text):
                figure+=1
                para.text=re.sub(r'^Hình 2\.\d+', 'Hình 2.'+str(figure),para.text)
                name='ba_figure_'+str(figure); self.bookmark(para,name); captions.append((para.text,name,0))
            if para.style.name in ['Heading 1','Heading 2','Heading 3']:
                name='ba_section_'+str(len(headings)+1);self.bookmark(para,name)
                headings.append((para.text,name,int(para.style.name[-1])-1))
        self.anchor=self.original[56]
        for title,name,level in headings:
            self.navigation_line(title,name,level,pages)
        self.anchor=self.original[153]
        for title,name,level in captions:
            self.navigation_line(title,name,0,pages,size=11)
        (WORK / 'data/navigation.json').write_text(json.dumps({'headings':headings,'captions':captions},ensure_ascii=False,indent=2))

    def navigation_line(self,title,name,level,pages,size=12):
        para=self.paragraph(size=size)
        para.paragraph_format.left_indent=Inches(level*.15)
        para.paragraph_format.space_after=Pt(3)
        para.paragraph_format.line_spacing=1.0
        para.paragraph_format.tab_stops.add_tab_stop(Inches(6.25),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
        link=OxmlElement('w:hyperlink');link.set(qn('w:anchor'),name)
        run=OxmlElement('w:r'); props=OxmlElement('w:rPr');fs=OxmlElement('w:sz');fs.set(qn('w:val'),str(size*2));props.append(fs);run.append(props)
        text=OxmlElement('w:t');text.text=title;run.append(text);link.append(run);para._p.append(link)
        para.add_run('\t'+str(pages.get(title,'…'))).font.size=Pt(size)

    def main(self):
        for element in self.original[739:1156]:
            for node in element.xpath('.//w:t'):
                node.text = re.sub(r'UC-[A-Z]+-\d+(?:\.\d+)*',lambda match: ', '.join(REDIRECTS.get(match.group(0), ['UC-ADM-04.1'] if match.group(0)=='UC-ADM-04' else [match.group(0)])),node.text or '')
        self.set_original_text(173,'Phạm vi nghiệp vụ gồm năm tác nhân con người: Khách vãng lai, Chủ trọ, Nhân viên vận hành, Người tìm phòng hoặc người thuê và Quản trị viên nền tảng. Người tìm phòng và người đang thuê là hai bối cảnh của cùng vai trò; quyền xem dữ liệu thuê chỉ có khi được liên kết đúng hồ sơ. Nhân viên thực hiện công việc theo quyền và khu được chủ trọ giao; quyền quản trị không thay thế quyền vận hành của chủ trọ.')
        self.set_original_text(199,'Phạm vi báo cáo gồm năm tác nhân đã xác định. Nhân viên chỉ sử dụng chức năng được cấp và khu được giao; người tìm phòng chưa có quan hệ thuê không được xem dữ liệu thuê của người khác. Chủ trọ và quản trị viên phải đáp ứng điều kiện thẩm quyền, sở hữu, trạng thái và tính độc lập của từng nghiệp vụ.')
        self.set_original_text(958,'Các ảnh trong phần này minh họa bốn vai trò đã có ảnh ở báo cáo trước. Chưa bổ sung bộ ảnh riêng cho Nhân viên vận hành; ảnh minh họa không thay thế ma trận phân quyền và không chứng minh mọi Use Case đã được nghiệm thu sử dụng thực tế.')
        self.replace(178,189,self.requirements)
        self.replace(217,639,self.scenarios)
        self.replace(640,738,self.diagrams)
        self.replace(1156,1163,self.appendix_a)
        self.replace(1164,1167,self.appendix_b)
        self.replace(1168,1172,self.appendix_c)
        self.navigation()
        self.doc.core_properties.modified = __import__('datetime').datetime(2026,10,29,12,0,0)
        self.doc.save(DESTINATION)
        assert hashlib.sha256(SOURCE.read_bytes()).hexdigest()==self.original_hash
        print(str(DESTINATION))
        print('Tables:',len(self.doc.tables),'Diagrams:',len(DIAGRAMS))

if __name__ == '__main__':
    Report().main()
