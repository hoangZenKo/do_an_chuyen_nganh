from pathlib import Path
from zipfile import ZipFile
from collections import Counter
from datetime import datetime, timezone
import hashlib
import json
import re
from docx import Document
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[3]
BA = ROOT / 'BA_Documentation'
WORK = BA / '09_Code_Documentation_Audit'

def load(path):
    return json.loads(path.read_text())

def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()

checks = []
def check(name, passed, detail):
    checks.append({'name': name, 'passed': bool(passed), 'detail': detail})

catalog = load(BA / '02_Functional_Decomposition_Use_Case_Catalog/data/catalog.json')
old_catalog = load(WORK / 'data/baselines/catalog-before-audit.json')
trace = load(BA / '02_Functional_Decomposition_Use_Case_Catalog/data/traceability.json')
content = load(BA / '07_Report_Completion/data/business-report-content.json')
old_content = load(WORK / 'data/baselines/business-report-content-before-audit.json')
corrections = load(WORK / 'data/word-corrections.json')
images = load(BA / '08_Use_Case_Black_White/data/image-manifest.json')
code = load(WORK / 'data/independent-code-audit.json')
probes = load(WORK / 'data/semantic-probes.json')
original_protected = load(WORK / 'data/baselines/protected-hashes-stage4-original.json')
check('catalog_ids_groups_actors_status_preserved', [(item['id'], item['businessGroup'], item['primaryActorIds'], item['supportingActorIds'], item['status']) for item in catalog['useCases']] == [(item['id'], item['businessGroup'], item['primaryActorIds'], item['supportingActorIds'], item['status']) for item in old_catalog['useCases']], '162 UC; không tự nâng trạng thái xác minh hoặc đổi tác nhân')
source_changes = [file for file, expected in trace['sourceHashes'].items() if not (ROOT / file).exists() or digest(ROOT / file) != expected]
check('application_source_integrity', len(trace['sourceHashes']) == 846 and not source_changes, {'files': len(trace['sourceHashes']), 'changed': source_changes})
check('independent_code_audit', code['totals']['useCases'] == 162 and not code['missing'] and not code['aclMismatches'] and not code['hashChanges'] and all(anchor['lineWithinDeclaration'] for case in code['useCases'] for anchor in case['anchors']) and all(ref['reachable'] for case in code['useCases'] for ref in case['serviceReachability']), code['totals'])
check('targeted_semantic_reproductions', probes['passed'] == 6 and all(item['passed'] for item in probes['results']), probes['scope'])
new_fields = {case['id']: case['fields'] for case in content['useCases']}
old_fields = {case['id']: case['fields'] for case in old_content['useCases']}
actual_changes = [{'id': identifier, 'field': field, 'before': old_fields[identifier][field], 'after': value} for identifier, fields in new_fields.items() for field, value in fields.items() if old_fields[identifier][field] != value]
check('five_uc_twenty_three_field_amendments', len(actual_changes) == 23 and actual_changes == corrections['changedFields'], corrections['changedUseCases'])
word = Document(corrections['output'])
word_fields = {}
for table in word.tables:
    if len(table.rows) == 21 and table.rows[0].cells[0].text == 'Thuộc tính':
        word_fields[table.rows[1].cells[1].text] = {row.cells[0].text: row.cells[1].text for row in table.rows[1:]}
check('word_all_162_uc_tables', len(word_fields) == 162 and word_fields == new_fields and all(len(fields) == 20 for fields in word_fields.values()), '162 bảng, mỗi bảng 20 thuộc tính khớp dữ liệu đã đính chính')
check('source_word_preserved_and_final_hash', digest(Path(corrections['source'])) == corrections['sourceSha256'] and digest(Path(corrections['output'])) == corrections['outputSha256'], {'source': corrections['source'], 'output': corrections['output']})
jargon = [{'id': identifier, 'field': field} for identifier, fields in new_fields.items() for field, text in fields.items() if re.search(r'\b(?:Controller|Service|DTO|JWT|API|endpoint|backend|frontend|Prisma|transaction|repository)\b', text, re.I)]
check('report_uc_business_language', not jargon, {'technical_terms_found': jargon, 'scope': 'Kiểm tra từ khóa hỗ trợ; không thay thế đọc biên tập'})
replaced = []
for index, paragraph in enumerate(word.paragraphs):
    drawings = paragraph._p.xpath('.//w:drawing')
    if not drawings:
        continue
    caption = next((candidate.text for candidate in word.paragraphs[index + 1:index + 4] if candidate.text.startswith('Hình 2.')), '')
    match = re.match(r'Hình 2\.(\d+)\b', caption)
    if match and 3 <= int(match.group(1)) <= 32:
        entry = images[int(match.group(1)) - 3]
        relation = drawings[0].xpath('.//a:blip')[0].get(qn('r:embed'))
        blob = word.part.related_parts[relation].blob
        replaced.append(hashlib.sha256(blob).hexdigest() == digest(BA / '08_Use_Case_Black_White' / entry['png']))
check('thirty_black_white_diagrams_embedded', len(replaced) == 30 and all(replaced), {'verified_images': len(replaced)})
coverage = Counter(identifier for entry in images if entry['kind'] != 'overall' for identifier in entry['ucIds'])
check('module_diagrams_cover_162_once', len(images) == 30 and set(coverage) == set(new_fields) and set(coverage.values()) == {1}, {'image_count': len(images), 'module_uc_count': len(coverage)})
source_image_checks = []
for entry in images:
    candidates = [BA / '07_Report_Completion/diagrams' / entry['file'], BA / '05_Use_Case_Diagrams' / entry['file']]
    source_image_checks.append(any(path.exists() and digest(path) == entry['sourceSha256'] for path in candidates))
check('diagram_sources_match_image_manifest', all(source_image_checks), {'matched': sum(source_image_checks)})
links = [(entry['id'], link['actorId'], link['ucId'], link['limited'], link['label']) for entry in images for link in entry['links']]
diagram_audit = load(WORK / 'data/diagram-code-audit.json')
proof_links = [(entry['diagram'], entry['actorId'], entry['ucId'], entry['limited'], entry['label']) for entry in diagram_audit['checks']]
check('diagram_association_evidence', Counter(links) == Counter(proof_links) and len(links) == 272 and sum(link[3] for link in links) == 7 and all(entry['evidence'] and not entry['fullAcceptance'] for entry in diagram_audit['checks']), {'associations': len(links), 'limited': 7, 'full_acceptance_claimed': False})
check('independent_two_point_lines', all(len(segment['points']) == 2 for entry in images for segment in entry['associationSegments']), 'Mỗi liên kết một nét riêng; không đường gom')
pagination = load(WORK / 'data/audited-pagination.json')
nav = {}
for paragraph in word.paragraphs:
    hyperlinks = paragraph._p.xpath('./w:hyperlink')
    if hyperlinks:
        anchor = hyperlinks[0].get(qn('w:anchor'))
        values = paragraph._p.xpath('./w:r/w:t')
        if anchor in pagination['pageMap'] and values:
            nav[anchor] = values[-1].text
check('navigation_159_page_references', len(nav) == 159 and nav == pagination['pageMap'] and not pagination['missing'], {'entries': len(nav)})
layout = load(WORK / 'data/layout-validation.json')
check('rendered_layout_bounds', layout['pdfPageCount'] == 413 and layout['pageImageCount'] == 413 and not layout['blankPages'] and not layout['textOutsidePage'], layout)
check('historical_stage4_hash_baseline_not_rebased', digest(BA / '06_Detailed_Use_Cases/data/protected-hashes.json') == digest(WORK / 'data/baselines/protected-hashes-stage4-original.json'), 'Giữ mốc lịch sử, không đổi hash để che sửa tài liệu')
for folder in ['02_Functional_Decomposition_Use_Case_Catalog', '05_Use_Case_Diagrams']:
    result = load(BA / folder / 'Validation_Result.json')
    check('saved_validation_' + folder, result['status'] == 'PASSED', {'status': result['status'], 'scope': 'Nhất quán/truy vết, không chứng nhận toàn bộ ngữ nghĩa'})
markdown = load(BA / '06_Detailed_Use_Cases/Markdown_Parse_Result.json')
check('markdown_162_spec_tables', markdown['success'] and len(markdown['files']) == 162 and all(digest(BA / '06_Detailed_Use_Cases' / item['file']) == item['sha256'] for item in markdown['files']), '162 bảng cú pháp hợp lệ và đúng hash lúc kiểm tra')
tests = load(WORK / 'data/unit-tests-result.json')
flutter = load(WORK / 'data/flutter-summary.json')
check('executed_backend_tests', tests['success'] and tests['numPassedTests'] == 736 and tests['numFailedTests'] == 0 and tests['numPendingTests'] == 7, {'passed': tests['numPassedTests'], 'skipped': tests['numPendingTests'], 'scope': 'Unit/mock, không E2E'})
check('executed_flutter_tests', flutter['success'] and flutter['passed'] == 678 and flutter['failed'] == 0 and flutter['skipped'] == 0, {'passed': flutter['passed'], 'files': len(flutter['files'])})
change_log = load(WORK / 'data/documentation-changes.json')
amendments = change_log['historicalProtectedDifferences'] + change_log['additionalAmendedDocuments']
check('amendment_log_matches_current_documents', all((ROOT / entry['file']).exists() and digest(ROOT / entry['file']) == entry['currentSha256'] for entry in amendments), {'amended_documents': len(amendments)})
visual = {'contact_sheets_reviewed': [f'QA/final-word/contact-{number}.jpg' for number in range(1, 10)], 'pages_reviewed_at_readable_scale': [118, 274, 298], 'scope': 'Đọc 9 ảnh tổng quan phủ 413 trang và 3 trang phóng lớn; không nhận đã đọc thủ công mọi câu của báo cáo'}
result = {'checkedAt': datetime.now(timezone.utc).isoformat(), 'status': 'PASSED' if all(item['passed'] for item in checks) else 'FAILED', 'checks': checks, 'visualReview': visual, 'limits': ['162 UC kiểm tra truy vết và đầu mối quyền; 43 UC đọc sâu các quy tắc ưu tiên, chưa rà hết từng phát biểu thủ công', 'Chưa nghiệm thu DB, nhà cung cấp bên ngoài hoặc mọi luồng đầu-cuối', 'Sơ đồ lớp/tương tác và phần ngoài UC trong báo cáo được giữ, chưa kiểm chứng lại'], 'certifies100Percent': False}
(WORK / 'data/audit-validation.json').write_text(json.dumps(result, ensure_ascii=False, indent=2) + '\n')
print(json.dumps({'status': result['status'], 'checks': len(checks), 'failures': [item for item in checks if not item['passed']]}, ensure_ascii=False, indent=2))
assert result['status'] == 'PASSED'

