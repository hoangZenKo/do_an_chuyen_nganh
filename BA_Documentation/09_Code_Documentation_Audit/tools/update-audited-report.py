from pathlib import Path
from copy import deepcopy
from zipfile import ZipFile
import hashlib
import json
import re
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

ROOT = Path(__file__).resolve().parents[3]
BA = ROOT / 'BA_Documentation'
WORK = BA / '09_Code_Documentation_Audit'
DESKTOP = Path('/Users/hoang/Desktop/ĐỒ ÁN')
SOURCE = DESKTOP / 'Bao_cao_Quan_ly_Nha_tro_V4_2026-10-29_BA.docx'
DESTINATION = DESKTOP / 'Bao_cao_Quan_ly_Nha_tro_V5_2026-10-30_Doi_chieu_Code.docx'
DATA = json.loads((BA / '07_Report_Completion/data/business-report-content.json').read_text())
IMAGES = json.loads((BA / '08_Use_Case_Black_White/data/image-manifest.json').read_text())
LOOKUP = {item['id']: item for item in DATA['useCases']}


def set_cell(cell, value):
    template = deepcopy(cell.paragraphs[0]._p)
    for node in list(template):
        if node.tag != qn('w:pPr'):
            template.remove(node)
    properties = cell.paragraphs[0].runs[0]._r.rPr if cell.paragraphs[0].runs else None
    for paragraph in list(cell._tc):
        if paragraph.tag != qn('w:tcPr'):
            cell._tc.remove(paragraph)
    for line in value.split('\n'):
        paragraph = deepcopy(template)
        run = OxmlElement('w:r')
        if properties is not None:
            run.append(deepcopy(properties))
        text = OxmlElement('w:t')
        text.text = line
        run.append(text)
        paragraph.append(run)
        cell._tc.append(paragraph)


def main():
    original_hash = hashlib.sha256(SOURCE.read_bytes()).hexdigest()
    document = Document(SOURCE)
    changes = []
    found = set()
    for table in document.tables:
        if len(table.rows) != 21 or table.rows[0].cells[0].text != 'Thuộc tính':
            continue
        identifier = table.rows[1].cells[1].text
        record = LOOKUP[identifier]
        found.add(identifier)
        for row in table.rows[1:]:
            field = row.cells[0].text
            value = record['fields'][field]
            if row.cells[1].text != value:
                changes.append({'id': identifier, 'field': field, 'before': row.cells[1].text, 'after': value})
                set_cell(row.cells[1], value)
    assert found == set(LOOKUP), len(found)
    changed_ids = sorted({item['id'] for item in changes})
    assert changed_ids == ['UC-ADM-18.1', 'UC-OPT-07.3', 'UC-SYS-05.1', 'UC-TNT-05', 'UC-TNT-12.1'], changed_ids
    paragraphs = document.paragraphs
    replaced = []
    for position, paragraph in enumerate(paragraphs):
        drawings = paragraph._p.xpath('.//w:drawing')
        if not drawings:
            continue
        following = next((candidate.text for candidate in paragraphs[position + 1:position + 4] if candidate.text.startswith('Hình 2.')), '')
        match = re.match(r'Hình 2\.(\d+)\b', following)
        if not match or not 3 <= int(match.group(1)) <= 32:
            continue
        order = int(match.group(1)) - 2
        entry = IMAGES[order - 1]
        assert entry['order'] == order
        picture = BA / '08_Use_Case_Black_White' / entry['png']
        relation, _ = document.part.get_or_add_image(str(picture))
        assert len(drawings) == 1
        drawing = drawings[0]
        blips = drawing.xpath('.//a:blip')
        assert len(blips) == 1
        blips[0].set(qn('r:embed'), relation)
        with Image.open(picture) as image:
            ratio = image.height / image.width
        width = int(6.25 * 914400)
        height = int(width * ratio)
        if height > int(7.75 * 914400):
            height = int(7.75 * 914400)
            width = int(height / ratio)
        for extent in drawing.xpath('.//wp:extent | .//a:xfrm/a:ext'):
            extent.set('cx', str(width))
            extent.set('cy', str(height))
        replaced.append({'caption': following, 'image': entry['png'], 'sourceId': entry['id']})
    assert len(replaced) == 30, len(replaced)
    pagination = WORK / 'data/audited-pagination.json'
    navigation_updates = 0
    if pagination.exists():
        page_map = json.loads(pagination.read_text())['pageMap']
        for paragraph in document.paragraphs:
            links = paragraph._p.xpath('./w:hyperlink')
            if not links:
                continue
            anchor = links[0].get(qn('w:anchor'))
            if anchor not in page_map:
                continue
            values = paragraph._p.xpath('./w:r/w:t')
            if values:
                values[-1].text = page_map[anchor]
                navigation_updates += 1
    document.save(DESTINATION)
    assert hashlib.sha256(SOURCE.read_bytes()).hexdigest() == original_hash
    with ZipFile(DESTINATION) as archive:
        assert archive.testzip() is None
    reopened = Document(DESTINATION)
    verified = {}
    for table in reopened.tables:
        if len(table.rows) == 21 and table.rows[0].cells[0].text == 'Thuộc tính':
            identifier = table.rows[1].cells[1].text
            values = {row.cells[0].text: row.cells[1].text for row in table.rows[1:]}
            assert values == LOOKUP[identifier]['fields'], identifier
            verified[identifier] = True
    assert len(verified) == 162
    result = {'source': str(SOURCE), 'output': str(DESTINATION), 'sourceSha256': original_hash, 'sourcePreserved': True, 'changedUseCases': changed_ids, 'changedFields': changes, 'replacedImages': replaced, 'verifiedTables': len(verified), 'navigationEntriesUpdated': navigation_updates, 'outputSha256': hashlib.sha256(DESTINATION.read_bytes()).hexdigest()}
    (WORK / 'data/word-corrections.json').write_text(json.dumps(result, ensure_ascii=False, indent=2) + '\n')
    print(json.dumps({'output': str(DESTINATION), 'changedUseCases': changed_ids, 'changedFields': len(changes), 'replacedImages': len(replaced), 'verifiedTables': len(verified)}, ensure_ascii=False, indent=2))

if __name__ == '__main__':
    main()

