import csv
import hashlib
import json
import re
import sys
import unicodedata
from pathlib import Path
from zipfile import ZipFile
from docx import Document
from lxml import etree
from pypdf import PdfReader
from report_content import BASE, ROOT, SNAPSHOT, build_content, SOURCES, ROUTES
from auth_specs import SPECS
from catalog_data import cases, CROSSWALK, TECH

DOCX=ROOT/'docs/Bao_cao_Quan_ly_Nha_tro_UC_AS_IS_Dot_1_2026-10-29.docx'
ORIGINAL=Path('/Users/hoang/Desktop/ĐỒ ÁN/Bao_cao_Quan_ly_Nha_tro_V3_2026-10-28.docx')
build_content()
assert hashlib.sha256(ORIGINAL.read_bytes()).hexdigest()=='198de27e5e6f96734a12125d2c8e3dc1e9c653e82cab2f30dabe8b402cdb3c31'
manifest=json.loads((BASE/'evidence/source-manifest.json').read_text())
for entry in manifest['files']:
    assert hashlib.sha256((SNAPSHOT/entry['file']).read_bytes()).hexdigest()==entry['sha256'],entry['file']
assert len(manifest['files'])==853
assert len(cases())==164 and len(SPECS)==9
old=json.loads((ROOT/'docs/report_2026-10-28_v3/audit.json').read_text())['cases']
new_ids={c['id'] for c in cases()}
assert all(CROSSWALK.get(c['id']) or c['id'] in new_ids for c in old)
bound={b for c in cases() for b in c['bindings'] if not b.startswith('LOCAL.')}
assert len(bound)==393 and set(ROUTES)==bound|set(TECH)
assert all((SNAPSHOT/e['file']).exists() and 0<e['line']<=len((SNAPSHOT/e['file']).read_text().splitlines()) for e in SOURCES.values())
document=Document(DOCX)
headings=[p.text for p in document.paragraphs]
letters=list('ABCDEFGHIJ')
for letter in letters:
    assert sum(p.startswith(letter+' ') for p in headings)==9,(letter,'Missing UC section')
for spec in SPECS:
    assert spec['flow'] and spec['fields'] and spec['errors'] and spec['scenarios']
    assert len(spec['scenarios'])>=4
    assert set(spec['sources'])<=set(SOURCES)
rules=[num for s in SPECS for num,_,_ in s['rules']]
assert len(rules)==len(set(rules))==26
with ZipFile(DOCX) as archive:
    assert archive.testzip() is None
    xml=etree.fromstring(archive.read('word/styles.xml'))
    namespace={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    assert not xml.xpath('//w:pBdr',namespaces=namespace)
assert len(list(csv.reader((BASE/'Requirements_Traceability_Matrix.tsv').open(),delimiter='\t')))==406
backend=(BASE/'evidence/auth-backend-tests.log').read_text()
flutter=(BASE/'evidence/auth-flutter-tests.log').read_text()
assert 'Tests:       31 passed' in backend and '+48: All tests passed!' in flutter
schema=(SNAPSHOT/'backend/prisma/schema.prisma').read_text()
assert len(re.findall(r'^model ',schema,re.M))==110
assert len(re.findall(r'^enum ',schema,re.M))==30
result=dict(original_unchanged=True,snapshot_hashes_verified=853,uc_catalog=164,
            completed_ucs=9,pending_ucs=155,source_locators=len(SOURCES),routes_covered=405,
            old_uc_ids_accounted=131,business_rules=26,acceptance_scenarios=sum(len(s['scenarios']) for s in SPECS),
            backend_tests_passed=31,flutter_tests_passed=48,sections_A_to_J=True,
            e2e_executed=False)
if len(sys.argv)>1:
    render=Path(sys.argv[1]);reader=PdfReader(next(render.glob('*.pdf')))
    normalize=lambda text:re.sub(r'\s+',' ',unicodedata.normalize('NFC',text)).strip()
    texts=[normalize(p.extract_text()) for p in reader.pages]
    start=next(i for i,text in enumerate(texts) if i>2 and '4 Đặc tả chi tiết nhóm tài khoản' in text)
    measured={}
    for paragraph in document.paragraphs:
        if paragraph.style.name=='Heading 1' and paragraph.text!='Mục lục' or paragraph.style.name=='Heading 2' and paragraph.text.startswith('UC-SYS-'):
            first=start if paragraph.text.startswith('UC-SYS-') else 2
            hits=[i+1 for i in range(first,len(texts)) if normalize(paragraph.text) in texts[i]]
            assert hits,paragraph.text
            measured[paragraph.text]=hits[0]
    (BASE/'evidence/page-map-measured.json').write_text(json.dumps(measured,ensure_ascii=False,indent=2))
    current=json.loads((BASE/'evidence/page-map.json').read_text())
    result.update(rendered_pages=len(texts),toc_matches_rendered=current==measured,
                  sparse_pages=[i+1 for i,text in enumerate(texts) if len(text)<200])
    (BASE/'evidence/layout-validation.json').write_text(json.dumps(result,ensure_ascii=False,indent=2))
    print(json.dumps(result,ensure_ascii=False,indent=2))
    if current!=measured:sys.exit(2)
else:print(json.dumps(result,ensure_ascii=False,indent=2))
