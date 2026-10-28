import csv
import hashlib
import json
import re
from collections import Counter
from pathlib import Path
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from report_content import BASE, ROOT, SNAPSHOT, SOURCES, ROUTES, CLASSES, BLOCKS, build_content
from catalog_data import cases, GROUPS, TECH, CROSSWALK
from auth_specs import SPECS

ORIGINAL = Path('/Users/hoang/Desktop/ĐỒ ÁN/Bao_cao_Quan_ly_Nha_tro_V3_2026-10-28.docx')
OUTPUT = ROOT / 'docs/Bao_cao_Quan_ly_Nha_tro_UC_AS_IS_Dot_1_2026-10-29.docx'
ORIGINAL_HASH = '198de27e5e6f96734a12125d2c8e3dc1e9c653e82cab2f30dabe8b402cdb3c31'
assert hashlib.sha256(ORIGINAL.read_bytes()).hexdigest() == ORIGINAL_HASH
build_content()

def save_support_files():
    text=['# Đặc tả Use Case AS IS đợt 1', '', 'Ngày 29/10/2026. Bản thay thế phần UC của báo cáo cũ; không phải xác nhận hoàn tất 164 UC.', '']
    for block in BLOCKS:
        if block[0]=='heading':text += ['#'*min(6,block[1]+1)+' '+block[2], '']
        elif block[0]=='paragraph':text += [block[1], '']
        elif block[0]=='table':
            escape=lambda value:str(value).replace('|','\\|').replace('\n','<br>')
            text += ['| '+' | '.join(map(escape,block[1]))+' |', '| '+' | '.join('---' for _ in block[1])+' |']
            text += ['| '+' | '.join(map(escape,row))+' |' for row in block[2]]
            text += ['']
    (BASE/'BA_AS_IS_Dot_1.md').write_text('\n'.join(text))
    owners={key:[] for key in ROUTES}
    for case in cases():
        for binding in case['bindings']:
            if binding in owners:owners[binding].append(case['id'])
    assert all(owners[key] or key in TECH for key in ROUTES)
    assert len(cases())==164 and len(SPECS)==9
    with (BASE/'Requirements_Traceability_Matrix.tsv').open('w') as stream:
        writer=csv.writer(stream,delimiter='\t')
        writer.writerow(['Mã nguồn','Use Case hoặc SP','HTTP','Đường dẫn','Controller.method','File:dòng','Roles metadata','AdminPermissions metadata','RequirePermissions metadata','DTO params','Service calls','Mức xác nhận'])
        for key,route in ROUTES.items():
            calls=[]
            for call in route['calls']:
                cls=CLASSES.get(call['type'])
                method=next((m for m in cls['methods'] if m['name']==call['method']),None) if cls else None
                calls.append((call['type'] or call['field'])+'.'+call['method']+((' '+cls['file']+':'+str(method['line'])) if method else ' (chưa phân giải tĩnh)'))
            writer.writerow([route['ref'], '; '.join(owners[key]) or TECH[key],route['verb'],route['path'],key,route['file']+':'+str(route['line']), route['roles'],route['adminPermissions'],route['staffPermissions'],json.dumps(route['parameters'],ensure_ascii=False),'; '.join(calls),'A–J' if any(c.startswith('UC-SYS-') for c in owners[key]) else 'Kiểm kê wiring; cần xác nhận hành vi từng UC'])
    old=json.loads((ROOT/'docs/report_2026-10-28_v3/audit.json').read_text())['cases']
    ids={case['id'] for case in cases()}
    cross={entry['id']:CROSSWALK.get(entry['id'],[entry['id']] if entry['id'] in ids else []) for entry in old}
    assert all(cross.values()) and len(cross)==131
    report=dict(date='2026-10-29', original=str(ORIGINAL),original_sha256=ORIGINAL_HASH,
        snapshot_files=853,catalog_ucs=len(cases()),detailed_ucs=[s['code'] for s in SPECS],
        pending_ucs=155,routes=len(ROUTES),controllers=len({r['class'] for r in ROUTES.values()}),
        routes_mapped_to_uc=sum(bool(v) for v in owners.values()),technical_routes=len(TECH),
        groups=dict(Counter(c['group'] for c in cases())),crosswalk=cross,evidence=list(SOURCES.values()),
        tests=dict(backend_suites_passed=7,backend_tests_passed=31,flutter_tests_passed=48,e2e_executed=False))
    (BASE/'evidence/revision-validation.json').write_text(json.dumps(report,ensure_ascii=False,indent=2))
    catalog=[]
    for case in cases():
        case['status']='xác nhận' if case['group']=='SYS' else ('chỉ triển khai một phần' if case['code'] in ('TEN-05','TEN-09.2') else 'chưa đủ bằng chứng đối chiếu chi tiết')
        catalog.append(case)
    (BASE/'Use_Case_Catalog.json').write_text(json.dumps(catalog,ensure_ascii=False,indent=2))
    snippets=[]
    for source in SOURCES.values():
        lines=(SNAPSHOT/source['file']).read_text().splitlines()
        assert source['line']<=len(lines)
        snippets += ['## '+source['id']+' '+source['label'], source['file']+':'+str(source['line']), '```']
        snippets += [str(number+1)+': '+lines[number] for number in range(max(0,source['line']-1),min(len(lines),source['line']+8))]
        snippets += ['```','']
    (BASE/'evidence/source-excerpts.md').write_text('\n'.join(snippets))

def set_font(run,size=12,bold=False):
    run.font.name='Times New Roman';run.font.size=Pt(size);run.font.bold=bold;run.font.color.rgb=RGBColor(0,0,0)
    fonts=run._r.get_or_add_rPr().find(qn('w:rFonts'))
    if fonts is None:fonts=OxmlElement('w:rFonts');run._r.get_or_add_rPr().append(fonts)
    for key in ('ascii','hAnsi','cs','eastAsia'):fonts.set(qn('w:'+key),'Times New Roman')

def compact_paragraph(paragraph,table=False):
    paragraph.paragraph_format.space_after=Pt(3 if table else 6)
    paragraph.paragraph_format.space_before=Pt(0)
    paragraph.paragraph_format.line_spacing=1.08 if table else 1.16
    paragraph.paragraph_format.first_line_indent=Pt(0)
    paragraph.paragraph_format.widow_control=True

def add_table(doc,headers,rows,widths):
    widths=widths or [15.9/len(headers)]*len(headers)
    if headers[-1]=='TT':widths=[3.7,4.9,6.2,1.1]
    if headers[0]=='Nguồn' and len(headers)==3:widths=[1.7,8.6,5.6]
    tbl=doc.add_table(rows=1, cols=len(headers));tbl.autofit=False
    for grid,width in zip(tbl.columns,widths):grid.width=Cm(width)
    pr=tbl._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        item=OxmlElement('w:'+edge);item.set(qn('w:val'),'single');item.set(qn('w:sz'),'4');item.set(qn('w:color'),'D9D9D9');borders.append(item)
    pr.append(borders)
    margins=OxmlElement('w:tblCellMar')
    for edge,twips in [('top','65'),('bottom','65'),('left','90'),('right','90')]:
        item=OxmlElement('w:'+edge);item.set(qn('w:w'),twips);item.set(qn('w:type'),'dxa');margins.append(item)
    pr.append(margins)
    for index,values in enumerate([headers]+list(rows)):
        row=tbl.rows[0] if index==0 else tbl.add_row()
        trpr=row._tr.get_or_add_trPr()
        prevent=OxmlElement('w:cantSplit');trpr.append(prevent)
        if index==0:trpr.append(OxmlElement('w:tblHeader'))
        for cell,value,width in zip(row.cells,values,widths):
            cell.width=Cm(width);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            shading=OxmlElement('w:shd');shading.set(qn('w:fill'),'FFFFFF');cell._tc.get_or_add_tcPr().append(shading)
            paragraph=cell.paragraphs[0];compact_paragraph(paragraph,True)
            if index==0:paragraph.paragraph_format.keep_with_next=True
            if headers[0]=='Tiêu chí':paragraph.paragraph_format.keep_with_next=index<len(rows)
            paragraph.alignment=WD_ALIGN_PARAGRAPH.LEFT
            set_font(paragraph.add_run(str(value)),11 if len(headers)>=4 else 11.5,index==0)

def make_document():
    document=Document()
    section=document.sections[0]
    section.page_width=Cm(21);section.page_height=Cm(29.7)
    section.left_margin=Cm(3);section.right_margin=Cm(2)
    section.top_margin=Cm(2.5);section.bottom_margin=Cm(2.5)
    section.header_distance=Cm(1.1);section.footer_distance=Cm(1.1)
    for style in document.styles:
        if style.type==1:
            style.font.name='Times New Roman';style.font.color.rgb=RGBColor(0,0,0)
            fonts=style.element.get_or_add_rPr().find(qn('w:rFonts'))
            if fonts is None:fonts=OxmlElement('w:rFonts');style.element.get_or_add_rPr().append(fonts)
            for attribute in list(fonts.attrib):
                if attribute.endswith('Theme'):del fonts.attrib[attribute]
            for key in ('ascii','hAnsi','eastAsia','cs'):fonts.set(qn('w:'+key),'Times New Roman')
        for border in list(style.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
    normal=document.styles['Normal'];normal.font.size=Pt(12)
    normal.paragraph_format.space_after=Pt(6);normal.paragraph_format.line_spacing=1.16
    for level in (1,2,3):
        style=document.styles['Heading '+str(level)];style.font.size=Pt({1:16,2:14,3:12}[level]);style.font.bold=True
        style.paragraph_format.space_before=Pt(12 if level<3 else 9);style.paragraph_format.space_after=Pt(6)
        style.paragraph_format.keep_with_next=True
    title=document.styles['Title'];title.font.size=Pt(22);title.font.bold=True
    footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    set_font(footer.add_run('AS IS  Đợt 1  |  '),10)
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    document.core_properties.title='Đặc tả Use Case AS IS hệ thống quản lý nhà trọ đợt 1'
    document.core_properties.subject='Sửa phần Use Case báo cáo V3 theo snapshot source ngày 29/10/2026'
    document.core_properties.author='Phân tích ngược từ mã nguồn dự án'
    document.add_paragraph('ĐẶC TẢ USE CASE AS IS\nHỆ THỐNG QUẢN LÝ NHÀ TRỌ',style='Title')
    paragraph=document.add_paragraph('Bản sửa đổi phần Use Case của báo cáo V3\nĐợt 1  |  Ngày 29 tháng 10 năm 2026')
    compact_paragraph(paragraph)
    document.add_paragraph('Mục đích của bản sửa',style='Heading 2')
    document.add_paragraph('Thay mô tả Use Case tổng quát bằng nội dung truy vết được về code hiện có. Bản này là tài liệu UC độc lập để thay phần 2.4.1 của báo cáo cũ; không sửa các chương hay hình ảnh ngoài phạm vi UC và không ghi đè tệp gốc.')
    document.add_paragraph('Phần đã hoàn thành',style='Heading 2')
    document.add_paragraph('Báo cáo hệ thống và ma trận quyền; catalog 164 mục tiêu nghiệp vụ; 9 Use Case tài khoản đặc tả đầy đủ A–J; 2 tiến trình hỗ trợ phiên; quy tắc, truy vết và danh sách thiếu bằng chứng.')
    document.add_paragraph('Giới hạn cần đọc trước khi sử dụng',style='Heading 2')
    document.add_paragraph('155 Use Case còn lại mới có catalog và mapping API, chưa phải đặc tả A–J đã nghiệm thu. Không dùng tài liệu này để kết luận toàn bộ 164 UC đã hoàn tất. Kết quả kiểm thử đợt này là 31 backend và 48 Flutter tests; chưa chạy E2E với DB thật.')
    document.add_paragraph('Nguồn chuẩn là snapshot mã nguồn ngày 29/10/2026. Mọi số dòng và nhận định AS-IS phải được hiểu theo snapshot này; báo cáo V3 chỉ là tài liệu để đối chiếu và sửa.')
    document.add_page_break()
    document.add_paragraph('Mục lục',style='Heading 1')
    toc=[b for b in BLOCKS if b[0]=='heading' and (b[1]==1 or (b[1]==2 and b[2].startswith('UC-SYS-')))]
    pages_file=BASE/'evidence/page-map.json'
    pages=json.loads(pages_file.read_text()) if pages_file.exists() else {}
    for entry in toc:
        paragraph=document.add_paragraph();compact_paragraph(paragraph)
        paragraph.paragraph_format.left_indent=Cm(0.4 if entry[1]==2 else 0)
        paragraph.paragraph_format.tab_stops.add_tab_stop(Cm(15.6),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
        set_font(paragraph.add_run(entry[2]+'\t'+str(pages.get(entry[2],'—'))),11.5)
    document.add_page_break()
    for block in BLOCKS:
        if block[0]=='heading':
            paragraph=document.add_paragraph(block[2],style='Heading '+str(block[1]))
            if block[1]==1 and not block[2].startswith('1 '):paragraph.paragraph_format.page_break_before=True
        elif block[0]=='paragraph':
            paragraph=document.add_paragraph(block[1]);compact_paragraph(paragraph)
        elif block[0]=='table':add_table(document,*block[1:])
        elif block[0]=='pagebreak':pass
    document.save(OUTPUT)
    return document

save_support_files()
document=make_document()
(BASE/'artifact.md').write_text('''# Hồ sơ chỉnh sửa tài liệu

- Tệp nguồn là báo cáo V3 của người dùng; SHA-256 giữ nguyên: '''+ORIGINAL_HASH+'''.
- Đầu ra là bản UC độc lập đợt 1 thay phần kịch bản UC; không phải bản thay toàn bộ báo cáo hoặc bản đặc tả 164 UC hoàn tất.
- Kế thừa ngôn ngữ học thuật của mẫu: A4, Times New Roman, đen trắng, lề trái 3 cm/phải 2 cm, bảng có lưới. Không sao chép logo, tên tác giả hay số trang cũ.
- Nội dung chỉnh sửa: báo cáo AS-IS, tác nhân/quyền, phân rã, catalog, 9 UC A–J, 2 SP, BR, truy vết và điểm chưa rõ. Các mục còn thiếu được ghi rõ.
- Vùng biên tập tái tạo: tools/auth_specs.py và tools/report_content.py; build_revision.py tạo DOCX và Markdown cùng nguồn dữ liệu.
- Ma trận API đầy đủ ở Requirements_Traceability_Matrix.tsv; đoạn nguồn định vị ở evidence/source-excerpts.md.
- Nguồn và test snapshot không bao gồm secrets hoặc dữ liệu DB; không sửa source ứng dụng.
''')
print(json.dumps(dict(output=str(OUTPUT),paragraphs=len(document.paragraphs),tables=len(document.tables),sources=len(SOURCES)),ensure_ascii=False))
