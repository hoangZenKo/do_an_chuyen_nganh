import json
import re
from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION, WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from build import MODEL, BASE, ROOT, Figure, message_implementation, reference_rows, scenario_implementation_rows


def configure(section, kind='portrait'):
    if kind in ['figure', 'figure_landscape']:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width, section.page_height = Mm(297), Mm(420)
        if kind == 'figure_landscape':
            section.orientation = WD_ORIENT.LANDSCAPE
            section.page_width, section.page_height = Mm(420), Mm(297)
        margin = 14
    elif kind == 'landscape':
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = Mm(297), Mm(210)
        margin = 14
    else:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width, section.page_height = Mm(210), Mm(297)
        margin = 18
    section.top_margin = section.bottom_margin = Mm(10 if kind == 'landscape' else margin)
    section.left_margin = section.right_margin = Mm(margin)
    section.header_distance = section.footer_distance = Mm(6)


def font_run(run, size=None):
    run.font.name = 'Times New Roman'
    run.font.color.rgb = RGBColor(0, 0, 0)
    if size:
        run.font.size = Pt(size)
    properties = run._element.get_or_add_rPr()
    fonts = properties.rFonts
    if fonts is None:
        fonts = OxmlElement('w:rFonts')
        properties.insert(0, fonts)
    for attribute in list(fonts.attrib):
        if 'Theme' in attribute:
            del fonts.attrib[attribute]
    for attribute in ['ascii', 'hAnsi', 'eastAsia', 'cs']:
        fonts.set(qn('w:'+attribute), 'Times New Roman')


def paragraph(document, text, style=None):
    result = document.add_paragraph(text, style)
    for run in result.runs:
        font_run(run)
    return result


def table(document, headings, rows, widths, size=11):
    result = document.add_table(rows=1, cols=len(headings))
    result.style = 'Table Grid'
    result.autofit = False
    for column, width in zip(result.columns, widths):
        column.width = Mm(width)
    header = result.rows[0]
    repeat = OxmlElement('w:tblHeader')
    header._tr.get_or_add_trPr().append(repeat)
    for cell, text, width in zip(header.cells, headings, widths):
        cell.width = Mm(width)
        cell.text = text
        shading = OxmlElement('w:shd')
        shading.set(qn('w:fill'), 'FFFFB9')
        cell._tc.get_or_add_tcPr().append(shading)
    for row in rows:
        cells = result.add_row().cells
        for cell, text, width in zip(cells, row, widths):
            cell.width = Mm(width)
            cell.text = str(text)
    for row_index, row in enumerate(result.rows):
        properties = row._tr.get_or_add_trPr()
        properties.append(OxmlElement('w:cantSplit'))
        for cell in row.cells:
            for text in cell.paragraphs:
                text.paragraph_format.space_before = Pt(0.5)
                text.paragraph_format.space_after = Pt(0.5)
                text.paragraph_format.line_spacing = 1.0
                for run in text.runs:
                    font_run(run, size)
                    if row_index == 0:
                        run.bold = True
    spacer = paragraph(document, '')
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.space_after = Pt(0)
    spacer.paragraph_format.line_spacing = Pt(1)
    spacer.paragraph_format.keep_with_next = True
    return result


def section(document, kind='portrait'):
    created = document.add_section(WD_SECTION.NEW_PAGE)
    configure(created, kind)
    return created


def create():
    document = Document()
    configure(document.sections[0])
    for style in document.styles:
        if style.type == 1 or style.type == 2:
            style.font.name = 'Times New Roman'
            style.font.color.rgb = RGBColor(0, 0, 0)
            properties = style.element.get_or_add_rPr()
            fonts = properties.rFonts
            if fonts is None:
                fonts = OxmlElement('w:rFonts')
                properties.insert(0, fonts)
            for attribute in list(fonts.attrib):
                if 'Theme' in attribute:
                    del fonts.attrib[attribute]
            for attribute in ['ascii','hAnsi','eastAsia','cs']:
                fonts.set(qn('w:'+attribute), 'Times New Roman')
            for border in list(style.element.findall('.//'+qn('w:pBdr'))):
                border.getparent().remove(border)
    document.styles['Normal'].font.size = Pt(12)
    document.styles['Normal'].paragraph_format.space_after = Pt(7)
    document.styles['Normal'].paragraph_format.line_spacing = 1.12
    document.styles['Title'].font.size = Pt(24)
    document.styles['Heading 1'].font.size = Pt(16)
    document.styles['Heading 2'].font.size = Pt(14)
    document.styles['Heading 3'].font.size = Pt(12)
    for name in ['Heading 1','Heading 2','Heading 3']:
        document.styles[name].paragraph_format.keep_with_next = True
    document.core_properties.title = 'Các biểu đồ tuần tự của hệ thống quản lý nhà trọ'
    document.core_properties.subject = 'Đặc tả tương tác cho 10 kịch bản và 15 Use Case'
    document.core_properties.author = ''
    document.core_properties.keywords = 'UML sequence alt opt loop rental room'
    footer = document.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    page_field = OxmlElement('w:fldSimple')
    page_field.set(qn('w:instr'), 'PAGE')
    footer._p.append(page_field)
    paragraph(document, 'Các biểu đồ tuần tự của hệ thống quản lý nhà trọ', 'Title')
    paragraph(document, '10 kịch bản nghiệp vụ trọng tâm 22 hình tương tác chi tiết')
    paragraph(document, 'Phiên bản 1.2 chuẩn hóa tên thành phần và văn phong ngày 7 tháng 10 năm 2026')
    paragraph(document, '1 Mục đích và phạm vi', 'Heading 1')
    paragraph(document, 'Tài liệu phục vụ báo cáo Kỹ thuật phần mềm, mô tả thứ tự tương tác giữa Admin, Chủ trọ, Người thuê, giao diện, các bộ phận xử lý nghiệp vụ và kho dữ liệu. Mười kịch bản liên quan 15 Use Case được phân rã thành 22 hình để thể hiện luồng chính, nhánh thay thế, ngoại lệ và thao tác tùy chọn.')
    paragraph(document, 'Bộ hình tập trung vào những luồng đã chọn cho báo cáo, không tương đương 67 sơ đồ cho toàn bộ 67 Use Case. Mỗi hình đi kèm bảng thông điệp; các điểm cần xác nhận được giữ riêng để tránh coi hành vi hiện tại là chính sách đã chốt.')
    paragraph(document, '2 Quy ước và cách đọc', 'Heading 1')
    for text in [
        'Giữ form vàng nhạt FFFFB9, viền đỏ nâu 800000, chữ Tahoma, nền trắng. Không có khung bao ngoài hoặc tab sd. Đây là hình mô phỏng phong cách StarUML cổ điển, không phải dự án mdj.',
        'Actor dùng hình người, boundary cho giao diện và Controller, control cho Service, cylinder cho kho dữ liệu. Mũi tên liền là lời gọi; nét đứt là phản hồi; thanh vàng mảnh là thực thi ở mức khái quát.',
        'alt chỉ chọn một nhánh; opt chạy khi điều kiện đúng; loop lặp theo tập dữ liệu; break kết thúc tương tác bao ngoài; ref dẫn tới tương tác được phân rã. Không đọc các nhánh alt như những bước chạy nối tiếp.',
        'Lifeline dùng tên thành phần hoặc đối tượng: Controller/Service giữ tên lớp thực tế; giao diện và kho dữ liệu dùng danh từ tiếng Việt. Thông điệp ưu tiên ngôn ngữ nghiệp vụ; băm mật khẩu được gọi đúng thuật ngữ. Tên hàm, thuật toán, đường dẫn API, mã lỗi, biến và trạng thái mã nguồn nằm trong bảng chi tiết triển khai khi cần. Thông điệp đánh số riêng trong từng hình.',
        'Các trang hình dùng A3 dọc hoặc ngang theo độ dài tương tác để giữ chữ nội dung khoảng 10 điểm trở lên; bảng thông điệp dùng A4 ngang. Nên dùng SVG khi đưa hình sang báo cáo để giữ nét chữ.'
    ]:
        paragraph(document, text)
    section(document)
    paragraph(document, '3 Điều kiện và ngoại lệ dùng chung', 'Heading 1')
    paragraph(document, 'Hệ thống kiểm tra phiên đăng nhập, tình trạng sử dụng tài khoản, vai trò, quyền thao tác và dữ liệu đầu vào trước khi xử lý nghiệp vụ. Biểu mẫu sai chưa gửi yêu cầu. Khi không đủ điều kiện, tương tác không đi tiếp nhánh thành công.')
    paragraph(document, 'Lỗi trước khi lưu thành công hủy các thay đổi trong lần ghi nhận. Lỗi sau khi lưu hoặc mất phản hồi không chứng minh dữ liệu đã bị hủy; cần tra cứu trước khi gửi lại. Ghi nhật ký không làm hủy kết quả nghiệp vụ đã lưu thành công.')
    table(document, ['Nội dung nghiệp vụ', 'Chi tiết triển khai khi cần'], [
        ['Kiểm tra quyền và dữ liệu', 'Guards kiểm tra access token, tài khoản hoạt động, vai trò và quyền; lỗi xác thực 401, thiếu quyền 403, DTO không hợp lệ 400 trước service.'],
        ['Ghi dữ liệu thống nhất', 'BEGIN COMMIT ROLLBACK biểu diễn giao dịch Prisma. Lỗi trước commit hoàn tác phần ghi trong giao dịch; lỗi sau commit không chứng minh đã rollback.'],
        ['Ghi nhật ký', 'AuditService.recordSafe bắt lỗi audit sau nghiệp vụ, không hoàn tác giao dịch nghiệp vụ đã commit.']
    ], [48,126], 10.5)
    paragraph(document, '4 Danh mục sơ đồ', 'Heading 1')
    table(document, ['Mã', 'Kịch bản', 'Hình'], [[scenario['id'], scenario['title'], ', '.join(view['id'] for view in scenario['views'])] for scenario in MODEL['scenarios']], [20, 106, 48], 11)
    paragraph(document, 'Hình mẫu tạo hợp đồng SD-01 trong bộ demo cũ chỉ là lịch sử duyệt form. Trong bộ chính thức, SD-01 là đăng nhập và SD-04 là tạo hợp đồng.')
    for scenario in MODEL['scenarios']:
        section(document)
        paragraph(document, scenario['id']+' '+scenario['title'], 'Heading 1')
        paragraph(document, 'Use Case '+', '.join(scenario['ucs']))
        for heading, text in [('Điều kiện trước',scenario['pre']), ('Dữ liệu đầu vào',scenario['input']), ('Điều kiện sau',scenario['post'])]:
            paragraph(document, heading, 'Heading 2')
            paragraph(document, text)
        paragraph(document, 'Nhánh nghiệp vụ và giới hạn', 'Heading 2')
        for note in scenario['notes']:
            paragraph(document, note, 'List Bullet')
        paragraph(document, 'Liên hệ triển khai và căn cứ kiểm chứng', 'Heading 2')
        table(document, ['Nội dung', 'Chi tiết triển khai khi cần'], scenario_implementation_rows(scenario), [48,126], 10.5)
        for view in scenario['views']:
            figure = Figure(view)
            kind = 'figure_landscape' if figure.height <= 1400 else 'figure'
            current = section(document, kind)
            paragraph(document, view['id']+' '+view['title'], 'Heading 1')
            available_width = current.page_width-current.left_margin-current.right_margin
            available_height = current.page_height-current.top_margin-current.bottom_margin-Mm(25)
            scale = min(available_width/figure.width, available_height/figure.height)
            minimum_font = 24*scale/12700
            if minimum_font < 10:
                raise ValueError(view['id']+' diagram text too small '+str(minimum_font))
            picture_paragraph = document.add_paragraph()
            picture_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            picture_paragraph.paragraph_format.space_after = Pt(0)
            picture = picture_paragraph.add_run().add_picture(str(BASE/(view['file']+'.png')), width=int(scale*figure.width))
            picture._inline.docPr.set('descr', view['id']+' '+view['title']+' với nhánh UML alt opt loop hoặc break theo nghiệp vụ')
            picture._inline.docPr.set('title', view['id'])
            section(document, 'landscape')
            paragraph(document, view['id']+' Diễn giải thông điệp', 'Heading 2')
            rows = []
            for message in figure.messages:
                sender = view['participants'][message['from']]['name']
                receiver = view['participants'][message['to']]['name']
                branch = '; '.join(re.sub(r'^(?:alt|opt|loop|break) ', '', guard) for guard in message['guards'])
                business = message.get('business_detail', message['label'])
                if branch:
                    business += '\nĐiều kiện: ' + branch + '.'
                implementation = message_implementation(message) or 'Không cần bổ sung'
                rows.append([message['number'],sender+' → '+receiver,business,implementation])
            table(document,['Bước','Tương tác','Diễn giải nghiệp vụ','Chi tiết triển khai khi cần'],rows,[14,55,100,100],10)
    section(document)
    paragraph(document, '5 Tài nguyên chỉnh sửa và truy vết', 'Heading 1')
    for text in [
        'SEQUENCE_DIAGRAM_SPECIFICATION.md chứa cùng 22 hình, bảng thông điệp, căn cứ nguồn và đối chiếu luồng chính thay thế ngoại lệ từ đặc tả Use Case.',
        'sequence-diagrams/full/SEQUENCE_DIAGRAMS.drawio gồm 22 trang chỉnh sửa được; mỗi hình cũng có tệp drawio riêng, SVG vector và PNG nền trắng. Hình học và chữ trong drawio là đối tượng riêng, không phải ảnh nhúng.',
        'sequence-diagrams/full/model.json tách nhãn nghiệp vụ với thông tin triển khai. build.py tái tạo hình và Markdown; build_document.py tạo Word. manifest.json ghi kích thước, thông điệp, các khung nhánh và SHA256 của tài nguyên.',
        'Mục 2.4.4 của báo cáo Markdown được cập nhật bằng cùng bộ hình. Báo cáo Word và PDF tổng, tài liệu Use Case và sơ đồ lớp hiện có không bị ghi đè.',
        'Các điểm PX vẫn cần xác nhận, đặc biệt thu cọc ban đầu, đồng thời xét duyệt, chính sách bảo trì và ngữ nghĩa báo cáo. Tài liệu sơ đồ không đồng nghĩa các giới hạn này đã được sửa trong ứng dụng.'
    ]:
        paragraph(document, text)
    output = ROOT / 'docs/ba/SEQUENCE_DIAGRAM_SPECIFICATION.docx'
    document.save(output)
    print('Saved '+str(output))


if __name__ == '__main__':
    create()
