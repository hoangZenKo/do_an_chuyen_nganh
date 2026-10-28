import copy
import json
import re
import subprocess
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Emu, Mm
from PIL import Image


BASE = Path(__file__).resolve().parent
ROOT = BASE.parents[2]
MODEL = json.loads((BASE / 'class-model.json').read_text())
LAYOUTS = json.loads((BASE / 'diagram-manifest.json').read_text())


REPLACEMENTS = {
    'Hộp tiêu đề màu xám': '<<reference>>',
    'Sơ đồ lớp tổng quan của hệ thống quản lý nhà trọ': 'Sơ đồ lớp miền nghiệp vụ tổng quan (Domain Class Diagram)',
    'Sơ đồ lớp tổng quan': 'Sơ đồ lớp miền nghiệp vụ tổng quan (Domain Class Diagram)',
    'Tổng quan các lớp nghiệp vụ': 'Sơ đồ lớp miền nghiệp vụ tổng quan (Domain Class Diagram)',
    'Sơ đồ tổng quan gồm 13 lớp trọng tâm và 14 quan hệ được chọn.': 'Sơ đồ lớp miền nghiệp vụ tổng quan (Domain Class Diagram). Sơ đồ gồm 13 lớp trọng tâm và 14 quan hệ được chọn.',
    'Đường nét đứt có mũi tên': 'Chú thích <<IDRef>>',
    'Đường nét đứt': 'Chú thích <<IDRef>>',
    'đường nét đứt': 'chú thích <<IDRef>>',
    'Ba tham chiếu nét đứt': 'Ba tham chiếu bằng mã',
    'Ba đường tham chiếu theo mã': 'Ba trường tham chiếu theo mã',
    'ba đường nét đứt': 'chú thích <<IDRef>> cho ba trường mã',
    'ba tham chiếu theo mã của lượt gửi thông báo được thể hiện riêng bằng đường nét đứt': 'ba tham chiếu theo mã của lượt gửi thông báo được thể hiện riêng bằng chú thích <<IDRef>>, không có đường nối',
    'Ba tham chiếu nét đứt của NotificationDelivery': 'Ba trường IDRef của NotificationDelivery',
    'các mã landlordId, tenantId, invoiceId được thể hiện bằng tham chiếu nét đứt': 'các mã landlordId, tenantId, invoiceId được thể hiện bằng chú thích <<IDRef>>, không vẽ đường nối',
    'Giữ đường nét đứt và không gán bội số.': 'Giữ chú thích <<IDRef>>, không vẽ đường nối và không gán bội số.',
    'Sơ đồ sử dụng ba đường nét đứt để thể hiện ngữ cảnh liên quan; không gán bội số hoặc diễn giải như khóa ngoại đã được bảo đảm.': 'Sơ đồ ghi các trường mã và chú thích <<IDRef>>, không vẽ đường nối tới các lớp ngữ cảnh. Tham chiếu bằng mã định danh trong dữ liệu; không phải structural association UML.',
    'được vẽ bằng đường nét đứt.': 'được mô tả bằng chú thích <<IDRef>>, không vẽ đường nối UML.',
}


def replace_text(text):
    for before, after in sorted(REPLACEMENTS.items(), key=lambda pair: -len(pair[0])):
        text = text.replace(before, after)
    return text


def apply_text(filename, content):
    destination = ROOT / filename
    old = destination.read_text() if destination.exists() else None
    if old == content:
        return
    operation = ('*** Update File: ' if old is not None else '*** Add File: ') + filename + '\n'
    patch = '*** Begin Patch\n' + operation
    if old is not None:
        patch += '@@\n' + ''.join('-' + line + '\n' for line in old.splitlines())
    patch += ''.join('+' + line + '\n' for line in content.splitlines()) + '*** End Patch\n'
    subprocess.run(['apply_patch'], input=patch, text=True, cwd=ROOT, check=True)


def fields_text(name, typed=True):
    fields = MODEL['models'][name]['selected_fields']
    return '; '.join(field + ': ' + MODEL['models'][name]['fields'][field]['type'] for field in fields) if typed else ', '.join(field for field in fields if field != 'id')


def update_markdown():
    filename = 'docs/ba/CLASS_DIAGRAM_SPECIFICATION.md'
    text = replace_text((ROOT / filename).read_text())
    text = text.replace('| <<reference>> |', '| `<<reference>>` |')
    for name in MODEL['included']:
        pattern = r'(^\| `' + name + r'`<br>[^\n]*?\| [^\n]*?\| )[^\n]*( \|$)'
        text = re.sub(pattern, lambda match: match.group(1) + '; '.join('`' + field + ': ' + MODEL['models'][name]['fields'][field]['type'] + '`' for field in MODEL['models'][name]['selected_fields']) + match.group(2), text, flags=re.M)
    text = text.replace('| Chú thích <<IDRef>> | Tham chiếu theo mã ở lớp nguồn đến loại đối tượng liên quan. Không phải quan hệ khóa ngoại đã khai báo và không gán bội số. |', '| Chú thích `<<IDRef>>` | Tham chiếu bằng mã định danh trong dữ liệu; không phải structural association UML. Không vẽ connector, không gán bội số; đây là quy ước chú thích, không phải UML Dependency. |')
    correction_heading = '\n## 10 Hiệu chỉnh trình bày ngày 7 tháng 10 năm 2026'
    text = text.split(correction_heading)[0]
    text += correction_heading + '\n\nBản hiệu chỉnh giữ 23 lớp và 39 association cấu trúc. Bỏ FK trùng association khỏi hình, lược tokenHash khỏi thuộc tính báo cáo, bổ sung thành phần tiền hóa đơn và danh sách ảnh minh chứng. Các ràng buộc unique có căn cứ được ghi trên hình; nhãn liên kết là tên vai trò từ lớp nguồn, còn PlantUML ghi tên vai trò tại các đầu association. Các lớp nhắc lại dùng cùng tên và thuộc tính của cùng classifier. CD-05 đặt Expense dưới MaintenanceTicket và dùng các tuyến nối không giao nhau; CD-06 dùng chú thích <<IDRef>>, không còn connector tham chiếu ID. Các nhãn bội số và vai trò trong Draw.io gắn với connector tương ứng.\n\nMột Expense được phép có nhiều MaintenanceTicket ở mức schema vì expenseId không unique. Luồng complete hiện tạo khoản chi mới, không cung cấp thao tác chọn một Expense đã tồn tại để chia sẻ. Việc cho phép chia sẻ khoản chi và đồng bộ actualCost với Expense.amount vẫn cần xác nhận, không được suy ra từ structural multiplicity. Các điểm PX-01, PX-02, PX-03, PX-14, PX-19, PX-24, PX-25, PX-26 và PX-32 không bị quyết định thay trong lần hiệu chỉnh này.\n\nXem `CLASS_DIAGRAM_AUDIT.md` để tra ma trận 39 quan hệ, 3 IDRef, danh sách 23 lớp, changelog, cấu trúc 11 hình và checklist. Không thay schema, code hay đặc tả Use Case.\n'
    apply_text(filename, text)
    proposal_name = 'PROPOSAL_QUAN_LY_TRO.md'
    proposal = (ROOT / proposal_name).read_text()
    start = proposal.index('### 2.4.3')
    end = proposal.index('### 2.4.4', start)
    section = replace_text(proposal[start:end])
    section = section.replace('| <<reference>> |', '| `<<reference>>` |')
    section = re.sub(r'^Dấu vân tay SHA256[^\n]*\n\n', '', section, flags=re.M)
    section = section.replace('`reversalOfPaymentId: String?`;', '').replace('`expenseId: String?`;', '').replace('`tokenHash: String`;', '')
    apply_text(proposal_name, proposal[:start] + section + proposal[end:])


def paragraphs(document):
    yield from document.paragraphs
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs


def update_word(filename):
    document = Document(BASE.parent / filename)
    for paragraph in paragraphs(document):
        replaced = replace_text(paragraph.text)
        if replaced != paragraph.text:
            properties = copy.deepcopy(paragraph.runs[0]._r.rPr) if paragraph.runs and paragraph.runs[0]._r.rPr is not None else None
            paragraph.text = replaced
            if properties is not None:
                paragraph.runs[0]._r.insert(0, properties)
    for table in document.tables:
        if len(table.columns) != 3 or not any('Thuộc tính' in cell.text for cell in table.rows[0].cells):
            continue
        for row in table.rows[1:]:
            name = row.cells[0].text.split('\n')[0].strip('` ')
            if name not in MODEL['included']:
                continue
            cell = row.cells[2]
            properties = copy.deepcopy(cell.paragraphs[0].runs[0]._r.rPr) if cell.paragraphs[0].runs and cell.paragraphs[0].runs[0]._r.rPr is not None else None
            cell.text = fields_text(name, filename == 'CLASS_DIAGRAM_SPECIFICATION.docx')
            if properties is not None:
                cell.paragraphs[0].runs[0]._r.insert(0, properties)
    image_order = [diagram['id'].lower() + '.png' for diagram in MODEL['diagrams']]
    if filename == 'CLASS_DIAGRAM_REPORT_SECTION.docx':
        image_order = [image_order[0], image_order[-1], *image_order[1:-1]]
    if len(document.inline_shapes) != len(image_order):
        raise ValueError(filename + ': unexpected number of figures')
    image_names = {'/word/media/image' + str(index) + '.png': name for index, name in enumerate(image_order, 1)}
    for shape in document.inline_shapes:
        relationship = shape._inline.graphic.graphicData.pic.blipFill.blip.embed
        part = document.part.related_parts[relationship]
        name = image_names.get(str(part.partname))
        if name is None:
            raise ValueError(filename + ': unknown figure part ' + str(part.partname))
        part._blob = (BASE / name).read_bytes()
        paragraph = shape._inline.xpath('ancestor::w:p')[0]
        while paragraph.getparent() is not document.element.body:
            paragraph = paragraph.getparent()
        siblings = list(document.element.body)
        section_properties = None
        for sibling in siblings[siblings.index(paragraph):]:
            candidates = sibling.xpath('.//w:sectPr') if sibling.tag != qn('w:sectPr') else [sibling]
            if candidates:
                section_properties = candidates[0]
                break
        page_size = section_properties.find(qn('w:pgSz'))
        margins = section_properties.find(qn('w:pgMar'))
        page_width, page_height = int(page_size.get(qn('w:w'))), int(page_size.get(qn('w:h')))
        if name in ['cd-06.png', 'od-02.png'] and page_width > page_height:
            page_size.set(qn('w:w'), str(page_height))
            page_size.set(qn('w:h'), str(page_width))
            page_size.attrib.pop(qn('w:orient'), None)
            page_width, page_height = page_height, page_width
        available_width = (page_width - int(margins.get(qn('w:left'))) - int(margins.get(qn('w:right')))) / 56.692913
        available_height = (page_height - int(margins.get(qn('w:top'))) - int(margins.get(qn('w:bottom')))) / 56.692913 - 34
        image = Image.open(BASE / name)
        width = min(available_width, available_height * image.width / image.height)
        shape.width, shape.height = Mm(width), Mm(width * image.height / image.width)
    advisory = 'Hiệu chỉnh ngày 7 tháng 10 năm 2026: các mã ngữ cảnh NotificationDelivery dùng chú thích <<IDRef>>, không vẽ đường nối UML. Cấu trúc 23 lớp và 39 association không thay đổi. Expense có thể được nhiều yêu cầu bảo trì tham chiếu ở mức schema; luồng hoàn tất tạo khoản chi mới, còn chính sách chia sẻ khoản chi và đồng bộ chi phí vẫn cần xác nhận. Các định danh liên kết bị lược khỏi hình vẫn được giữ trong schema và mã nguồn.'
    if filename == 'CLASS_DIAGRAM_REPORT_SECTION.docx':
        for paragraph in document.paragraphs:
            if paragraph.text.startswith(('Dấu vân tay SHA256', 'Hiệu chỉnh ngày 7 tháng 10 năm 2026:')):
                paragraph._element.getparent().remove(paragraph._element)
    elif not any(paragraph.text == advisory for paragraph in document.paragraphs):
        document.add_paragraph(advisory)
    document.save(BASE.parent / filename)


if __name__ == '__main__':
    update_markdown()
    for filename in ['CLASS_DIAGRAM_REPORT_SECTION.docx', 'CLASS_DIAGRAM_SPECIFICATION.docx']:
        update_word(filename)
    print('Updated class section only; no source, schema, Use Case, total Word or total PDF changes.')
