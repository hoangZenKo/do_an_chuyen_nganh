from pathlib import Path
from collections import Counter
from zipfile import ZipFile
import json,re,hashlib,unicodedata,sys
from lxml import etree
from docx import Document
from pypdf import PdfReader
ROOT=Path('/Users/hoang/Rental_room_management');BASE=Path('/tmp/rental-report-20261028');SUPPORT=ROOT/'docs/report_2026-10-28'
FINAL=ROOT/'docs/Bao_cao_Quan_ly_Nha_tro_hoan_chinh_2026-10-28.docx'
render=BASE/(sys.argv[1] if len(sys.argv)>1 else 'final-render4')
doc=Document(FINAL);data=json.loads((BASE/'audit.json').read_text());diagrams=json.loads((BASE/'diagram-manifest.json').read_text())
reader=PdfReader(str(next(render.glob('*.pdf'))));texts=[page.extract_text() or '' for page in reader.pages]
ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
with ZipFile(FINAL) as archive:
    xml=etree.fromstring(archive.read('word/document.xml'))
uc_tables=[]
for table in doc.tables:
    if len(table.columns)==2 and table.cell(0,0).text=='Thuộc tính' and table.cell(0,1).text=='Đặc tả':
        rows={row.cells[0].text:row.cells[1].text for row in table.rows[1:]}
        uc_tables.append((re.match(r'UC-[A-Z]+-\d+(?:\.\d+)?',rows['Mã và tên']).group(),rows))
expected_fields={'Mã và tên','Tác nhân','Mục tiêu','Kích hoạt','Điều kiện trước','Đầu vào','Luồng chính','Luồng thay thế','Ngoại lệ','Điều kiện sau','Quy tắc nghiệp vụ','Căn cứ triển khai'}
ids=[item[0] for item in uc_tables];expected_ids=[case['id'] for case in data['cases']]
checks={
    'uc_specs_exactly_132':len(ids)==132 and set(ids)==set(expected_ids),
    'uc_ids_unique':len(ids)==len(set(ids)),
    'uc_fields_complete':all(set(rows)==expected_fields and all(value.strip() for value in rows.values()) for key,rows in uc_tables),
    'uc_source_files_exist':all((ROOT/case['source']).is_file() for case in data['cases']),
    'frozen_source_evidence_matches_audit':all(hashlib.sha256((SUPPORT/'evidence'/case['source']).read_bytes()).hexdigest()==case['source_hash'] for case in data['cases']),
    'diagram_sources_and_pngs_exist':all((ROOT/d['source']).exists() and (BASE/'diagrams'/(d['id']+'.png')).is_file() for d in diagrams),
    'diagram_count_64':len(diagrams)==64,
    'tables_white_only':all(node.get('{'+ns['w']+'}fill') in ('FFFFFF','auto',None) for node in xml.xpath('//w:tbl//w:shd',namespaces=ns)),
    'table_borders_black':all(node.get('{'+ns['w']+'}color')=='000000' for node in xml.xpath('//w:tblBorders/*',namespaces=ns)),
    'all_table_headers_repeat':all(table.rows[0]._tr.xpath('./w:trPr/w:tblHeader') for table in doc.tables),
    'two_template_sections':len(doc.sections)==2,
    'template_page_sizes_preserved':[round(section.page_height.inches,2) for section in doc.sections]==[14,11] and all(round(section.page_width.inches,2)==8.5 for section in doc.sections),
    'all_render_pages_present':len(list(render.glob('page-*.png')))==len(reader.pages),
    'no_blank_body_pages':all(len(text.strip())>30 for text in texts),
    'no_schema_or_route_appendix':not any(phrase in '\n'.join(texts) for phrase in ['Danh mục mô hình dữ liệu','Quan hệ và trạng thái dữ liệu','Danh mục endpoint kỹ thuật']),
    'crosswalk_all_67':len(json.loads((BASE/'uc-crosswalk.json').read_text()))==67,
}
entries=json.loads((BASE/'toc-entries.json').read_text());toc_pages=json.loads((BASE/'toc-pages.json').read_text())
def norm(text):return re.sub(r'\s+',' ',unicodedata.normalize('NFC',text)).strip()
toc_actual={};missing=[]
for anchor,title,level in entries:
    hits=[i for i,text in enumerate(texts) if i>=3 and norm(title) in norm(text)]
    if hits:toc_actual[title]=hits[0]
    else:missing.append(title)
checks['toc_measured_pages_match']=not missing and toc_actual==toc_pages
checks['toc_no_unfilled_entries']=all(str(value).isdigit() for value in toc_pages.values()) and len(toc_pages)==len(entries)
checks['toc_links_resolve']=all(node.get('{'+ns['w']+'}anchor') in {bookmark.get('{'+ns['w']+'}name') for bookmark in xml.xpath('//w:bookmarkStart',namespaces=ns)} for node in xml.xpath('//w:hyperlink[@w:anchor]',namespaces=ns))
originals={}
for filename in ['CSE703014_Report_Template_VI_V1.docx','Bao_cao_da_chinh_form.docx']:
    file=ROOT/'docs'/filename; inventory=json.loads((BASE/(file.stem+'-parts.json')).read_text())
    with ZipFile(file) as archive:
        current={part.filename:hashlib.sha256(archive.read(part.filename)).hexdigest() for part in archive.infolist()}
    checks['original_unchanged_'+file.stem]=current=={part['part']:part['sha256'] for part in inventory}
    originals[filename]=hashlib.sha256(file.read_bytes()).hexdigest()
report={
    'date':'2026-10-28',
    'output':str(FINAL),
    'sha256':hashlib.sha256(FINAL.read_bytes()).hexdigest(),
    'render_directory':str(render),
    'rendered_pages':len(reader.pages),
    'use_cases':len(ids),
    'uc_tables_fields':12,
    'tables':len(doc.tables),
    'inline_images':len(doc.inline_shapes),
    'diagrams':dict(Counter(d['kind'] for d in diagrams)),
    'toc_entries':len(entries),
    'checks':checks,
    'all_checks_passed':all(checks.values()),
    'originals_sha256':originals,
    'tests':{'backend':{'passed':672,'skipped':7,'exit_code':0},'backend_build':{'exit_code':0},'flutter_features_data':{'passed':503,'failed':12,'exit_code':1},'admin_capture':{'passed':51,'scope':'four focused widget files; fixtures'}},
    'limitations':['Administrative names, student IDs, class, supervisor, and actual task assignments not provided.','No DB-reset E2E, migration, provider/device acceptance, real payments or production deployment performed.','Flutter test failures are documented; no claim that the application passed all tests.'],
}
SUPPORT.mkdir(exist_ok=True)
(SUPPORT/'verification.json').write_text(json.dumps(report,ensure_ascii=False,indent=2))
(BASE/'verification.json').write_text(json.dumps(report,ensure_ascii=False,indent=2))
print(json.dumps({k:v for k,v in report.items() if k not in ['originals_sha256','limitations']},ensure_ascii=False,indent=2))
if not all(checks.values()):sys.exit(1)
