from pathlib import Path
from collections import Counter
from copy import deepcopy
from zipfile import ZipFile
import sys, json, re, math, hashlib, textwrap, shutil
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont

ROOT=Path('/Users/hoang/Rental_room_management')
TMP=Path('/tmp/rental-report-20261028')
sys.path.insert(0,str(TMP))
from report_data import GROUPS, MODEL_GROUPS, SEQUENCES
from narrative import INTRO,NFR,CONSTRAINTS,BUSINESS_RULES,GLOSSARY,FR_GROUPS,QUALITY_SCENARIOS,CHANGE_LOG
DATA=json.loads((TMP/'audit.json').read_text())
CASES=DATA['cases']; MODELS=DATA['models']; ROUTES=DATA['routes']
DIAGRAMS=json.loads((TMP/'diagram-manifest.json').read_text())
FINAL=ROOT/'docs/Bao_cao_Quan_ly_Nha_tro_hoan_chinh_2026-10-28.docx'
TOC_PAGES=json.loads((TMP/'toc-pages.json').read_text()) if (TMP/'toc-pages.json').exists() else {}
DOC=Document(ROOT/'docs/CSE703014_Report_Template_VI_V1.docx')
body=DOC.element.body
for child in list(body):
    if child.tag!=qn('w:sectPr'):body.remove(child)

def remove_children(parent,tag):
    for el in list(parent.findall(qn(tag))):parent.remove(el)

for style in DOC.styles:
    if style.type!=1:continue
    style.font.name='Times New Roman';style.font.color.rgb=RGBColor(0,0,0)
    if style.element.pPr is not None:
        remove_children(style.element.pPr,'w:numPr')
        remove_children(style.element.pPr,'w:pBdr')
    if style.name.startswith('Heading'):
        level=int(style.name.split()[-1]) if style.name.split()[-1].isdigit() else 4
        style.font.size=Pt(16 if level==1 else 14 if level==2 else 13 if level==3 else 12)
        style.font.bold=True
        style.paragraph_format.space_before=Pt(14 if level<=2 else 10)
        style.paragraph_format.space_after=Pt(8)
        style.paragraph_format.keep_with_next=True
        style.paragraph_format.line_spacing=1.1
        style.paragraph_format.left_indent=Pt(0)
        style.paragraph_format.first_line_indent=Pt(0)
        style.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
normal=DOC.styles['Normal'];normal.font.size=Pt(13)
normal.paragraph_format.space_after=Pt(7);normal.paragraph_format.line_spacing=1.12
normal.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
normal.paragraph_format.first_line_indent=Inches(0.25)
for style_name in ('Title',):
    style=DOC.styles[style_name];style.font.name='Times New Roman';style.font.color.rgb=RGBColor(0,0,0)

def font(run,size=13,bold=None,italic=None):
    run.font.name='Times New Roman';run.font.size=Pt(size);run.font.color.rgb=RGBColor(0,0,0)
    if bold is not None:run.bold=bold
    if italic is not None:run.italic=italic
    rpr=run._r.get_or_add_rPr();fonts=rpr.find(qn('w:rFonts'))
    if fonts is None:fonts=OxmlElement('w:rFonts');rpr.append(fonts)
    for key in ['ascii','hAnsi','eastAsia','cs']:fonts.set(qn('w:'+key),'Times New Roman')
    lang=OxmlElement('w:lang');lang.set(qn('w:val'),'vi-VN');rpr.append(lang)

TOC=[];BOOKMARKS=[]
def heading(text,level=1,new_page=False,toc=True):
    p=DOC.add_paragraph(style=f'Heading {level}')
    if new_page:p.paragraph_format.page_break_before=True
    run=p.add_run(text);font(run,16 if level==1 else 14 if level==2 else 13 if level==3 else 12,True)
    index=len(BOOKMARKS)+1;name=f'sec_{index:04}'
    start=OxmlElement('w:bookmarkStart');start.set(qn('w:id'),str(index));start.set(qn('w:name'),name)
    end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),str(index));p._p.insert(0,start);p._p.append(end)
    BOOKMARKS.append((name,text,level))
    if toc and level<=3:TOC.append((name,text,level))
    return p

def para(text='',bold=False,italic=False,small=False,align=None,keep=False):
    p=DOC.add_paragraph();font(p.add_run(text),11.5 if small else 13,bold,italic)
    if align is not None:p.alignment=align
    p.paragraph_format.keep_with_next=keep
    return p

def bullet(text):
    p=para('• '+text);p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.left_indent=Inches(.12)
    return p

def table(headers,rows,widths=None,size=11.5,keep_rows=True):
    t=DOC.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    if widths is None:widths=[6.5/len(headers)]*len(headers)
    for col,width in zip(t.columns,widths):col.width=Inches(width)
    pr=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'5');e.set(qn('w:color'),'000000');borders.append(e)
    pr.append(borders)
    margins=OxmlElement('w:tblCellMar')
    for side,val in [('top','65'),('bottom','65'),('left','85'),('right','85')]:
        e=OxmlElement('w:'+side);e.set(qn('w:w'),val);e.set(qn('w:type'),'dxa');margins.append(e)
    pr.append(margins)
    layout=OxmlElement('w:tblLayout');layout.set(qn('w:type'),'fixed');pr.append(layout)
    allrows=[headers]+list(rows)
    for idx,values in enumerate(allrows):
        row=t.rows[0] if idx==0 else t.add_row()
        rowpr=row._tr.get_or_add_trPr()
        if idx==0:
            repeat=OxmlElement('w:tblHeader');repeat.set(qn('w:val'),'true');rowpr.append(repeat)
        if keep_rows:rowpr.append(OxmlElement('w:cantSplit'))
        for cell,value,width in zip(row.cells,values,widths):
            cell.width=Inches(width);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
            shd=OxmlElement('w:shd');shd.set(qn('w:val'),'clear');shd.set(qn('w:fill'),'FFFFFF');cell._tc.get_or_add_tcPr().append(shd)
            p=cell.paragraphs[0];p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0)
            p.paragraph_format.line_spacing=1.0;p.paragraph_format.keep_with_next=False;p.alignment=WD_ALIGN_PARAGRAPH.LEFT
            font(p.add_run(str(value)),size,bold=(idx==0))
    DOC.add_paragraph().paragraph_format.space_after=Pt(0)
    return t

def picture(path,caption,max_width=6.4,max_height=6.65):
    image=Image.open(path);ratio=image.width/image.height
    width=min(max_width,max_height*ratio)
    p=DOC.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Pt(0)
    p.paragraph_format.space_before=Pt(5);p.paragraph_format.space_after=Pt(5);p.paragraph_format.keep_with_next=True
    p.add_run().add_picture(str(path),width=Inches(width))
    cap=para(caption,italic=True,small=True,align=WD_ALIGN_PARAGRAPH.CENTER)
    cap.paragraph_format.first_line_indent=Pt(0)
    return width

def section_footer(section,start=None):
    section.header.is_linked_to_previous=False;section.footer.is_linked_to_previous=False
    for p in section.header.paragraphs:p.text=''
    p=section.footer.paragraphs[0];p.text='';p.alignment=WD_ALIGN_PARAGRAPH.RIGHT;p.paragraph_format.first_line_indent=Pt(0)
    r=p.add_run();font(r,11)
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');p._p.append(field)
    if start is not None:
        for old in section._sectPr.findall(qn('w:pgNumType')):section._sectPr.remove(old)
        num=OxmlElement('w:pgNumType');num.set(qn('w:start'),str(start));num.set(qn('w:fmt'),'decimal');section._sectPr.append(num)

cover=DOC.sections[0];cover.page_width=Inches(8.5);cover.page_height=Inches(14)
for attr in ['top_margin','bottom_margin','left_margin','right_margin']:setattr(cover,attr,Inches(1))
cover.header_distance=cover.footer_distance=Inches(.5);cover.different_first_page_header_footer=True
for h in [cover.header,cover.first_page_header,cover.footer,cover.first_page_footer]:
    for p in h.paragraphs:p.text=''
for text,size,bold in [('TRƯỜNG ĐẠI HỌC PHENIKAA',20,True),('KHOA CÔNG NGHỆ THÔNG TIN',19,True)]:
    p=para(text,align=WD_ALIGN_PARAGRAPH.CENTER);p.paragraph_format.first_line_indent=Pt(0)
    p.paragraph_format.space_after=Pt(15)
    for run in p.runs:font(run,size,bold)
p=DOC.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Pt(0)
p.add_run().add_picture(str(TMP/'phenikaa-logo.png'),width=Inches(2.5))
p.paragraph_format.space_after=Pt(48)
for text,size in [('BÁO CÁO ĐỒ ÁN LIÊN NGÀNH',19),('HỆ THỐNG QUẢN LÝ NHÀ TRỌ',19),('Rental Room Management',16)]:
    p=DOC.add_paragraph(style='Title' if text=='HỆ THỐNG QUẢN LÝ NHÀ TRỌ' else 'Normal');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Pt(0)
    p.paragraph_format.space_after=Pt(16)
    font(p.add_run(text),size,True)
para('Bản đối chiếu mã nguồn và hoàn thiện theo mẫu V1',italic=True,align=WD_ALIGN_PARAGRAPH.CENTER)
para('Cập nhật nội dung ngày 28 tháng 10 năm 2026',align=WD_ALIGN_PARAGRAPH.CENTER)
p=para('Thông tin hành chính',bold=True,align=WD_ALIGN_PARAGRAPH.CENTER);p.paragraph_format.space_before=Pt(35)
for text in ['Mã nhóm, thành viên, mã sinh viên và lớp: chưa có trong nguồn cung cấp.','Giảng viên hướng dẫn: chưa có trong nguồn cung cấp.']:
    para(text,small=True,align=WD_ALIGN_PARAGRAPH.CENTER)
para('Không tự suy đoán hoặc gán danh tính từ siêu dữ liệu tài liệu.',italic=True,small=True,align=WD_ALIGN_PARAGRAPH.CENTER)
p=para('Hà Nội, tháng 10 năm 2026',align=WD_ALIGN_PARAGRAPH.CENTER);p.paragraph_format.space_before=Pt(45)

body_section=DOC.add_section(WD_SECTION_START.NEW_PAGE)
body_section.page_width=Inches(8.5);body_section.page_height=Inches(11)
for attr in ['top_margin','bottom_margin','left_margin','right_margin']:setattr(body_section,attr,Inches(1))
body_section.header_distance=body_section.footer_distance=Inches(.5);body_section.different_first_page_header_footer=False
section_footer(body_section,1)
toc_title=para('MỤC LỤC',bold=True,align=WD_ALIGN_PARAGRAPH.CENTER)
toc_title.paragraph_format.space_after=Pt(18)
toc_anchor=DOC.add_paragraph()

heading('Phân công thực hiện đồ án',1,new_page=True)
para('Mẫu V1 yêu cầu bảng phân công theo từng sinh viên. Tài liệu và repository được cung cấp chưa có danh sách thành viên, mã sinh viên, lớp hoặc phân công được xác nhận. Vì vậy, báo cáo giữ mục này nhưng không tự gán tác giả hoặc khẳng định đóng góp cá nhân. Bảng dưới là cấu trúc trách nhiệm kỹ thuật cần chủ nhiệm nhóm xác nhận trước khi nộp.')
table(['Mã SV / họ tên / lớp','Gói công việc cần xác nhận','Minh chứng có thể đối chiếu'],[
('Chưa được cung cấp','Phân tích yêu cầu, UC và truy vết','Báo cáo, đặc tả, commit/issue nếu có'),
('Chưa được cung cấp','Flutter và trải nghiệm theo vai trò','lib/, test/, ảnh widget'),
('Chưa được cung cấp','Backend, schema và toàn vẹn nghiệp vụ','backend/src/, prisma/, unit/E2E'),
('Chưa được cung cấp','Kiểm chứng, vận hành và hoàn thiện tài liệu','Log test, runbook, bảng lỗi, nguồn hình')],[1.55,2.5,2.45])
para('Phạm vi đóng góp phải được xác nhận bằng hồ sơ nhóm. Mốc thời gian tham chiếu và ma trận kỹ năng ở chương 3 là kế hoạch đề xuất, không thay thế biên bản phân công hoặc nhật ký làm việc thực tế.',small=True)

heading('1 Giới thiệu',1,new_page=True)
for title,paragraphs in INTRO:
    heading(title,2)
    for text in paragraphs:para(text)
    if title.startswith('1.2'):
        table(['Phương án','Điểm phù hợp','Giới hạn cần xử lý'],[
            ('Sổ và ghi chú','Khởi đầu đơn giản, ít thiết bị','Tra cứu, cộng dồn và lịch sử phụ thuộc người ghi.'),
            ('Bảng tính','Công thức, lọc và tổng hợp thuận tiện','Khó gắn quyền theo hồ sơ/hợp đồng và kiểm soát ghi đồng thời.'),
            ('Tin nhắn','Trao đổi nhanh, quen thuộc','Không có sổ giao dịch có trạng thái và ràng buộc đầy đủ.'),
            ('Hệ thống tập trung của đồ án','Liên kết dữ liệu và quyền theo luồng thuê','Cần vận hành API/DB, quản lý bảo mật, backup và nghiệm thu tích hợp.')],[1.5,2.3,2.7])

heading('2 Thiết kế và triển khai',1,new_page=True)
heading('2.1 Các yêu cầu chức năng',2)
para(f'Kiểm kê hiện tại xác định {len(CASES)} đặc tả UC thuộc 10 nhóm nghiệp vụ. UC là mục tiêu tác nhân, không phải mỗi HTTP endpoint là một UC riêng. Một UC có thể dùng nhiều route để đọc, xác nhận và lưu; các endpoint hỗ trợ kỹ thuật được liệt kê độc lập trong phụ lục. Yêu cầu chức năng dưới đây được triển khai hoặc hỗ trợ trong source; mức kiểm chứng phải đọc cùng mục 2.6.')
table(['Mã','Nhóm yêu cầu và phạm vi','UC liên quan'],[(code,title+'\n'+description,group+' ('+str(sum(c['group']==group for c in CASES))+' UC)') for code,title,description,group in FR_GROUPS],[.78,4.6,1.12])
heading('2.1.1 Tác nhân và phạm vi truy cập',3)
table(['Tác nhân','Mục tiêu','Giới hạn'],[
('Khách','Tìm/lọc và xem tin công khai','Không có role trong DB; không xem hồ sơ thuê riêng.'),
('Người tìm phòng','Yêu thích, lịch xem, bộ lọc đã lưu','Dùng role TENANT; chưa mặc nhiên có Lease/Tenant liên kết.'),
('Người thuê','Xem dữ liệu của mình, proof, bảo trì và lifecycle','User–Tenant được liên kết; không xem tenant khác.'),
('Chủ trọ','Vận hành toàn bộ dữ liệu thuộc mình','Không dùng ownership của chủ trọ khác; settlement/tariff approval theo quy tắc.'),
('Staff','Thực hiện chức năng được giao theo khu','Permission + propertyIds + controller allowlist; không thừa hưởng mọi luồng LANDLORD.'),
('Admin','Quản trị tài khoản/nội dung/hỗ trợ/vận hành','Capability, MFA/step-up theo policy; inventory không sửa tài chính của chủ trọ.'),
('Provider và worker','Xử lý tác vụ tích hợp hoặc định kỳ','Tác nhân kỹ thuật; phụ thuộc cấu hình, không thay vai trò con người.')],[1.05,2.65,2.8])
heading('2.1.2 Thuật ngữ và quy tắc nghiệp vụ',3)
table(['Thuật ngữ','Định nghĩa trong dự án'],GLOSSARY,[1.45,5.05])
table(['Mã','Quy tắc','Diễn giải'],BUSINESS_RULES,[.65,1.65,4.2])

heading('2.2 Các yêu cầu phi chức năng',2,new_page=True)
para('Các yêu cầu sau là tiêu chí thiết kế và kiểm chứng. Không tự đặt số liệu SLA, tải đồng thời, độ trễ, độ sẵn sàng hoặc chứng nhận accessibility khi repository chưa có phép đo tương ứng. Cột căn cứ giúp xác định cơ chế hiện hữu và nơi cần bổ sung kiểm thử.')
table(['Mã / tiêu chí','Yêu cầu','Cơ chế và mức bằng chứng'],[(code+'\n'+name,goal,evidence) for code,name,goal,evidence in NFR],[1.08,2.55,2.87])
heading('2.3 Các ràng buộc',2,new_page=True)
for title,texts in CONSTRAINTS.items():
    heading(title,3)
    for text in texts:para(text)
table(['Hạng mục cần dự toán','Cơ sở xác định','Trạng thái số liệu'],[
('Phát triển và kiểm thử','Giờ công, máy phát triển, thiết bị thử','Chưa có ngân sách/biên bản được cung cấp'),
('API và cơ sở dữ liệu','CPU/RAM, dung lượng, backup và lưu lượng','Ước tính sau khi có cấu hình và phép đo tải'),
('Tệp và delivery','Dung lượng lưu, retention, email/Zalo/provider','Chưa xác minh giá hoặc chi phí thực trả'),
('AI và quan sát vận hành','Mô hình, phần cứng, log/metric/alert','Phụ thuộc cấu hình thử nghiệm và release')],[1.8,2.65,2.05])

heading('2.4 Mô hình hệ thống và thiết kế giải pháp',2,new_page=True)
para('Thiết kế được đối chiếu từ AppModule, controller/service, các model Flutter và schema Prisma, thay vì lấy sơ đồ cũ làm nguồn sự thật. Mô hình trình bày ở ba mức: cấu trúc triển khai, tương tác mục tiêu người dùng và thực thể dữ liệu. Trọng tâm trình bày là các miền nghiệp vụ được mô tả qua UC, lớp thực thể và tương tác; route kỹ thuật chỉ dùng để truy vết, không phải mỗi route là một chức năng độc lập.')
for d in [d for d in DIAGRAMS if d['kind']=='architecture']:
    picture(TMP/'diagrams'/(d['id']+'.png'),d['id']+' '+d['title'],max_height=5.7)
    para(d['description'],small=True)
heading('2.4.1 Các kịch bản của hệ thống',3,new_page=True)
para('Mỗi đặc tả gồm mã, tác nhân, mục tiêu, kích hoạt, điều kiện trước/sau, dữ liệu đầu vào, luồng chính, thay thế, ngoại lệ, quy tắc và căn cứ triển khai. Mã mới phân nhóm SYS, DIS, OPT, TNT, BIL, OPS, TEN, STF, ADM và SUP. Việc đánh lại mã nhằm loại bỏ số chương bị lẫn với tài liệu con; phụ lục A nêu thay đổi so với bản dở.')
para('Ký hiệu A là nhánh hợp lệ hoặc hủy trước khi xác nhận; E là lỗi dữ liệu/quyền/trạng thái. Các phản hồi 400/401/403/404/409/429 chỉ được diễn giải theo contract từng endpoint, không áp đồng loạt một mã cho mọi lỗi. Luồng đọc không có tác dụng ghi sổ; luồng ghi nhiều thực thể phải theo đơn vị transaction của service.')
for group,meta in GROUPS.items():
    heading(meta[0],4,toc=False)
    table(['Mã UC','Tên mục tiêu'],[(c['id'],c['title']) for c in CASES if c['group']==group],[1.16,5.34])

for case in CASES:
    heading(case['id']+' '+case['title'],4,new_page=True,toc=False)
    actors=case['actor']
    if case['group']=='TEN':pre='Phiên TENANT hợp lệ; tài nguyên thuộc hồ sơ liên kết; quyền được kiểm tra lại tại máy chủ.'
    elif case['group']=='ADM':pre='Phiên ADMIN hợp lệ, capability tương ứng; đáp ứng step-up/MFA nếu policy của hành động yêu cầu.'
    elif case['group']=='SUP':pre='Phiên ADMIN hợp lệ, capability và phạm vi hồ sơ hỗ trợ hoặc vận hành tương ứng; đáp ứng policy của hành động.'
    elif case['group']=='SYS':pre='Truy cập được ứng dụng/API; có tài khoản/token/lời mời tương ứng với luồng. Luồng công khai không yêu cầu phiên sẵn có.'
    elif case['group']=='DIS':pre='Tin/phòng đủ điều kiện theo luồng; các thao tác cá nhân yêu cầu phiên đúng vai trò, không chỉ xem công khai.'
    else:pre='Tài nguyên thuộc chủ trọ; tác nhân có quyền nghiệp vụ. Staff còn phải đúng khu được giao và controller hỗ trợ scope.'
    if case['id']=='UC-SUP-06':pre='Phiên LANDLORD hợp lệ; khu được hỏi thuộc chủ trọ; kiểm tra giới hạn tần suất và quyền truy cập nguồn tại máy chủ.'
    if case['id']=='UC-ADM-18':pre='Chủ trọ có phiên LANDLORD để nộp/xem hồ sơ của mình; bước xét duyệt dùng phiên ADMIN có capability xác minh và support.read. Chứng cứ gắn đúng ticket hỗ trợ.'
    trigger=case['steps'][0]+'.'
    steps='\n'.join(f'{i+1}. {s.rstrip(".")}.' for i,s in enumerate(case['steps']))
    if case['id'] in {'UC-DIS-01','UC-DIS-02','UC-OPT-01','UC-OPT-02','UC-OPT-05','UC-TNT-01','UC-TNT-04','UC-TNT-13','UC-BIL-04','UC-BIL-08.2','UC-BIL-20','UC-BIL-21','UC-TEN-01','UC-TEN-02','UC-TEN-03','UC-STF-03','UC-ADM-01','UC-ADM-02','UC-ADM-08','UC-ADM-10','UC-ADM-14'}:
        after='Hiển thị dữ liệu/kết quả đúng quyền và phạm vi; chỉ việc xem/xuất không tự thay đổi sổ tiền hoặc điều khoản.'
        alt='A1. Không có dữ liệu: trả trạng thái rỗng phù hợp. A2. Đổi bộ lọc hoặc chọn tài nguyên khác: truy vấn lại theo scope.'
    else:
        after='Thành công: lưu kết quả của hành động và hiển thị trạng thái mới. Hủy/lỗi: không coi yêu cầu chưa được xác nhận là đã hoàn thành.'
        alt='A1. Người dùng hủy trước xác nhận: không gửi lệnh ghi. A2. Dữ liệu đã đổi: đọc lại trạng thái rồi quyết định thao tác tiếp.'
    if 'requestKey' in case['inputs'] or 'idempot' in case['rules'].lower():alt='A1. Retry cùng requestKey và nội dung: nhận lại kết quả theo contract. A2. Kết quả mạng chưa rõ: đối chiếu bản ghi trước khi tạo lệnh mới.'
    if case['id']=='UC-SUP-06':
        alt='A1. Không chọn khu: truy vấn trong phạm vi chủ trọ. A2. Chưa đủ căn cứ: nêu giới hạn và đề nghị bổ sung thông tin, không suy đoán số liệu.'
        after='Hiển thị câu trả lời và nguồn được phép hoặc thông báo giới hạn; không tự thay đổi sổ tiền, hợp đồng hoặc dữ liệu nghiệp vụ.'
    rows=[('Mã và tên',case['id']+' — '+case['title']),('Tác nhân',actors),('Mục tiêu','Thực hiện '+case['title'][0].lower()+case['title'][1:]+' đúng quyền và trạng thái nghiệp vụ.'),('Kích hoạt',trigger),('Điều kiện trước',pre),('Đầu vào',case['inputs']+'.'),('Luồng chính',steps),('Luồng thay thế',alt),('Ngoại lệ',case['error']),('Điều kiện sau',after),('Quy tắc nghiệp vụ',case['rules']),('Căn cứ triển khai',case['source']+'\nMàn hình/contract và kiểm thử liên quan: xem phụ lục B.')]
    table(['Thuộc tính','Đặc tả'],rows,[1.2,5.3],size=11.5)

heading('2.4.2 Mô hình Use case',3,new_page=True)
para('Các sơ đồ dưới phân rã theo nhóm và đặt UC trong biên hệ thống. Tác nhân là vai trò ngoài hệ thống; database và Flutter không được dùng thay tác nhân nghiệp vụ. Đường nối actor–UC là association không hướng. Không vẽ include/extend chỉ để nối hai chức năng cùng màn hình; điều kiện trước và các thao tác kỹ thuật dùng chung được mô tả trong đặc tả.')
para('Các tên UC dài được xuống dòng trong hình nhưng mã giữ nguyên. Mỗi đường association nối đúng tác nhân của UC, không nối tất cả vai trò của nhóm vào mọi chức năng. Nhãn Staff còn phải đọc cùng giới hạn permission và scope trong đặc tả; đọc impact biểu phí không cho phép duyệt biểu phí. Chatbot được nối riêng với chủ trọ, không với Admin.')
for index,d in enumerate(d for d in DIAGRAMS if d['kind']=='usecase'):
    heading(d['id']+' '+d['title'],4,new_page=True,toc=False)
    picture(TMP/'diagrams'/(d['id']+'.png'),'Hình '+d['id']+' '+d['title'],max_height=6.55)
    ids=[]
    group=d['id'].split('-')[1];part=int(d['id'].split('-')[-1])-1
    ids=[c['id'] for c in CASES if c['group']==group][part*7:part*7+7]
    para('Đặc tả liên quan: '+', '.join(ids)+'.',small=True)

heading('2.4.3 Mô hình lớp và đối tượng',3,new_page=True)
para('Mô hình lớp mô tả cấu trúc các miền nghiệp vụ: tài khoản và phân quyền; khu/phòng và tìm thuê; hồ sơ/hợp đồng; hóa đơn/sổ thu; vận hành; quản trị/hỗ trợ. Hình giúp giải thích thực thể nào tham gia từng UC và các liên kết giữa chúng. Không dùng composition cho mọi liên kết chỉ vì việc xóa ở database có Cascade.')
para('Các sơ đồ CD chỉ trình bày các thực thể và thuộc tính chính cần để giải thích nghiệp vụ; không phải sơ đồ mọi class UI/controller/service và không chép toàn bộ schema database. Một thực thể có thể lặp trong nhiều hình để thể hiện điểm nối giữa các miền. Ký hiệu quan hệ được đối chiếu với service và dữ liệu hiện tại, không tự suy ra quan hệ chỉ vì tên field giống nhau.')
table(['Quan hệ / ràng buộc','Diễn giải'],[
('1 ở đầu đối tượng cha','Một bản ghi con bắt buộc tham chiếu đúng một cha qua khóa ngoại không nullable.'),('0..1 ở đầu đối tượng cha','Quan hệ tùy chọn, khóa ngoại có thể null.'),('0..* ở đầu đối tượng con','Cha có thể chưa có hoặc có nhiều bản ghi con; không giới hạn một hợp đồng lịch sử.'),('0..1 ở đầu con khi FK unique','Cha có tối đa một bản ghi con trong quan hệ duy nhất này.'),('Quy tắc theo trạng thái','Ví dụ một Lease ACTIVE/phòng thuộc logic service, không thay bội số lịch sử Room–Lease.'),('Quan hệ tự tham chiếu','Một Payment đảo tham chiếu khoản gốc; một Invoice thay thế giữ liên kết bản trước. Bỏ vòng lặp trong hình nhóm để tránh chồng nhãn; ý nghĩa trình bày ở UC-BIL-06 và UC-BIL-12.')],[2.2,4.3])
for d in [d for d in DIAGRAMS if d['kind']=='class']:
    heading(d['id']+' '+d['title'],4,new_page=True,toc=False)
    picture(TMP/'diagrams'/(d['id']+'.png'),'Hình '+d['id']+' '+d['title'],max_height=6.75)
    para(d['description'],small=True)
    modelnames=next(names for key,title,names in MODEL_GROUPS if key==d['id'])
    para('Các thực thể chính: '+', '.join(modelnames)+'.',small=True)
heading('Đối tượng minh họa và phép đối chiếu số dư',4,new_page=True,toc=False)
obj=next(d for d in DIAGRAMS if d['kind']=='object')
picture(TMP/'diagrams'/'OD-01.png','Hình OD-01 '+obj['title'],max_height=5.7)
para(obj['description'])
para('Công thức đối chiếu ví dụ: tiền receipt 3.000.000 = phân bổ đang hiệu lực 2.500.000 + khả dụng 500.000, khi chưa có hoàn tiền. Công nợ invoice 4.000.000 − 2.500.000 = 1.500.000. Nếu giải phóng toàn bộ phân bổ, một reversal âm được thêm và phần khả dụng tăng lại; chỉ sau đó mới có thể ghi nhận hoàn phần tương ứng. Đây là dữ liệu giả định để giải thích mô hình, không phải một giao dịch thực.')

heading('2.4.4 Các biểu đồ tuần tự',3,new_page=True)
para(str(len(SEQUENCES))+' sơ đồ SD chọn các luồng có rủi ro và tương tác liên tầng quan trọng nhất. Thời gian đi từ trên xuống, đường nét đứt là lifeline, thông điệp có số thứ tự. Các lifeline nhóm như Portal/Billing hoặc Prisma/Uploads là ranh giới mức thiết kế, không phải khẳng định một class cụ thể có đúng tên đó. Việc đổi từ người thuê sang chủ trọ hoặc maker sang checker luôn là một phiên đã xác thực khác.')
para('Sơ đồ trình bày luồng thành công; nhánh thay thế/ngoại lệ và điểm transaction được giải thích ngay trong hình và trong UC. Hai thông điệp có nhãn lựa chọn không được diễn giải là luôn cùng thực thi. Không đưa password/token/khóa bí mật vào payload minh họa. Tham chiếu mã cha, chẳng hạn UC-TNT-06, bao gồm các đặc tả con UC-TNT-06.1, UC-TNT-06.2 và UC-TNT-06.3 trong cùng nhóm mục tiêu.')
for d,spec in zip([d for d in DIAGRAMS if d['kind']=='sequence'],SEQUENCES):
    heading(d['id']+' '+d['title'],4,new_page=True,toc=False)
    picture(TMP/'diagrams'/(d['id']+'.png'),'Hình '+d['id']+' '+d['title'],max_height=4.5)
    para('UC: '+d['refs']+'. Căn cứ: '+d['source']+'.',small=True)
    table(['Bước','Diễn giải xử lý'],[(str(i+1),message[2]) for i,message in enumerate(spec[5])],[.6,5.9],size=10.5)

heading('2.4.5 Các màn hình giao diện người dùng',3,new_page=True)
para('Phần này mô tả đầy đủ nhóm màn hình qua route và trách nhiệm, kèm ảnh của các màn hình mới có cơ chế capture tái lập. Ảnh Admin được render từ Flutter widget tests của source hiện tại với ADMIN_QA_OUTPUT; dữ liệu là fixture. Ảnh billing có sẵn được ghi nguồn riêng. Không ảnh nào trong mục này được dùng thay bằng chứng E2E trên simulator/thiết bị hoặc provider thật.')
screen_rows=[
('Khởi động và tài khoản','Splash, onboarding, login, register, forgot, invite, bảo mật','Session restore; role routing; xử lý loading/expired/denied.'),
('Khám phá phòng','Explore, listing detail, saved, searches, viewing, review','Khách xem công khai; thao tác cá nhân đúng phiên TENANT.'),
('Dashboard chủ trọ','Dashboard, room map, chỉ số và action destinations','Property filter; chuyển tới đúng room/invoice ID.'),
('Khu/phòng','Danh sách, chi tiết, form, assets, inspection, tariff','Create/edit/lifecycle đúng quyền; không dùng basePrice làm giá tin quảng cáo.'),
('Người thuê/hợp đồng','Tenant list/detail/form, lease/expiry/PDF/lifecycle/settlement','Tách profile khỏi account, review khỏi execute và preview khỏi confirm.'),
('Billing','Invoice list/form/detail/bulk, tariffs, receipt ledger, reports','Snapshot dòng tiền; phân bổ/hoàn; rõ pending/success/unknown.'),
('Bảo trì/vận hành','Maintenance, expenses, action items, P1 checklist','Trạng thái và phân công không vượt scope hiện được hỗ trợ.'),
('Tenant portal','Overview, invoices/proofs/demo, leases, maintenance, inbox','Chỉ dữ liệu của mình; chưa có hồ sơ thuê không hiển thị dữ liệu giả.'),
('Admin workspace','Overview, accounts, inventory, moderation, support/governance','Capability, PII masking, search/pagination, read-only inventory.'),
('Admin registry','Xác minh chủ trọ, duyệt khu, gói và đăng ký dịch vụ','Minh chứng SUPPORT_CLEAN, snapshot/version, hạn mức; gán gói không ghi doanh thu.'),
('Công cụ bổ trợ','Staff scope/simulator, audit, notifications, chatbot','Phân biệt quyền đọc/ghi, provider và giới hạn LANDLORD cho chatbot.')]
table(['Nhóm','Màn hình/luồng','Hành vi trọng tâm'],screen_rows,[1.35,2.7,2.45])
screens=[
 ('ui/workspace-desktop.png','UI-01 Workspace Admin trên desktop','test/features/admin/admin_workspace_flow_test.dart'),
 ('ui/workspace-phone.png','UI-02 Workspace Admin trên màn hình hẹp','test/features/admin/admin_workspace_flow_test.dart'),
 ('ui/accounts-desktop.png','UI-03 Danh sách tài khoản và bộ lọc','test/features/admin/admin_accounts_panel_test.dart'),
 ('ui/account-detail-desktop.png','UI-04 Chi tiết tài khoản và hành động quản trị','test/features/admin/admin_account_detail_panel_test.dart'),
 ('ui/inventory-properties-desktop.png','UI-05 Inventory khu trọ và phân trang','test/features/admin/admin_inventory_panel_test.dart'),
 ('ui/inventory-rooms-desktop.png','UI-06 Inventory phòng có trạng thái tin đăng riêng','test/features/admin/admin_inventory_panel_test.dart'),
 ('ui/inventory-rooms-mobile-200.png','UI-07 Inventory phòng ở chữ lớn trên mobile','test/features/admin/admin_inventory_panel_test.dart'),
 ('ui/inventory-detail-mobile-200.png','UI-08 Chi tiết khu với preview phòng giới hạn','test/features/admin/admin_inventory_panel_test.dart'),
]
for filename,title,source in screens:
    img=TMP/filename
    if not img.exists():continue
    heading(title,4,new_page=True,toc=False)
    picture(img,'Hình '+title+' — dữ liệu fixture, widget render',max_height=6.4)
    para('Nguồn tái lập: '+source+'. Đợt capture mới chạy 51 test đạt; đây không phải ảnh vận hành dữ liệu thật.',small=True)
for filename,title in [('.impeccable/review/tariffs-phone-form.png','UI-09 Form đề xuất biểu phí trên mobile'),('.impeccable/review/tariffs-tablet.png','UI-10 Lịch sử và so sánh biểu phí trên tablet')]:
    img=ROOT/filename
    if not img.exists():continue
    heading(title,4,new_page=True,toc=False);picture(img,'Hình '+title+' — ảnh kiểm thử có sẵn trong workspace',max_height=6.5)
    para('Nguồn ảnh: '+filename+'. Được đối chiếu với code hiện tại, không phải một lần chụp nghiệm thu mới trên thiết bị. Không suy từ ảnh rằng toàn bộ backend/provider đã được xác minh.',small=True)

heading('2.5 Thiết kế triển khai và các điểm cần đối chiếu',2,new_page=True)
heading('2.5.1 Cấu trúc mã nguồn và hợp đồng API',3)
para('lib/main.dart khởi tạo ProviderScope; lib/app/app.dart định nghĩa router và theo dõi SessionState để refresh redirect. Model JSON ở lib/models và API/repository ở lib/data nối với provider của từng feature. Dio quản lý token/refresh trong tầng network/session; UI không nên tự suy lại quyền nghiệp vụ mà máy chủ đã quyết định.')
para('Backend khởi tạo AppModule, configureApp áp prefix, Helmet, request ID, validation và chuẩn hóa lỗi. Controllers khai báo role/permission, DTO kiểm tra dữ liệu và service xác minh scope/ownership/trạng thái. PrismaService quản lý truy cập dữ liệu. Các module vẫn chạy trong cùng một ứng dụng NestJS; transaction tiền nằm ở service và command layer, không ở biểu mẫu Flutter.')
table(['Vùng source','Trách nhiệm','Căn cứ'],[
('lib/app; features/auth; core/network','Điều hướng, phiên, retry và cache theo danh tính','app.dart; auth_navigation.dart; session/cache tests'),
('lib/models; lib/data; features/*','Contract JSON, API/repository, state và widget','Các API/model/provider và test/data, test/features'),
('backend/src/common; auth','Role/capability, property scope, JWT/MFA/session','property-scope.ts; guards; AuthService/AdminMfaService'),
('backend/src/billing; leases','Sổ giao dịch, biểu phí, hóa đơn và quyết toán','BillingLedgerService; BillingCommandsService; SettlementsService'),
('backend/prisma','Model, quan hệ, migration và bootstrap','schema.prisma; migrations; seed/safety documentation'),
('backend/src/platform; support; health','Moderation, finance, reports, SLA và observability','Controller/service và unit tests tương ứng')],[1.9,2.5,2.1])
heading('2.5.2 Tính tiền phân bổ và báo cáo',3)
para('Giá hóa đơn được đóng thành các InvoiceLineItem. Đơn giá hiệu lực, override phòng, tiện ích bậc thang, dòng giảm/phụ thu và nguồn dữ liệu phải được giữ ở thời điểm lập. Với các lệnh tiền mới, BillingCommandsService liên kết requestKey, hash nội dung và kết quả để tránh ghi lặp khi mạng không rõ thành công. Receipt có thể nhận tiền vượt một hóa đơn nhưng chỉ phân bổ đến công nợ hợp lệ; đó khác với Payment trực tiếp vốn bị chặn vượt số dư của hóa đơn.')
para('Báo cáo /reports/monthly không phải báo cáo dòng tiền tổng quát của toàn bộ ledger mới. Billed amount lấy hóa đơn của billingMonth, receivedAmount tổng Payment/PAYMENT+REVERSAL theo paidAt, expenseAmount theo expenseDate; tiền cọc tách riêng. Receipt chưa phân bổ và BillingRefund không được truy vấn trực tiếp trong report này. Vì vậy cần tránh diễn giải nhãn “Đã thu” là mọi tiền đã nhận hoặc “Lợi nhuận ước tính” là lợi nhuận kế toán đã chốt.')
table(['Chỉ tiêu','Cách lấy tại snapshot','Giới hạn diễn giải'],[
('Phải thu theo kỳ','Tổng totalAmount invoice thật, không CANCELLED, đúng billingMonth','Không cùng tập với tiền thu trong tháng.'),
('Đã thu của báo cáo','Tổng Payment.amount theo paidAt, gồm reversal âm, trên invoice thật không hủy','Không bao gồm mọi receipt chưa phân bổ; release có thể khác ngày receipt.'),
('Công nợ kỳ','Tổng max(totalAmount − paidAmount, 0) của invoice kỳ','Số dư hiện tại, không đóng băng tại cuối tháng quá khứ.'),
('Chi phí','Expense theo expenseDate và phạm vi','Ghi khoản chi không xác thực ngân hàng đã thanh toán.'),
('Lợi nhuận ước tính','receivedAmount − expenseAmount','Chênh lệch theo công thức hiện tại, không báo cáo lãi lỗ kế toán.'),
('Tiền cọc','Thu cọc và hoàn/khấu trừ DepositTransaction theo occurredAt','Không cộng cọc vào doanh thu cho thuê.'),
('MRR / occupancy','Lease hiện hành và phòng active qua analytics','Chọn tháng quá khứ không biến occupancy thành snapshot lịch sử.'),
('Collection rate','paidAmount hiện tại của cohort kỳ / tổng phát hành','Không là tiền thu riêng trong tháng; null khi mẫu số 0.'),
('CSV chủ trọ','API trả filename/mimeType/content; Flutter copy Clipboard','Không có lưu XLSX/PDF hoặc lịch báo cáo chủ trọ đã hoàn thành.')],[1.35,2.95,2.2])
heading('2.5.3 Phân quyền dữ liệu và tác vụ nền',3)
para('PropertyScopeInterceptor dùng AsyncLocalStorage để đưa danh tính Staff vào các helper scope. Danh sách controller được phép là whitelist cụ thể; scope rỗng không có nghĩa truy cập mọi khu. TenantScope còn giới hạn hồ sơ có lịch sử thuê/hóa đơn trong scope, do đó một hồ sơ trải nhiều khu có thể bị từ chối. Báo cáo cần mô tả chính sách hạn chế này thay vì hứa rằng Staff luôn xem được mọi chức năng của chủ trọ.')
para('Các tác vụ outbox, invoice automation, nhắc nợ, SLA và Admin report phụ thuộc flag/worker/provider, migration và trạng thái job. API tạo job hoặc hàng đợi không chứng minh job đã được xử lý. Hệ thống có mô hình delivery/heartbeat/incident và các runbook nhưng đợt làm báo cáo không bật worker, gửi tin thực hoặc thử hoàn tiền ngân hàng.')
heading('2.5.4 Các lỗi lệch tài liệu đã được hiệu chỉnh',3)
table(['Nội dung','Bản dở hoặc cách hiểu cũ','Bản đối chiếu mới'],CHANGE_LOG,[1.0,2.4,3.1])

heading('2.6 Kiểm chứng và đánh giá kết quả',2,new_page=True)
heading('2.6.1 Phương pháp và phạm vi kiểm tra',3)
para('Đợt đối chiếu ngày 28/10/2026 đọc trực tiếp controller/service, DTO, router/provider, schema và test. Các kiểm thử dưới đây được chạy trong workspace; kết quả chỉ có giá trị tại snapshot của từng log, không suy là mọi thay đổi sau đó đều đã đạt. Backend unit/build kết thúc khoảng 05:04 và Flutter khoảng 05:07 (UTC+7). Các thay đổi sau mốc test được đối chiếu tĩnh để hoàn thiện UC, không được cộng vào số ca đã kiểm chứng. Không chạy reset-based E2E, migrate, seed, gửi provider hoặc triển khai production.')
table(['Lệnh / phạm vi','Kết quả thực chạy','Diễn giải'],[
('backend: npm test -- --runInBand --silent','65 suite đạt; 1 suite bỏ qua. 672 test đạt, 7 bỏ qua.','Unit/component backend, gồm registry; không là mọi E2E DB.'),
('backend: npm run build','Đạt, exit code 0','TypeScript/NestJS build; không xác minh provider hoặc deployment.'),
('flutter test test/features test/data','503 đạt; 12 lỗi, exit code 1','Không phải toàn bộ test/; 11 lỗi dashboard_chatbot_entry_test.dart; 1 lỗi admin_operations_release_test.dart.'),
('Capture Admin: 4 file widget tests','51 test đạt, ảnh được tạo','Accounts, account detail, inventory và workspace, dùng fixture.'),
('E2E với disposable PostgreSQL','Không chạy trong đợt báo cáo','Có source/runbook; số liệu E2E cũ không được dùng như kết quả mới.'),
('Device/provider/backup restore','Chưa nghiệm thu trong đợt này','Không tuyên bố production-ready, delivered hoặc restore thành công.')],[2.35,2.15,2.0])
heading('2.6.2 Các kịch bản kiểm chứng trọng yếu',3)
para('Các tình huống sau là ma trận thiết kế kiểm thử và nghiệm thu. Một dòng có source test không có nghĩa tất cả lớp integration/provider đã đạt. Cần dùng đúng runner an toàn cho E2E, dữ liệu thử và bộ quyền thực tương ứng.')
table(['Mã / vùng','Tình huống','Kết quả mong đợi'],[(code+'\n'+area,scenario,expected) for code,area,scenario,expected in QUALITY_SCENARIOS],[1.0,2.6,2.9])
heading('2.6.3 Lỗi hiện hữu và giới hạn còn lại',3)
para('Đợt chạy mới ghi nhận 12 lỗi Flutter. Trong đó 11 lỗi thuộc test/features/dashboard/dashboard_chatbot_entry_test.dart, nhóm điều hướng nổi: giữ/rê tab, cuộn tới cuối, viewport 320/375/390/430 với inset 0/34 và callback. Một lỗi thuộc admin_operations_release_test.dart: ca legacy landlord status không tìm được nút Deactivate, nên chưa kiểm chứng được bước gửi lý do và version. Không kết luận nguyên nhân sản phẩm chỉ từ lỗi finder của test. Đợt làm tài liệu giữ nguyên code; xử lý và chạy lại các ca lỗi là công việc kỹ thuật riêng trước nghiệm thu.')
table(['Giới hạn','Tác động','Hướng xử lý tiếp'],[
('Báo cáo chưa bao phủ toàn bộ receipt/refund','Nhãn thu tiền có thể bị đọc sai là mọi dòng tiền','Bổ sung read model tiền mặt/receipt riêng và reconcile với số liệu hiện hữu.'),
('Analytics chưa là snapshot lịch sử/persisted read model','MRR/occupancy hiện tại không đại diện tháng quá khứ','Chốt định nghĩa thời điểm rồi thiết kế snapshot/materialization nếu cần.'),
('Báo cáo định kỳ chủ trọ/XLSX/PDF','API CSV thủ công không đáp ứng mọi dạng artifact','Tách khỏi AdminReportSchedule; bổ sung worker/artifact/retention theo đặc tả mới.'),
('Push/email/Zalo/scanner/AI','Có code nhưng thiếu nghiệm thu provider/device','Cấu hình môi trường thử, chứng cứ delivery/scanning và kiểm tra thu hồi quyền.'),
('Staff chưa hỗ trợ mọi module','Một số thao tác bị deny dù UI/decorator có permission','Giữ deny-by-default; mở từng luồng sau khi kiểm tra scope và test chéo khu.'),
('Danh sách lớn và phần nhập dữ liệu','Một số luồng còn take cố định/client pagination; chưa có import đầu kỳ tổng quát','Đo tải, bổ sung server pagination/import được audit và đối soát.'),
('An toàn phát hành','Unit/build không xác minh migrate/restore/rollback','Theo runbook, backup, rehearsal DB cô lập và owner rollback rõ.'),
('Thông tin hành chính đồ án','Thiếu tên/MSSV/giảng viên và phân công thật','Người dùng cung cấp để hoàn thiện bìa và bảng phân công; không tự hư cấu.')],[1.75,2.35,2.4])

heading('3 Một số thành phần khác của đồ án',1,new_page=True)
heading('3.1 Kế hoạch dự án',2)
para('Mã nguồn thể hiện nhiều mốc phát triển nhưng không đủ xác nhận lịch học phần, thời lượng thực tế hoặc phân công của từng sinh viên. Kế hoạch dưới đây là phương án tham chiếu để hoàn thiện và nghiệm thu hệ thống sau đợt rà soát, không phải bản ghi rằng nhóm đã hoàn thành đúng các tuần này. Ngày bắt đầu, số tuần và người chịu trách nhiệm cần được xác nhận theo kế hoạch học phần.')
phases=[('P1','Chốt phạm vi, tác nhân và UC',1,2,'Đặc tả, định nghĩa tiền/scope và tiêu chí nghiệm thu'),('P2','Ổn định quyền và vòng đời thuê',2,3,'Test chéo khu, settlement/lifecycle/occupants'),('P3','Đối soát hóa đơn và sổ thu',3,4,'Tariff/version, receipt/allocation/refund, reports'),('P4','Sửa lỗi UI và kiểm thử thiết bị',4,6,'Navigation, chữ lớn, tải tệp/PDF và phiên'),('P5','Nghiệm thu tích hợp và vận hành',5,7,'Provider, migration rehearsal, backup/restore, worker'),('P6','Hoàn thiện báo cáo và bàn giao',7,8,'UC/sơ đồ/ảnh, số liệu thật, demo và runbook')]
table(['Gói','Công việc','Tuần tham chiếu','Sản phẩm kiểm tra'],[(k,name,str(s)+'–'+str(e),deliverable) for k,name,s,e,deliverable in phases],[.55,2.1,1.0,2.85])
font_path='/System/Library/Fonts/Supplemental/Times New Roman.ttf'
im=Image.new('RGB',(1800,650),'white');draw=ImageDraw.Draw(im);ft=ImageFont.truetype(font_path,30);fsmall=ImageFont.truetype(font_path,26)
x0,y0=600,65;cw=140;rh=82
for week in range(8):draw.text((x0+week*cw+30,12),'T'+str(week+1),font=ft,fill='black')
for row,(key,name,start,end,_) in enumerate(phases):
    y=y0+row*rh;draw.text((10,y+12),key+' '+name,font=fsmall,fill='black')
    for week in range(8):draw.rectangle((x0+week*cw,y,x0+(week+1)*cw,y+rh),outline='black',width=2)
    bx=x0+(start-1)*cw+8;ex=x0+end*cw-8
    draw.rectangle((bx,y+15,ex,y+rh-15),fill='#EEEEEE',outline='black',width=2)
    for x in range(bx,ex,20):draw.line((x,y+rh-15,min(x+35,ex),y+15),fill='black',width=1)
im.save(TMP/'gantt.png');picture(TMP/'gantt.png','Hình KH-01 Gantt tham chiếu 8 tuần để hoàn thiện và nghiệm thu',max_height=3)
heading('3.2 Đảm bảo thực hiện đúng làm việc nhóm',2)
para('Quy trình phối hợp được đề xuất theo đơn vị nghiệp vụ thay vì chia rời UI và API không có contract. Mỗi công việc cần một người chịu trách nhiệm chính, một người kiểm tra và tiêu chí hoàn thành. Các điểm giao như invoice JSON, scope Staff và settlement fingerprint được chốt trước khi tích hợp. Báo cáo không gán tên người hoặc ghi số buổi họp vì chưa có bằng chứng được cung cấp.')
para('Các thay đổi đi qua bốn bước: xác nhận yêu cầu/UC; sửa một lát cắt có thể kiểm tra; kiểm thử và review quyền/dữ liệu; cập nhật tài liệu và bằng chứng. Khi có nhiều người cùng làm trên workspace, phải giữ thay đổi chưa commit của người khác, không reset hoặc xóa để làm sạch test. Nên dùng nhánh/PR có phạm vi rõ và không đưa secret hay dữ liệu thật vào repository.')
table(['Điểm kiểm soát','Cách áp dụng','Bằng chứng cần lưu'],[('Hợp đồng công việc','UC, tác nhân, đầu vào/đầu ra và ngoại lệ rõ','Issue/checklist và ma trận UC–API–test'),('Review hai chiều','UI contract cùng quy tắc backend/DB','Diff, review comment và test gắn thay đổi'),('Tích hợp','Không vượt scope, không làm mất thay đổi người khác','Nhánh/PR, log CI và kết quả chạy'),('Tài liệu','Sơ đồ đúng snapshot, không lấy số test cũ','Bản báo cáo, source diagrams và audit ngày'),('Rủi ro','Ghi blocker thật, không hạ tiêu chí để báo đạt','Bảng lỗi, người xử lý và điều kiện nghiệm thu')],[1.15,2.7,2.65])
heading('3.3 Các vấn đề về đạo đức và làm việc chuyên nghiệp',2)
para('Trách nhiệm chuyên môn quan trọng nhất của hệ thống là không làm sai sổ tiền và không mở dữ liệu ngoài quyền. Các quyết định như đảo payment thay vì xóa, lưu invoice revision và phân tách maker–checker giúp duy trì căn cứ khi có tranh chấp. Khi có lỗi kiểm thử hoặc thiếu nghiệm thu, người viết phải công khai phạm vi thay vì thay bằng một nhận định “hoàn thiện toàn bộ”.')
para('Về học thuật, hình vẽ và nội dung kế thừa phải có nguồn; tài liệu framework và chuẩn UML dùng để giải thích kỹ thuật, không làm bằng chứng rằng dự án đã đạt chuẩn hoặc hiệu năng tương ứng. Các hình mới được dựng từ schema/service hiện tại, ảnh widget ghi rõ fixture. Báo cáo không tự tạo kết quả khảo sát, phản hồi người dùng, doanh thu, thời gian tiết kiệm hoặc đánh giá giảng viên.')
para('Về xử lý nội dung và hỗ trợ, đánh giá của người thuê cần được kiểm duyệt công bằng, có lý do và khả năng truy vết/kháng nghị theo chức năng hiện hữu. Quyền reveal PII, reset tài khoản hoặc thao tác Admin nhạy cảm phải được giới hạn. Các chính sách pháp lý, lưu trữ dữ liệu và quyền truy cập cần được đơn vị triển khai rà soát riêng trước sử dụng thực tế; báo cáo kỹ thuật này không thay thế tư vấn pháp lý.')
heading('3.4 Tác động xã hội',2)
para('Lợi ích kỳ vọng là giảm dữ liệu rời rạc, tăng khả năng tự tra cứu và minh bạch trạng thái hóa đơn, bảo trì, hợp đồng. Với người thuê, việc biết chứng từ đang chờ duyệt khác với đã thanh toán giúp hạn chế hiểu lầm. Với chủ trọ, cấu trúc receipt/allocation và lịch sử quyết toán giúp giải thích dòng tiền thay vì chỉ nhìn một tổng cộng.')
para('Rủi ro xã hội gồm lộ thông tin cá nhân, đánh giá sai lệch, thông báo quá mức và loại trừ người ít kỹ năng số. Vì vậy cần giao diện tiếng Việt dễ đọc, cách xử lý thay thế khi mất mạng, kênh hỗ trợ rõ và nguyên tắc tối thiểu hóa dữ liệu. Không để khả năng dùng ứng dụng trở thành điều kiện duy nhất để người thuê được giải quyết yêu cầu chính đáng.')
para('Các tác động trên là phân tích thiết kế. Chưa có nghiên cứu triển khai thực địa trong nguồn để định lượng mức tiết kiệm thời gian, tỷ lệ giảm nợ hoặc sự hài lòng. Muốn kết luận cần pilot có đồng thuận, tiêu chí trước/sau, dữ liệu ẩn danh và báo cáo giới hạn mẫu.')
heading('3.5 Kế hoạch cho kiến thức mới và chiến lược học tập',2)
para('Nội dung học tập được ưu tiên từ tính đúng đến khả năng mở rộng. Tài liệu chính thức Flutter/Riverpod [1–2], NestJS [3–4], Prisma/PostgreSQL [5–6], OWASP [7–8] và UML [9] là nguồn tham khảo; khi khác với version dự án phải kiểm tra manifest/lockfile và test, không sao chép cơ chế của phiên bản mới vào source cũ một cách mặc định.')
table(['Chủ đề','Bài tập hoặc hoạt động','Tiêu chí đánh giá'],[('State và session Flutter','Tái hiện đổi tài khoản khi request chưa kết thúc; xử lý refresh','Không lẫn cache và không retry vô hạn'),('Transaction và cạnh tranh ghi','Tái hiện phân bổ/hoàn/quyết toán đồng thời trên DB test','Không âm số dư, không ghi trùng, rollback đúng'),('Quyền và privacy','Lập ma trận role/permission/property scope và test chéo khu','Từ chối đúng ở server, không chỉ ẩn nút'),('UML và truy vết','Đọc schema, vẽ bội số đúng và sequence theo transaction','UC–figure–service–test nhất quán'),('Tích hợp an toàn','Provider sandbox, file access/scanner và demo gateway','Không dùng tiền/thông tin thật; trạng thái delivery đúng'),('Vận hành','Rehearsal migrate, backup/restore và worker recovery','Bằng chứng tái lập trên môi trường cô lập'),('Viết báo cáo','Phân biệt dữ liệu đo, mục tiêu và giả định','Không hư cấu số liệu hoặc danh tính; nguồn minh bạch')],[1.35,2.95,2.2])

heading('4 Kết luận',1,new_page=True)
para('Đồ án đã có một cấu trúc ứng dụng xuyên suốt từ tìm phòng đến vận hành thuê và quản trị nền tảng trong mã nguồn hiện tại. Các điểm phát triển quan trọng là mô hình người dùng theo vai trò, scope Staff, lịch sử giá và hóa đơn, receipt/allocation/refund, quyết toán trả phòng, lifecycle execution và các công cụ Admin có capability. Báo cáo làm lại giữ đúng bố cục mẫu V1, thay các tài liệu con dán rời bằng hệ thống chương/UC/sơ đồ có truy vết.')
para('Về bằng chứng kiểm chứng của lần rà soát này, backend unit suite có 672 test đạt và 7 test bỏ qua; build NestJS đạt. Bộ Flutter features và data có 503 test đạt và 12 lỗi ở điều hướng dashboard và ca trạng thái chủ trọ Admin; ảnh capture Admin đến từ 51 widget test đạt với fixture. Các kết quả này hỗ trợ đánh giá chất lượng source nhưng không chứng minh mọi E2E, thiết bị, provider, worker hoặc production đã hoàn tất.')
para('Mức hoàn thiện thực tế còn phụ thuộc việc xử lý lỗi UI, định nghĩa/read model báo cáo bao phủ ledger mới, lịch báo cáo chủ trọ, tích hợp delivery/push/scanner, kiểm chứng migration–restore–rollback và nghiệm thu thiết bị. Thông tin nhóm, MSSV và giảng viên chưa có trong nguồn nên không được tự điền. Đây là các phần phải hoàn tất minh bạch thay vì đưa vào kết luận như thành quả đã đạt.')
para('Hướng phát triển tiếp theo nên ưu tiên đóng rủi ro dữ liệu và kiểm chứng hơn là thêm màn hình: hoàn thiện báo cáo đối chiếu tiền nhận/phân bổ/hoàn, chốt snapshot lịch sử, xử lý quyền chéo khu, sửa navigation và tái chạy bộ kiểm thử. Sau đó mới mở rộng provider và tối ưu ở quy mô đo được. Giá trị của hệ thống được đánh giá bằng khả năng giải thích đúng dữ liệu và tái lập bằng chứng, không chỉ số lượng chức năng.')

heading('5 Tài liệu tham khảo',1,new_page=True)
references=json.loads((TMP/'references.json').read_text())
ref_names=['Flutter — Guide to app architecture','Riverpod — Getting started','NestJS — Authorization','NestJS — Validation','Prisma ORM — Transactions','PostgreSQL 16 — Transaction Isolation','OWASP — Authorization Cheat Sheet','OWASP — File Upload Cheat Sheet','OMG — Unified Modeling Language 2.5.1']
for i,(name,ref) in enumerate(zip(ref_names,references),1):para(f'[{i}] {name}. {ref["url"]}. Truy cập ngày 28/10/2026.',small=True)
local_refs=[('Mã nguồn Flutter','lib/app, lib/core, lib/data, lib/features, lib/models và test/.'),('Mã nguồn backend','backend/src/, backend/test/, backend/prisma/schema.prisma và migrations/.'),('Tài liệu mẫu','docs/CSE703014_Report_Template_VI_V1.docx; chỉ dùng làm chuẩn cấu trúc/hình thức.'),('Bản báo cáo đang làm dở','docs/Bao_cao_da_chinh_form.docx; dùng để đối chiếu chênh lệch, không giữ số liệu/sơ đồ lỗi thời.'),('Tình trạng chức năng','docs/app-functional-delivery-status.md; docs/admin-functional-delivery-status.md; BILLING_DELIVERY_STATUS.md.'),('Tài khoản và inventory Admin','docs/admin-account-management.md; docs/admin-inventory-management.md; docs/admin-permission-matrix.md.'),('Báo cáo kiểm chứng mới','Log backend unit/build, Flutter features+data và widget capture ngày 28/10/2026; đối chiếu trong tài nguyên kèm source báo cáo.')]
for i,(title,description) in enumerate(local_refs,10):para(f'[{i}] {title}. {description} Snapshot kiểm tra ngày 28/10/2026.',small=True)

heading('Phụ lục A Đối chiếu nội dung cũ và mới',1,new_page=True)
para('Phụ lục này giải thích các thay đổi trọng yếu để tránh xem bản báo cáo mới như chỉ đổi định dạng. Bản dở có 67 UC, 133 bảng và 45 ảnh/media; mục lục cũ dùng số trang không phản ánh bản hiện tại. Các sơ đồ mới được dựng lại thay cho ảnh cũ, có source để sửa khi code đổi.')
table(['Vùng','Nội dung đối chiếu','Cách xử lý trong báo cáo'],CHANGE_LOG,[1.0,2.5,3.0])
heading('A.1 Truy vết 67 UC của bản dở sang báo cáo mới',2)
para('Mã UC đã được tổ chức lại theo mục tiêu và miền nghiệp vụ. Bảng sau giúp đối chiếu từng UC cũ với đặc tả hiện tại; một UC mới có thể kế thừa nhiều mục cũ hoặc phân rã mục cũ thành nhiều luồng. Các UC bổ sung ngoài bảng được liệt kê đầy đủ ở mục 2.4.1.')
crosswalk=json.loads((TMP/'uc-crosswalk.json').read_text())
table(['Mã cũ','Tên mục tiêu trong bản dở','Đặc tả hiện tại'],[(row['old_id'],row['old_title'],row['new_refs']) for row in crosswalk],[1.15,3.5,1.85],size=10.5)

heading('Phụ lục B Truy vết UC và căn cứ kiểm chứng',1,new_page=True)
para('Mỗi dòng nối UC với source nghiệp vụ, vùng Flutter và test liên quan. Test được liệt kê là nơi kiểm tra theo module, không tuyên bố phủ mọi nhánh của UC. Một số chức năng có backend nhưng UI/acceptance chưa đồng đều; giới hạn đã nêu trong đặc tả và mục 2.6. Các đường dẫn là đường dẫn tương đối repository để có thể dùng lại trên máy khác.')
rows=[]
for case in CASES:
    test=case['tests'][0] if case['tests'] else 'Chưa tìm thấy test riêng bằng kiểm kê tên; cần kiểm tra module.'
    ui=', '.join('lib/features/'+m+'/' for m in GROUPS[case['group']][2].split(','))
    rows.append((case['id']+'\n'+case['title'],case['source'],ui+'\n'+test))
table(['UC','Nguồn nghiệp vụ','UI và test liên quan'],rows,[1.9,2.15,2.45],size=10)

heading('Phụ lục C Bằng chứng và tài nguyên chỉnh sửa',1,new_page=True)
para('Bản Word này là kết quả đọc lại template và source, không ghi đè hai file gốc. Các sơ đồ có PNG/SVG và source DOT hoặc dữ liệu thông điệp; danh mục UC, inventory kỹ thuật lưu riêng, hash template/source và log test được lưu trong thư mục tài nguyên của báo cáo. Khi dự án thay đổi, cần chạy lại kiểm kê và kiểm thử, không chỉ sửa số trang trong Word.')
table(['Tài nguyên','Ý nghĩa'],[
('docs/report_2026-10-28/','Nguồn dựng, danh mục, sơ đồ và log kiểm chứng của bản này.'),('audit.json',''+str(len(CASES))+' UC và căn cứ source, được lưu riêng để đối chiếu.'),('diagram-manifest.json; diagrams/','Danh mục hình cùng PNG/SVG/DOT; sequence được dựng từ thông điệp trong report_data.py.'),('artifact.md','Quy tắc giữ mẫu V1, nguồn SHA-256 và các thay đổi hình thức có chủ đích.'),('backend-tests-latest.log; backend-build-latest.log','Kết quả kiểm thử unit backend và build vừa chạy.'),('flutter-tests-latest.log; flutter-failures.json','Kết quả Flutter features/data và 12 ca lỗi quan sát được.'),('ui-capture-tests.log; ui/','51 test capture đạt và ảnh fixture Admin.'),('verification.json','Kết quả kiểm tra cấu trúc DOCX, hình, bảng trắng, mã UC và tham chiếu sau dựng.')],[2.2,4.3])
para('Thông tin còn cần người dùng xác nhận trước khi nộp: mã nhóm, họ tên/MSSV/lớp/email thành viên, giảng viên hướng dẫn và phân công thực tế. Không có chỗ nào trong báo cáo tự gán các thông tin này dựa vào tên thư mục hoặc tác giả trong metadata.')

# A stable, linked TOC is filled with page numbers measured from the rendered PDF.
anchor=toc_anchor._p
for name,title,level in TOC:
    p=OxmlElement('w:p');ppr=OxmlElement('w:pPr')
    ind=OxmlElement('w:ind');ind.set(qn('w:left'),str((level-1)*240));ppr.append(ind)
    spacing=OxmlElement('w:spacing');spacing.set(qn('w:after'),'60');spacing.set(qn('w:line'),'240');spacing.set(qn('w:lineRule'),'auto');ppr.append(spacing)
    tabs=OxmlElement('w:tabs');tab=OxmlElement('w:tab');tab.set(qn('w:val'),'right');tab.set(qn('w:leader'),'dot');tab.set(qn('w:pos'),'9360');tabs.append(tab);ppr.append(tabs)
    p.append(ppr)
    link=OxmlElement('w:hyperlink');link.set(qn('w:anchor'),name)
    r=OxmlElement('w:r');rpr=OxmlElement('w:rPr');sz=OxmlElement('w:sz');sz.set(qn('w:val'),'24');rpr.append(sz)
    rf=OxmlElement('w:rFonts');rf.set(qn('w:ascii'),'Times New Roman');rf.set(qn('w:hAnsi'),'Times New Roman');rpr.append(rf)
    if level==1:rpr.append(OxmlElement('w:b'))
    r.append(rpr);text=OxmlElement('w:t');text.text=title;r.append(text);link.append(r)
    r=OxmlElement('w:r');r.append(OxmlElement('w:tab'));link.append(r)
    r=OxmlElement('w:r');text=OxmlElement('w:t');text.text=str(TOC_PAGES.get(title,'—'));r.append(text);link.append(r);p.append(link)
    anchor.addnext(p);anchor=p

settings=DOC.settings.element
for item in list(settings.findall(qn('w:updateFields'))):settings.remove(item)
update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');settings.append(update)
DOC.core_properties.title='Hệ thống quản lý nhà trọ'
DOC.core_properties.subject='Báo cáo đồ án liên ngành theo mẫu V1 — đối chiếu mã nguồn ngày 28/10/2026'
DOC.core_properties.author='';DOC.core_properties.last_modified_by=''
DOC.core_properties.keywords='Rental Room Management; UC; UML; Flutter; NestJS; báo cáo đồ án'
DOC.save(FINAL)
(TMP/'toc-entries.json').write_text(json.dumps(TOC,ensure_ascii=False,indent=2))
(TMP/'bookmarks.json').write_text(json.dumps(BOOKMARKS,ensure_ascii=False,indent=2))
print('Saved',FINAL)
print('UC:',len(CASES),'tables:',len(DOC.tables),'images:',len(DOC.inline_shapes),'TOC entries:',len(TOC))

support=ROOT/'docs/report_2026-10-28';support.mkdir(exist_ok=True)
for filename in ['report_data.py','narrative.py','audit.py','diagrams.py','render_diagrams.cjs','build_report.py','audit.json','diagram-manifest.json','artifact.md','references.json','flutter-failures.json','backend-tests.log','backend-build.log','flutter-tests.log','backend-tests-latest.log','backend-build-latest.log','flutter-tests-latest.log','ui-capture-tests.log','toc-entries.json','toc-pages.json','uc-crosswalk.json','phenikaa-logo.png']:
    source=TMP/filename
    if source.exists():shutil.copy2(source,support/filename)
for subdir in ['diagrams','ui']:
    if (TMP/subdir).exists():shutil.copytree(TMP/subdir,support/subdir,dirs_exist_ok=True)
shutil.copy2(TMP/'gantt.png',support/'gantt.png')
