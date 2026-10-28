from pathlib import Path
from collections import Counter
from zipfile import ZipFile
import hashlib
import json
import re
import sys
from docx import Document
from lxml import etree
from measure_pages import measured
from ba_specs import specification

BASE = Path(__file__).resolve().parent
ROOT = Path('/Users/hoang/Rental_room_management')
FINAL = ROOT / 'docs/Bao_cao_Quan_ly_Nha_tro_V2_2026-10-28.docx'
SUPPORT = ROOT / 'docs/report_2026-10-28_v2'
render = Path(sys.argv[1])
data = json.loads((BASE / 'audit.json').read_text())
diagrams = json.loads((BASE / 'diagram-manifest.json').read_text())
doc = Document(FINAL)
with ZipFile(FINAL) as archive:
    xml = etree.fromstring(archive.read('word/document.xml'))
ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
all_text = '\n'.join(xml.xpath('//w:t/text()', namespaces=ns))
spec_tables = []
for index, table in enumerate(doc.tables):
    if len(table.columns) == 2 and len(table.rows) > 1 and table.cell(1, 0).text == 'Mã và tên':
        match = re.match(r'UC-[A-Z]+-\d+(?:\.\d+)?', table.cell(1, 1).text)
        if match:
            spec_tables.append((index, match.group()))
ids = [key for index, key in spec_tables]
toc, figures, main, texts = measured(render)
checks = {
    '131_unique_in_scope_specifications': len(ids) == 131 and len(set(ids)) == 131
        and set(ids) == {case['id'] for case in data['cases']},
    'BA_metadata_mainflow_outcomes_complete': all(
        len(doc.tables[index].rows) == 9 and len(doc.tables[index + 1].rows) >= 7
        and len(doc.tables[index + 2].rows) == 9
        and all(cell.text.strip() for table in doc.tables[index:index + 3]
                for row in table.rows for cell in row.cells) for index, key in spec_tables),
    'four_human_actors': {actor for case in data['cases'] for actor in case['diagram_actors']}
        == {'Khách vãng lai', 'Người thuê', 'Chủ trọ', 'Admin'},
    'no_staff_in_report': not re.search(r'\bstaff\b|UC-STF|FR-STF', all_text, re.I),
    'all_BA_cases_have_detailed_steps': all(len(specification(case)['steps']) >= 6 for case in data['cases']),
    'source_evidence_hashes': all(hashlib.sha256((SUPPORT / 'evidence' / case['source']).read_bytes()).hexdigest()
        == case['source_hash'] for case in data['cases']),
    '83_diagram_panels': len(diagrams) == 83,
    'all_used_diagrams_available': all((SUPPORT / 'diagrams' / (entry['id'] + extension)).is_file()
        for entry in diagrams for extension in ['.svg', '.png']),
    '24_stick_actor_UC_diagrams': sum(entry['kind'] == 'usecase' for entry in diagrams) == 24
        and all('<circle' in (BASE / 'diagrams' / (entry['id'] + '.svg')).read_text()
                for entry in diagrams if entry['kind'] == 'usecase'),
    '38_yellow_sequence_panels': sum(entry['kind'] == 'sequence' for entry in diagrams) == 38
        and all('#FFFFB9' in (BASE / 'diagrams' / (entry['id'] + '.svg')).read_text()
                for entry in diagrams if entry['kind'] == 'sequence'),
    'all_tables_white': all(node.get('{'+ns['w']+'}fill') in ('FFFFFF', 'auto', None)
        for node in xml.xpath('//w:tbl//w:shd', namespaces=ns)),
    'all_table_borders_black': all(node.get('{'+ns['w']+'}color') == '000000'
        for node in xml.xpath('//w:tblBorders/*', namespaces=ns)),
    'all_table_headers_repeat': all(table.rows[0]._tr.xpath('./w:trPr/w:tblHeader') for table in doc.tables),
    'three_sections_A4': len(doc.sections) == 3 and all(abs(section.page_height.cm - 29.7) < .03
        and abs(section.page_width.cm - 21) < .03 for section in doc.sections),
    'academic_margins': all(abs(section.left_margin.cm - 3) < .03 and abs(section.right_margin.cm - 2) < .03
        for section in doc.sections),
    'cover_double_border_only': len(xml.xpath('//w:sectPr/w:pgBorders', namespaces=ns)) == 1
        and all(node.get('{'+ns['w']+'}val') == 'double' for node in xml.xpath('//w:pgBorders/*', namespaces=ns)),
    'all_render_pages_present': len(list(render.glob('page-*.png'))) == len(texts),
    'no_empty_pages': all(len(text.strip()) > 25 for text in texts),
    'toc_measured_pages_match': toc == json.loads((BASE / 'toc-pages.json').read_text()),
    'figure_list_measured_pages_match': figures == json.loads((BASE / 'figure-pages.json').read_text()),
    '114_numbered_figures': len(figures) == 114 and len(doc.inline_shapes) == 114,
    '28_unique_UI_screens': len(set(re.findall(r'UI-\d+', all_text))) == 28,
    'no_raw_schema_appendix': not any(title in all_text for title in ['Danh mục mô hình dữ liệu', 'Danh mục endpoint kỹ thuật']),
    '67_legacy_cases_crosswalk': len(json.loads((BASE / 'uc-crosswalk.json').read_text())) == 67,
}
anchors = {node.get('{'+ns['w']+'}name') for node in xml.xpath('//w:bookmarkStart', namespaces=ns)}
checks['navigation_anchors_resolve'] = all(node.get('{'+ns['w']+'}anchor') in anchors
    for node in xml.xpath('//w:hyperlink[@w:anchor]', namespaces=ns))
for filename in ['CSE703014_Report_Template_VI_V1.docx', 'Bao_cao_da_chinh_form.docx']:
    path = ROOT / 'docs' / filename
    inventory = json.loads((BASE / (path.stem + '-parts.json')).read_text())
    with ZipFile(path) as archive:
        current = {part.filename: hashlib.sha256(archive.read(part.filename)).hexdigest() for part in archive.infolist()}
    checks['original_unchanged_' + path.stem] = current == {part['part']: part['sha256'] for part in inventory}
for filename, expected in [('ui-capture-v2.log', '+11: All tests passed!'),
                           ('ui-tenant-capture-v2.log', '+27: All tests passed!'),
                           ('ui-landlord-capture-v2.log', '+5: All tests passed!')]:
    checks['capture_' + filename] = expected in (BASE / filename).read_text()
report = {'date': '2026-10-28', 'output': str(FINAL), 'sha256': hashlib.sha256(FINAL.read_bytes()).hexdigest(),
          'pages': len(texts), 'body_starts_physical_page': main + 1, 'use_cases': len(ids),
          'diagrams': dict(Counter(entry['kind'] for entry in diagrams)), 'figures': len(figures),
          'tables': len(doc.tables), 'checks': checks, 'all_checks_passed': all(checks.values()),
          'limitations': ['Administrative identities and actual assignments were not provided.',
                         'Widget fixtures are not device, provider, E2E or production acceptance.',
                         'Earlier product tests had 12 Flutter failures; later source updates were checked statically.']}
for folder in [BASE, SUPPORT]:
    (folder / 'verification.json').write_text(json.dumps(report, ensure_ascii=False, indent=2))
print(json.dumps(report, ensure_ascii=False, indent=2))
if not report['all_checks_passed']:
    sys.exit(1)
