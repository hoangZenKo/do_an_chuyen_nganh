import copy
import hashlib
import json
import math
import re
from pathlib import Path

import numpy as np
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.text.paragraph import Paragraph
from lxml import etree
from PIL import Image

ROOT = Path('/Users/hoang/Rental_room_management')
BASE = ROOT / 'output/report_submission_2026-10-30'
DIAGRAMS = ROOT / 'output/sequence_ba_report_2026-10-30'
SOURCE = next(Path('/Users/hoang/Desktop').glob('*/Report_finnal.docx'))
OUTPUT = BASE / 'Report_final_BA_Da_chinh_form_2026-10-30.docx'
CROPS = BASE / 'diagram_pages'
NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
FONT = 'Times New Roman'
PORTRAIT_WIDTH = Cm(16)
LANDSCAPE_WIDTH = Cm(24.7)
manifest = {'input': str(SOURCE), 'input_sha256': hashlib.sha256(SOURCE.read_bytes()).hexdigest(), 'diagrams': [], 'toc': [], 'figures': []}


def text_of(element):
    return ''.join(element.xpath('.//w:t/text()'))


def child(parent, name):
    found = parent.find(qn(name))
    if found is None:
        found = OxmlElement(name)
        parent.append(found)
    return found


def set_font(run, size=None, bold=None):
    run.font.name = FONT
    fonts = child(run._element.get_or_add_rPr(), 'w:rFonts')
    for key in ['ascii', 'hAnsi', 'eastAsia', 'cs']:
        fonts.set(qn('w:' + key), FONT)
    for key in list(fonts.attrib):
        if 'Theme' in key:
            del fonts.attrib[key]
    run.font.color.rgb = RGBColor(0, 0, 0)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def paragraph_before(anchor, text='', style=None):
    element = OxmlElement('w:p')
    anchor.addprevious(element)
    paragraph = Paragraph(element, doc._body)
    if style:
        paragraph.style = style
    if text:
        paragraph.add_run(text)
    return paragraph


def replace_text(paragraph, text):
    paragraph.clear()
    paragraph.add_run(text)


def bookmark(paragraph, name):
    global bookmark_id
    bookmark_id += 1
    start = OxmlElement('w:bookmarkStart')
    start.set(qn('w:id'), str(bookmark_id))
    start.set(qn('w:name'), name)
    end = OxmlElement('w:bookmarkEnd')
    end.set(qn('w:id'), str(bookmark_id))
    paragraph._p.insert(1 if paragraph._p.find(qn('w:pPr')) is not None else 0, start)
    paragraph._p.append(end)


def format_paragraph(paragraph, size=13, align=WD_ALIGN_PARAGRAPH.JUSTIFY, first=0.75, line=1.3, before=0, after=6):
    paragraph.alignment = align
    formatting = paragraph.paragraph_format
    formatting.first_line_indent = Cm(first)
    formatting.left_indent = Cm(0)
    formatting.right_indent = Cm(0)
    formatting.line_spacing = line
    formatting.space_before = Pt(before)
    formatting.space_after = Pt(after)
    formatting.widow_control = True
    for run in paragraph.runs:
        set_font(run, size)


def field_run(paragraph, kind, instruction=None):
    run = paragraph.add_run()
    set_font(run, 1)
    if kind == 'instruction':
        element = OxmlElement('w:instrText')
        element.set(qn('xml:space'), 'preserve')
        element.text = instruction
    else:
        element = OxmlElement('w:fldChar')
        element.set(qn('w:fldCharType'), kind)
    run._r.append(element)


def toc_entry(anchor, entry, style):
    paragraph = paragraph_before(anchor, style=style)
    paragraph.paragraph_format.tab_stops.add_tab_stop(PORTRAIT_WIDTH, WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    link = OxmlElement('w:hyperlink')
    link.set(qn('w:anchor'), entry['anchor'])
    link.set(qn('w:history'), '1')
    run = OxmlElement('w:r')
    run_properties = OxmlElement('w:rPr')
    fonts = OxmlElement('w:rFonts')
    fonts.set(qn('w:ascii'), FONT)
    fonts.set(qn('w:hAnsi'), FONT)
    run_properties.append(fonts)
    run.append(run_properties)
    text = OxmlElement('w:t')
    text.text = entry['title']
    run.append(text)
    link.append(run)
    paragraph._p.append(link)
    paragraph.add_run('\t')
    page_run = paragraph.add_run('1')
    page_run._r.set(qn('w:rsidR'), '00000001')
    entry['page'] = '1'
    return paragraph


def configure_style(name, size, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.LEFT, before=0, after=6, line=1.15, first=0, keep=False):
    style = doc.styles[name] if name in doc.styles else doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    style.font.name = FONT
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = italic
    style.font.color.rgb = RGBColor(0, 0, 0)
    fonts = child(style.element.get_or_add_rPr(), 'w:rFonts')
    for key in ['ascii', 'hAnsi', 'eastAsia', 'cs']:
        fonts.set(qn('w:' + key), FONT)
    for key in list(fonts.attrib):
        if 'Theme' in key:
            del fonts.attrib[key]
    formatting = style.paragraph_format
    formatting.alignment = align
    formatting.first_line_indent = Cm(first)
    formatting.space_before = Pt(before)
    formatting.space_after = Pt(after)
    formatting.line_spacing = line
    formatting.keep_with_next = keep
    formatting.widow_control = True
    return style


def image_slices(key):
    path = DIAGRAMS / 'PNG' / f'Sequence_Diagram_{key}.png'
    image = Image.open(path).convert('RGB')
    width, height = image.size
    svg = etree.parse(str(DIAGRAMS / 'source' / f'{key}.svg'))
    ratio = width / float(svg.getroot().get('width').removesuffix('px'))
    headers = [float(node.get('y', '0')) for node in svg.xpath('//*[local-name()="g" and contains(@class,"participant-head")]//*[local-name()="text"]')]
    assert headers, key
    header_end = math.ceil((max(headers) + 22) * ratio)
    points_per_pixel = LANDSCAPE_WIDTH.pt / width
    first_max = int(330 / points_per_pixel)
    following_max = int(405 / points_per_pixel) - header_end - 12
    pixels = np.asarray(image)
    dark = np.sum(np.min(pixels, axis=2) < 105, axis=1)
    text_rows = np.zeros(height, dtype=bool)
    for node in svg.xpath('//*[local-name()="text"]'):
        if not (node.text or '').strip():
            continue
        baseline = float(node.get('y', '0'))
        font_size = float(node.get('font-size', '42.5'))
        top = max(0, math.floor((baseline - font_size * 0.95) * ratio))
        bottom = min(height, math.ceil((baseline + font_size * 0.22) * ratio))
        text_rows[top:bottom] = True
    cuts = [0]
    position = 0
    while position < height:
        capacity = first_max if position == 0 else following_max
        target = min(height, position + capacity)
        if target < height:
            lower = max(position + capacity // 2, target - 230)
            candidates = []
            for row in range(lower + 4, target - 4):
                if max(dark[row - 4:row + 5]) < 30 and not text_rows[row - 4:row + 5].any():
                    candidates.append(row)
            if not candidates:
                raise ValueError(f'No safe cut for {key} at {target}')
            target = candidates[-1]
        cuts.append(target)
        position = target
    if len(cuts) > 2 and (cuts[-1] - cuts[-2]) * points_per_pixel < 35:
        del cuts[-2]
    paths = []
    for index, (top, bottom) in enumerate(zip(cuts, cuts[1:])):
        if index == 0:
            crop = image.crop((0, top, width, bottom))
        else:
            crop = Image.new('RGB', (width, header_end + 12 + bottom - top), 'white')
            crop.paste(image.crop((0, 0, width, header_end)), (0, 0))
            crop.paste(image.crop((0, top, width, bottom)), (0, header_end + 12))
        destination = CROPS / f'{key}_phan_{index + 1:02}.png'
        crop.save(destination, dpi=(300, 300))
        paths.append(destination)
    manifest['diagrams'].append({'id': key, 'source_sha256': hashlib.sha256(path.read_bytes()).hexdigest(), 'parts': len(paths), 'cuts': cuts, 'pixels': [width, height], 'repeated_header_pixels': header_end})
    return paths


BASE.mkdir(exist_ok=True)
CROPS.mkdir(exist_ok=True)
doc = Document(SOURCE)
body = doc._element.body
bookmark_id = max(int(value) for value in body.xpath('.//w:bookmarkStart/@w:id')) + 100
original_texts = [text_of(element) for element in body if element.tag == qn('w:tbl')]
original_table_count = len(doc.tables)

configure_style('Normal', 13, align=WD_ALIGN_PARAGRAPH.JUSTIFY, line=1.3, first=0.75)
for level, size in [(1, 16), (2, 14), (3, 13), (4, 13), (5, 12)]:
    configure_style(f'Heading {level}', size, bold=True, before=12 if level < 3 else 9, after=6, keep=True)
configure_style('BAHeading4', 13, bold=True, before=10, after=6, keep=True)
configure_style('Report Figure Caption', 11, align=WD_ALIGN_PARAGRAPH.CENTER, before=4, after=6)
configure_style('Report Figure Continuation', 11, align=WD_ALIGN_PARAGRAPH.CENTER, before=4, after=6)
configure_style('Report Sequence Note', 11, align=WD_ALIGN_PARAGRAPH.LEFT, line=1.12, after=4)
configure_style('Report Source Note', 10, align=WD_ALIGN_PARAGRAPH.LEFT, line=1.1, after=6)
configure_style('Report Contents Title', 16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, before=0, after=15)
for level in range(1, 4):
    style = configure_style(f'TOC {level}', 12 if level == 1 else 11.5, bold=level == 1, line=1.05, after=1.5)
    style.paragraph_format.left_indent = Cm((level - 1) * 0.5)
configure_style('Report Figure List', 11, line=1.1, after=3)

for section in doc.sections:
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(2.5), Cm(2.5)
    section.left_margin, section.right_margin = Cm(3), Cm(2)
    for footer in [section.footer, section.first_page_footer]:
        for paragraph in footer.paragraphs:
            format_paragraph(paragraph, 11, WD_ALIGN_PARAGRAPH.CENTER, first=0, line=1, after=0)
            for run in paragraph.runs:
                run.bold, run.italic = False, False

main_started = False
for paragraph in doc.paragraphs:
    text = paragraph.text.strip()
    if text == 'Phân công thực hiện đồ án':
        main_started = True
    if not main_started:
        continue
    style = paragraph.style.style_id
    if paragraph._p.xpath('.//w:drawing'):
        format_paragraph(paragraph, 1, WD_ALIGN_PARAGRAPH.CENTER, first=0, line=1, before=3, after=3)
        paragraph.paragraph_format.keep_with_next = True
        continue
    if style.startswith('Heading') or style == 'BAHeading4':
        for run in paragraph.runs:
            set_font(run, None, True)
            run.font.size = None
            run.italic = False
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.first_line_indent = Cm(0)
        paragraph.paragraph_format.keep_with_next = True
        paragraph.paragraph_format.keep_together = True
        if style == 'Heading1' and re.match(r'^[1-5] ', text):
            paragraph.paragraph_format.page_break_before = True
    elif re.match(r'^Hình\s+\d+\.\d+\b', text):
        paragraph.style = doc.styles['Report Figure Caption']
        format_paragraph(paragraph, 11, WD_ALIGN_PARAGRAPH.CENTER, first=0, line=1.1, before=4)
        paragraph.paragraph_format.keep_with_next = False
        paragraph.paragraph_format.keep_together = True
    elif paragraph._p.find(qn('w:pPr')) is not None and paragraph._p.find(qn('w:pPr')).find(qn('w:sectPr')) is not None:
        format_paragraph(paragraph, 1, first=0, line=1, after=0)
    elif text:
        format_paragraph(paragraph)

for table in doc.tables:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [int(column.width or 1) for column in table.columns]
    total = sum(widths)
    for column, width in zip(table.columns, widths):
        column.width = round(int(PORTRAIT_WIDTH) * width / total)
    properties = table._tbl.tblPr
    table_width = child(properties, 'w:tblW')
    table_width.set(qn('w:w'), str(round(PORTRAIT_WIDTH.twips)))
    table_width.set(qn('w:type'), 'dxa')
    indent = child(properties, 'w:tblInd')
    indent.set(qn('w:w'), '0')
    indent.set(qn('w:type'), 'dxa')
    borders = child(properties, 'w:tblBorders')
    for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        border = child(borders, 'w:' + edge)
        for attribute, value in [('val', 'single'), ('sz', '4'), ('color', '000000')]:
            border.set(qn('w:' + attribute), value)
    margins = child(properties, 'w:tblCellMar')
    for edge, value in [('top', '60'), ('bottom', '60'), ('left', '80'), ('right', '80')]:
        margin = child(margins, 'w:' + edge)
        margin.set(qn('w:w'), value)
        margin.set(qn('w:type'), 'dxa')
    seen_cells = set()
    for row_index, row in enumerate(table.rows):
        row_properties = row._tr.get_or_add_trPr()
        for height in list(row_properties.findall(qn('w:trHeight'))):
            row_properties.remove(height)
        child(row_properties, 'w:cantSplit')
        if row_index == 0:
            child(row_properties, 'w:tblHeader')
        for cell in row.cells:
            if cell._tc in seen_cells:
                continue
            seen_cells.add(cell._tc)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            if row_index == 0:
                child(cell._tc.get_or_add_tcPr(), 'w:shd').set(qn('w:fill'), 'F2F2F2')
            for paragraph in cell.paragraphs:
                format_paragraph(paragraph, 11.5, WD_ALIGN_PARAGRAPH.LEFT, first=0, line=1.08, after=2)
                paragraph.paragraph_format.keep_with_next = False
                paragraph.paragraph_format.keep_together = False
                if row_index == 0:
                    for run in paragraph.runs:
                        run.bold = True

for paragraph in doc.paragraphs[:11]:
    for instruction in list(paragraph._p.xpath('.//w:instrText')):
        if 'INCLUDEPICTURE' in (instruction.text or ''):
            instruction.getparent().remove(instruction)
            for field in list(paragraph._p.xpath('.//w:fldChar')):
                field.getparent().remove(field)
    for run in paragraph.runs:
        set_font(run)

for index, paragraph in enumerate(doc.paragraphs[:-1]):
    following = doc.paragraphs[index + 1]
    if paragraph.style.style_id in ['Heading1', 'Heading2', 'Heading3'] and following.style.name == 'BA Heading 4':
        following.paragraph_format.page_break_before = False

sequence_audit = json.loads((DIAGRAMS / 'source/verification.json').read_text())
check_lines = (DIAGRAMS / 'DOI_CHIEU_KIEM_TRA.md').read_text().splitlines()
outcomes = {}
for line in check_lines:
    if re.match(r'^\| SD-\d+', line):
        columns = [value.strip() for value in line.strip('|').split('|')]
        outcomes[columns[0]] = columns[2]
preconditions = [
    'Tài khoản đã đăng ký; người dùng mở màn hình đăng nhập hoặc ứng dụng đang khôi phục phiên sử dụng.',
    'Người tìm thuê đã đăng nhập; tin đăng còn nhận khách; chủ trọ có quyền xử lý lịch và nhận thuê.',
    'Người thực hiện có quyền lập hợp đồng; đã chọn phòng, hồ sơ người thuê và nhập điều khoản thuê.',
    'Chủ trọ hoặc nhân viên có quyền đề xuất; người xét duyệt được cấp quyền duyệt biểu phí tương ứng.',
    'Có hợp đồng thuộc phạm vi quản lý; đã chọn kỳ lập hóa đơn và có chỉ số điện nước cần thiết.',
    'Có quyền ghi thu, người thuê và các hóa đơn thuộc cùng phạm vi; đã nhập số tiền thực nhận.',
    'Có quyền xử lý khoản thu; phân bổ cần gỡ còn hiệu lực hoặc khoản thu còn tiền chưa phân bổ.',
    'Người thuê có hóa đơn được phép gửi chứng từ; người xét duyệt có quyền quản lý hóa đơn đó.',
    'Có quyền sửa hóa đơn và lý do thay đổi; hóa đơn thuộc phạm vi quản lý và còn đủ điều kiện xử lý.',
    'Có quyền quyết toán hợp đồng; đã nhập ngày trả phòng, chỉ số cuối và các khoản phí cần đối chiếu.',
    'Người thuê có hồ sơ và hợp đồng liên quan; chủ trọ có quyền xét duyệt và thực hiện thay đổi thuê.',
    'Người thuê gửi sự cố của phòng đang thuê; chủ trọ hoặc nhân viên được cấp quyền xử lý tương ứng.',
    'Có quyền xem báo cáo; đã chọn tháng và khu trọ thuộc phạm vi quản lý.',
    'Quản trị viên có quyền xử lý tài khoản khách hàng; có trạng thái tài khoản hiện hành và lý do xử lý.',
    'Quản trị viên có quyền tra cứu; thao tác kiểm duyệt cần quyền riêng và đúng hồ sơ tin đăng.',
    'Quản trị viên có quyền đối soát và sao kê hợp lệ; người kiểm tra đề nghị thủ công khác người lập.',
    'Chủ trọ đã đăng nhập; trợ lý được bật và có tài liệu hướng dẫn thuộc phạm vi được phép đọc.',
    'Quản trị viên có quyền xét duyệt; hồ sơ xác minh và chứng cứ liên quan có thể được tra cứu.',
    'Quản trị viên có quyền quản lý gói; đã chọn chủ trọ, gói, thời hạn và lý do thay đổi.',
]
sequence_heading = next(p for p in doc.paragraphs if p.text.startswith('2.4.4 ') and p.style.style_id == 'Heading3')
following_heading = next(p for p in doc.paragraphs if p.text.startswith('2.4.5 ') and p.style.style_id == 'Heading3')
old_sequence = []
cursor = sequence_heading._p.getnext()
while cursor is not following_heading._p:
    old_sequence.append(cursor)
    cursor = cursor.getnext()
old_refs = {}
old_image_relationships = []
current_key = None
for element in old_sequence:
    text = text_of(element)
    if re.match(r'^SD-\d+', text):
        current_key = text[:5]
    if text.startswith('UC:'):
        old_refs[current_key] = text
    old_image_relationships.extend(element.xpath('.//a:blip/@r:embed'))
    body.remove(element)

intro = [
    'Phần này trình bày 19 kịch bản tuần tự ở mức phân tích nghiệp vụ. Mỗi sơ đồ thể hiện tác nhân, giao diện, xử lý nghiệp vụ và cơ sở dữ liệu. Kiểm tra quyền, điều kiện và tính toán thuộc phần xử lý; cơ sở dữ liệu tiếp nhận yêu cầu tra cứu hoặc lưu và trả kết quả. Các thao tác cùng mục tiêu có thể được tóm tắt, không mô tả chi tiết từng lời gọi trong mã nguồn.',
    'Thời gian diễn tiến từ trên xuống. Mũi tên nét liền thể hiện yêu cầu, nét đứt thể hiện phản hồi; điều kiện được đặt trong ngoặc vuông. Khung alt chọn một nhánh phù hợp, opt thực hiện khi đủ điều kiện, loop biểu diễn lặp và break dừng kịch bản khi có lỗi. Các bước do những người khác nhau thực hiện có thể diễn ra ở các thời điểm khác nhau, không phải chuỗi thao tác tự động.',
    'Sơ đồ giữ màu vàng nhạt, nét đỏ nâu và đường sống xám theo mẫu báo cáo; cơ sở dữ liệu dùng hình trụ. Hình dài được trình bày qua các trang nối tiếp trên khổ A4 ngang, lặp lại tên các thành phần để dễ theo dõi và giữ nguyên thứ tự thông điệp. Các phần nối tiếp thuộc cùng một sơ đồ, không phải kịch bản độc lập. Với đề nghị đối soát thủ công, hai người lập và kiểm tra là hai tài khoản khác nhau cùng vai trò quản trị viên.',
]
for text in intro:
    paragraph = paragraph_before(following_heading._p, text)
    format_paragraph(paragraph)

for index, item in enumerate(sequence_audit['diagrams']):
    key, title = item['id'], item['title']
    heading = paragraph_before(following_heading._p, f'{key} {title}', 'Heading 4')
    heading.paragraph_format.page_break_before = True
    for text in ['Điều kiện trước: ' + preconditions[index], 'Kết quả và giới hạn: ' + outcomes[key]]:
        paragraph = paragraph_before(following_heading._p, text, 'Report Sequence Note')
        paragraph.paragraph_format.keep_with_next = True
    paths = image_slices(key)
    caption_base = f'Hình 2.{52 + index} {key} {title}'
    for part_index, path in enumerate(paths):
        paragraph = paragraph_before(following_heading._p)
        if part_index:
            paragraph.paragraph_format.page_break_before = True
        paragraph.paragraph_format.keep_with_next = True
        format_paragraph(paragraph, 1, WD_ALIGN_PARAGRAPH.CENTER, first=0, line=1, after=0)
        run = paragraph.add_run()
        with Image.open(path) as page_image:
            image_width, image_height = page_image.size
        maximum_height = Pt(370 if part_index == len(paths) - 1 and old_refs.get(key) else 405 if part_index else 330)
        displayed_width = min(LANDSCAPE_WIDTH, int(maximum_height * image_width / image_height))
        shape = run.add_picture(str(path), width=displayed_width)
        shape._inline.docPr.set('descr', f'{key}; phần {part_index + 1}/{len(paths)}; tác nhân, giao diện, xử lý nghiệp vụ và cơ sở dữ liệu')
        caption_text = caption_base + (f' (phần {part_index + 1}/{len(paths)})' if len(paths) > 1 else '')
        caption = paragraph_before(following_heading._p, caption_text, 'Report Figure Caption' if part_index == 0 else 'Report Figure Continuation')
        caption.paragraph_format.keep_together = True
        caption.paragraph_format.keep_with_next = part_index == len(paths) - 1 and bool(old_refs.get(key))
    reference = old_refs.get(key, '')
    if reference:
        prefix, separator, evidence = reference.partition('. Căn cứ:')
        codes = list(dict.fromkeys(re.findall(r'UC-[A-Z]+-\d+(?:\.\d+)?', prefix)))
        reference = 'UC liên quan: ' + ', '.join(codes) + ('. Căn cứ mã nguồn: ' + evidence.strip() if separator else '')
        paragraph_before(following_heading._p, reference, 'Report Source Note')

for rel_id in old_image_relationships:
    if not body.xpath(f'.//a:blip[@r:embed="{rel_id}"]'):
        doc.part.drop_rel(rel_id)

last_section = body.find(qn('w:sectPr'))
portrait_section = copy.deepcopy(last_section)
section_break = paragraph_before(sequence_heading._p)
section_break._p.get_or_add_pPr().append(portrait_section)
format_paragraph(section_break, 1, first=0, line=1, after=0)
landscape_section = copy.deepcopy(last_section)
page_size = child(landscape_section, 'w:pgSz')
page_size.set(qn('w:w'), str(Cm(29.7).twips))
page_size.set(qn('w:h'), str(Cm(21).twips))
page_size.set(qn('w:orient'), 'landscape')
for section in [landscape_section, last_section]:
    numbering = child(section, 'w:pgNumType')
    numbering.attrib.pop(qn('w:start'), None)
    numbering.set(qn('w:fmt'), 'decimal')
section_break = paragraph_before(following_heading._p)
section_break._p.get_or_add_pPr().append(landscape_section)
format_paragraph(section_break, 1, first=0, line=1, after=0)
following_heading.paragraph_format.page_break_before = False

for shape in doc.inline_shapes:
    if shape._inline.docPr.get('descr', '').startswith('SD-'):
        continue
    maximum_width = PORTRAIT_WIDTH
    maximum_height = Cm(20.5)
    if shape.width > maximum_width or shape.height > maximum_height:
        ratio = min(maximum_width / shape.width, maximum_height / shape.height)
        shape.width, shape.height = round(shape.width * ratio), round(shape.height * ratio)

contents_title = next(p for p in doc.paragraphs if p.text == 'MỤC LỤC')
assignment = next(p for p in doc.paragraphs if p.text == 'Phân công thực hiện đồ án' and p.style.style_id == 'Heading1')
cursor = contents_title._p
while cursor is not assignment._p:
    following = cursor.getnext()
    body.remove(cursor)
    cursor = following
assignment.paragraph_format.page_break_before = True
for paragraph in doc.paragraphs:
    text = paragraph.text.strip()
    if paragraph.style.style_id in ['Heading1', 'Heading2', 'Heading3']:
        anchor = f'ReportHeading{len(manifest["toc"]) + 1:03}'
        bookmark(paragraph, anchor)
        manifest['toc'].append({'title': text, 'level': int(paragraph.style.style_id[-1]), 'anchor': anchor})
    if paragraph.style.name == 'Report Figure Caption':
        anchor = f'ReportFigure{len(manifest["figures"]) + 1:03}'
        bookmark(paragraph, anchor)
        manifest['figures'].append({'title': text, 'anchor': anchor})

title = paragraph_before(assignment._p, 'MỤC LỤC', 'Report Contents Title')
field_start = paragraph_before(assignment._p)
format_paragraph(field_start, 1, first=0, line=1, after=0)
field_run(field_start, 'begin')
field_run(field_start, 'instruction', ' TOC \\o "1-3" \\h \\z ')
field_run(field_start, 'separate')
for entry in manifest['toc']:
    last_entry = toc_entry(assignment._p, entry, f'TOC {entry["level"]}')
field_run(last_entry, 'end')
title = paragraph_before(assignment._p, 'DANH MỤC HÌNH', 'Report Contents Title')
title.paragraph_format.page_break_before = True
field_start = paragraph_before(assignment._p)
format_paragraph(field_start, 1, first=0, line=1, after=0)
field_run(field_start, 'begin')
field_run(field_start, 'instruction', ' TOC \\h \\z \\t "Report Figure Caption,1" ')
field_run(field_start, 'separate')
for entry in manifest['figures']:
    last_entry = toc_entry(assignment._p, entry, 'Report Figure List')
field_run(last_entry, 'end')

settings = doc.settings.element
child(settings, 'w:updateFields').set(qn('w:val'), 'false')
doc.core_properties.title = 'Phân tích và thiết kế hệ thống quản lý nhà trọ'
doc.core_properties.subject = 'Báo cáo đồ án chuyên ngành'
doc.core_properties.comments = ''
assert len(doc.tables) == original_table_count
assert [text_of(element) for element in body if element.tag == qn('w:tbl')] == original_texts
assert len(manifest['figures']) == 101, len(manifest['figures'])
assert len(manifest['diagrams']) == 19
assert len(body.xpath('.//w:sectPr')) == 5
assert len(body.xpath('.//wp:docPr[starts-with(@descr,"SD-")]')) == sum(item['parts'] for item in manifest['diagrams'])
assert len([p for p in doc.paragraphs if re.match(r'^SD-\d+', p.text)]) == 19
assert any('Phần này trình bày 19 kịch bản' in p.text for p in doc.paragraphs)
assert hashlib.sha256(SOURCE.read_bytes()).hexdigest() == manifest['input_sha256']
doc.save(OUTPUT)
(BASE / 'manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2))
print(json.dumps({'output': str(OUTPUT), 'tables': original_table_count, 'toc_entries': len(manifest['toc']), 'figure_entries': len(manifest['figures']), 'sequence_pages': sum(item['parts'] for item in manifest['diagrams'])}, ensure_ascii=False, indent=2))
