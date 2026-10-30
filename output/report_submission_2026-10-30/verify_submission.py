import hashlib
import json
import math
import re
import unicodedata
from pathlib import Path
from zipfile import ZipFile

import numpy as np
import pdfplumber
from docx import Document
from docx.oxml.ns import qn
from lxml import etree
from PIL import Image
from pypdf import PdfReader

BASE = Path(__file__).resolve().parent
ROOT = BASE.parent.parent
OUTPUT = BASE / 'Report_final_BA_Da_chinh_form_2026-10-30.docx'
PDF = Path('/tmp/report-final-format-2026-10-30/approved-final') / OUTPUT.with_suffix('.pdf').name
manifest = json.loads((BASE / 'manifest.json').read_text())
source = Path(manifest['input'])
NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}

def normalized(text):
    return re.sub(r'\s+', ' ', unicodedata.normalize('NFC', text)).strip()

def table_texts(document):
    return [''.join(table._tbl.xpath('.//w:t/text()')) for table in document.tables]

original = Document(source)
document = Document(OUTPUT)
assert hashlib.sha256(source.read_bytes()).hexdigest() == manifest['input_sha256']
assert len(document.tables) == 198
assert table_texts(original) == table_texts(document)
assert len(manifest['toc']) == 37 and len(manifest['figures']) == 101
assert len(document.sections) == 5
with ZipFile(OUTPUT) as package:
    assert package.testzip() is None
    xml = etree.fromstring(package.read('word/document.xml'))
    assert not xml.xpath('//w:ins|//w:del|//w:commentReference', namespaces=NS)
    assert not any('INCLUDEPICTURE' in value for value in xml.xpath('//w:instrText/text()', namespaces=NS))
    index_values = {}
    for paragraph in xml.xpath('//w:body/w:p', namespaces=NS):
        anchors = paragraph.xpath('w:hyperlink/@w:anchor', namespaces=NS)
        if anchors:
            index_values[anchors[0]] = paragraph.xpath('w:r/w:t/text()', namespaces=NS)[-1]
    rels = etree.fromstring(package.read('word/_rels/document.xml.rels'))
    rel_map = {item.get('Id'): item.get('Target') for item in rels}
    namespace = {'wp': 'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing', 'a': 'http://schemas.openxmlformats.org/drawingml/2006/main', 'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
    inserted = xml.xpath('//wp:inline[wp:docPr[starts-with(@descr,"SD-")]]', namespaces=namespace)
    assert len(inserted) == sum(entry['parts'] for entry in manifest['diagrams'])
    for inline in inserted:
        description = inline.find('wp:docPr', namespace).get('descr')
        key, part = re.match(r'(SD-\d+); phần (\d+)', description).groups()
        relationship = inline.xpath('.//a:blip/@r:embed', namespaces=namespace)[0]
        embedded = package.read('word/' + rel_map[relationship])
        crop_path = BASE / 'diagram_pages' / f'{key}_phan_{int(part):02}.png'
        assert hashlib.sha256(embedded).digest() == hashlib.sha256(crop_path.read_bytes()).digest()

for entry in manifest['diagrams']:
    key = entry['id']
    image_path = ROOT / 'output/sequence_ba_report_2026-10-30/PNG' / f'Sequence_Diagram_{key}.png'
    assert hashlib.sha256(image_path.read_bytes()).hexdigest() == entry['source_sha256']
    image = Image.open(image_path).convert('RGB')
    width, height = image.size
    assert entry['cuts'][0] == 0 and entry['cuts'][-1] == height
    svg = etree.parse(str(ROOT / 'output/sequence_ba_report_2026-10-30/source' / f'{key}.svg'))
    ratio = width / float(svg.getroot().get('width').removesuffix('px'))
    for node in svg.xpath('//*[local-name()="text"]'):
        if not (node.text or '').strip():
            continue
        baseline = float(node.get('y', '0'))
        font_size = float(node.get('font-size', '42.5'))
        top = math.floor((baseline - font_size * 0.95) * ratio)
        bottom = math.ceil((baseline + font_size * 0.22) * ratio)
        assert not any(top - 4 <= cut <= bottom + 4 for cut in entry['cuts'][1:-1]), (key, node.text)
    for index, (top, bottom) in enumerate(zip(entry['cuts'], entry['cuts'][1:])):
        crop = Image.open(BASE / 'diagram_pages' / f'{key}_phan_{index + 1:02}.png').convert('RGB')
        offset = entry['repeated_header_pixels'] + 12 if index else 0
        assert np.array_equal(np.asarray(crop)[offset:], np.asarray(image)[top:bottom])
        if index:
            assert np.array_equal(np.asarray(crop)[:entry['repeated_header_pixels']], np.asarray(image)[:entry['repeated_header_pixels']])

reader = PdfReader(PDF)
assert len(reader.pages) == manifest['pages']
texts = [normalized(page.extract_text()) for page in reader.pages]
index_errors = []
for entry in manifest['toc'] + manifest['figures']:
    physical = entry['physical_page']
    if normalized(entry['title']) not in texts[physical - 1] or index_values[entry['anchor']] != entry['page']:
        index_errors.append(entry)
assert not index_errors, index_errors
assert all(text for text in texts)
layout_warnings = []
with pdfplumber.open(PDF) as pdf:
    for index, page in enumerate(pdf.pages):
        if page.rotation:
            continue
        for char in page.chars:
            if not char['text'].strip():
                continue
            if char['x0'] < 20 or char['x1'] > page.width - 20 or char['top'] < 12 or char['bottom'] > page.height - 12:
                layout_warnings.append({'page': index + 1, 'text': char['text'], 'box': [char['x0'], char['top'], char['x1'], char['bottom']]})
        text = texts[index]
        if text.startswith('UC liên quan:'):
            layout_warnings.append({'page': index + 1, 'issue': 'isolated_sequence_reference'})
        if len(text) < 20 and not page.images:
            layout_warnings.append({'page': index + 1, 'issue': 'empty_or_heading_only'})
        if page.width > page.height:
            if index + 1 >= manifest['main_starts_at_physical_page']:
                assert page.images or 'Phần này trình bày 19 kịch bản' in text
        page.close()

assert not layout_warnings, layout_warnings[:20]
result = {
    'original_unchanged': True,
    'tables_preserved': 198,
    'sequence_diagrams': 19,
    'sequence_image_parts': len(inserted),
    'image_content_and_safe_cuts_verified': True,
    'toc_entries_verified': 37,
    'figure_entries_verified': 101,
    'rendered_pages': len(reader.pages),
    'layout_boundary_warnings': layout_warnings,
    'docx_sha256': hashlib.sha256(OUTPUT.read_bytes()).hexdigest(),
    'rendered_pdf': str(PDF),
}
(BASE / 'verification.json').write_text(json.dumps(result, ensure_ascii=False, indent=2))
print(json.dumps(result, ensure_ascii=False, indent=2))
