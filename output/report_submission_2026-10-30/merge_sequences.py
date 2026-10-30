import copy
import hashlib
import json
import math
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from PIL import Image, ImageFont

BASE = Path(__file__).resolve().parent
TARGET = BASE.parent / 'report_merged_sequence_2026-10-30'
SOURCE = BASE / 'Report_final_BA_Da_chinh_form_2026-10-30.docx'
OUTPUT = TARGET / 'Report_final_BA_Moi_so_do_1_hinh_2026-10-30.docx'
DIAGRAMS = BASE.parent / 'sequence_ba_report_2026-10-30' / 'PNG'
FONT = '/System/Library/Fonts/Supplemental/Times New Roman.ttf'


def paragraph_height(text, size, spacing, before=0, after=0):
    font = ImageFont.truetype(FONT, round(size * 4))
    lines = 1
    current = ''
    for word in text.split():
        candidate = f'{current} {word}'.strip()
        if current and font.getlength(candidate) > (Cm(16).pt - 10) * 4:
            lines += 1
            current = word
        else:
            current = candidate
    return lines * size * spacing + before + after


def build():
    TARGET.mkdir(exist_ok=True)
    document = Document(SOURCE)
    original_tables = [table._tbl.xml for table in document.tables]
    manifest = json.loads((BASE / 'manifest.json').read_text())
    manifest['edited_from'] = str(SOURCE)
    manifest['edited_from_sha256'] = hashlib.sha256(SOURCE.read_bytes()).hexdigest()
    manifest.pop('pages', None)
    for entry in manifest['toc'] + manifest['figures']:
        entry.pop('physical_page', None)
    for section in document.sections:
        if section.page_width > section.page_height:
            section.page_width = Cm(21)
            section.page_height = Cm(29.7)
            section._sectPr.find(qn('w:pgSz')).attrib.pop(qn('w:orient'), None)
    paragraphs = document.paragraphs
    start = next(index for index, paragraph in enumerate(paragraphs) if paragraph.style.style_id == 'Heading3' and paragraph.text.startswith('2.4.4 '))
    end = next(index for index, paragraph in enumerate(paragraphs) if paragraph.style.style_id == 'Heading3' and paragraph.text.startswith('2.4.5 '))
    section_paragraphs = paragraphs[start + 1:end]
    obsolete_relations = set()
    sequence = None
    groups = {}
    for paragraph in section_paragraphs:
        if paragraph.text.startswith('Sơ đồ giữ màu vàng nhạt'):
            paragraph.runs[0].text = 'Sơ đồ giữ màu vàng nhạt, nét đỏ nâu và đường sống xám theo mẫu báo cáo; cơ sở dữ liệu dùng hình trụ. Mỗi kịch bản được trình bày bằng một hình hoàn chỉnh, không chia thành các phần nối tiếp. Với đề nghị đối soát thủ công, hai người lập và kiểm tra là hai tài khoản khác nhau cùng vai trò quản trị viên.'
            for run in paragraph.runs[1:]:
                run.text = ''
        if re.match(r'^SD-\d+', paragraph.text):
            sequence = paragraph.text[:5]
            groups[sequence] = []
        if sequence:
            groups[sequence].append(paragraph)
    assert len(groups) == 19
    image_records = []
    for key, group in groups.items():
        image_paragraphs = [paragraph for paragraph in group if paragraph._p.xpath('.//a:blip')]
        captions = [paragraph for paragraph in group if paragraph.style.name in ['Report Figure Caption', 'Report Figure Continuation']]
        assert image_paragraphs and captions
        for paragraph in image_paragraphs:
            obsolete_relations.update(paragraph._p.xpath('.//a:blip/@r:embed'))
        first_image, first_caption = image_paragraphs[0], captions[0]
        for paragraph in image_paragraphs[1:] + captions[1:]:
            paragraph._p.getparent().remove(paragraph._p)
        first_image.clear()
        first_image.paragraph_format.page_break_before = False
        first_image.paragraph_format.keep_with_next = True
        first_image.paragraph_format.space_before = Pt(0)
        first_image.paragraph_format.space_after = Pt(0)
        first_image.paragraph_format.line_spacing = 1
        first_image.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in first_caption.runs:
            run.text = re.sub(r'\s*\(phần \d+/\d+\)', '', run.text)
        figure = next(entry for entry in manifest['figures'] if f' {key} ' in entry['title'])
        figure['title'] = re.sub(r'\s*\(phần \d+/\d+\)', '', figure['title'])
        overhead = 18
        for paragraph in group:
            if paragraph._p.getparent() is None or paragraph is first_image:
                continue
            if paragraph.style.style_id == 'Heading4':
                overhead += paragraph_height(paragraph.text, 13, 1.15, 12, 6)
            elif paragraph.style.name == 'Report Sequence Note':
                overhead += paragraph_height(paragraph.text, 11, 1.12, 0, 4)
            elif paragraph.style.name == 'Report Source Note':
                overhead += paragraph_height(paragraph.text, 10, 1.1, 0, 6)
            elif paragraph is first_caption:
                overhead += paragraph_height(paragraph.text, 11, 1.15, 4, 6)
            if paragraph.text:
                paragraph.paragraph_format.keep_with_next = paragraph.style.name != 'Report Source Note'
                paragraph.paragraph_format.keep_together = True
        image_path = DIAGRAMS / f'Sequence_Diagram_{key}.png'
        with Image.open(image_path) as image:
            width, height = image.size
        maximum_height = min(620, Cm(24.7).pt - overhead)
        displayed_width = min(Cm(16), round(Pt(maximum_height) * width / height))
        shape = first_image.add_run().add_picture(str(image_path), width=displayed_width)
        shape._inline.docPr.set('descr', f'{key}; sơ đồ hoàn chỉnh; tác nhân, giao diện, xử lý nghiệp vụ và cơ sở dữ liệu')
        image_records.append({'id': key, 'parts': 1, 'source_sha256': hashlib.sha256(image_path.read_bytes()).hexdigest(), 'pixels': [width, height], 'display_height_pt': shape.height.pt})
    for relationship in obsolete_relations:
        if not document._element.body.xpath(f'.//a:blip[@r:embed="{relationship}"]'):
            document.part.drop_rel(relationship)
    for paragraph in document.paragraphs:
        anchors = paragraph._p.xpath('w:hyperlink/@w:anchor')
        if anchors:
            for node in paragraph._p.xpath('w:hyperlink//w:t'):
                if node.text:
                    node.text = re.sub(r'\s*\(phần \d+/\d+\)', '', node.text)
    assert [table._tbl.xml for table in document.tables] == original_tables
    assert len(document._element.body.xpath('.//wp:docPr[starts-with(@descr,"SD-")]')) == 19
    assert not any(paragraph.style.name == 'Report Figure Continuation' for paragraph in document.paragraphs)
    manifest['diagrams'] = image_records
    document.save(OUTPUT)
    (TARGET / 'manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2))
    print(json.dumps({'output': str(OUTPUT), 'whole_sequence_figures': len(image_records)}, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    build()
